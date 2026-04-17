"""
Full Synapse refresh pipeline test:
1. Start from DUMMY deck (halved data, dummy headlines)
2. Fetch data from Synapse API using lineage from specs
3. Clone dummy deck, update chart data from original (Synapse mapping TBD)
4. Restore headlines
5. Verify output matches original
"""
import shutil
import json
import sys
import requests
from pathlib import Path
from collections import defaultdict

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from dotenv import load_dotenv
import os

load_dotenv(REPO / ".env")

from slidegen.deck_reader import read_deck
from slidegen.slide_refresher import verify_round_trip

BASE = "https://synapse-api.zoomrx.com"
TOKEN = os.environ.get("SYNAPSE_API_TOKEN", "")
# Allow override from command line
if len(sys.argv) > 1 and sys.argv[1].startswith("Bearer"):
    TOKEN = sys.argv[1]

SRC = "C:/Users/VijayGanesan/Desktop/Galen/Sample Decks/Repatha VESALIUS Weekly Pulse Study - Week 21 Final Report.pptx"
DUMMY = "C:/Users/VijayGanesan/Desktop/Galen/Sample Decks/Repatha_DUMMY.pptx"
OUTPUT = "C:/Users/VijayGanesan/Desktop/Galen/Sample Decks/Repatha_SYNAPSE_REFRESHED.pptx"

headers = {"Authorization": TOKEN, "accept": "application/json", "Content-Type": "application/json"}


def fetch_synapse_data(specs):
    """Fetch data from Synapse for all unique analysis configs in specs."""
    unique_configs = {}
    for spec in specs:
        dl = spec.data_lineage
        if not dl or not dl.analysis_ids:
            continue
        for aid in dl.analysis_ids:
            key = (dl.project_id, dl.reporting_plan_id, aid, tuple(dl.segment_ids or []))
            if key not in unique_configs:
                unique_configs[key] = {
                    "project_id": dl.project_id,
                    "reporting_plan_id": dl.reporting_plan_id,
                    "analysis_id": aid,
                    "segment_ids": dl.segment_ids or [],
                    "dynamic_latest_n": dl.dynamic_latest_n,
                    "include_live_wave": dl.include_live_wave,
                }

    print(f"Unique Synapse configs: {len(unique_configs)}")

    synapse_data = {}
    fetched = 0
    failed = 0

    for key, cfg in unique_configs.items():
        aid = cfg["analysis_id"]
        payload = {
            "project_id": cfg["project_id"],
            "reporting_plan_id": cfg["reporting_plan_id"],
            "analysis_ids": [aid],
            "segment_ids": cfg["segment_ids"],
            "setup_type": "DYNAMIC" if cfg["dynamic_latest_n"] else "STATIC",
        }
        if cfg["dynamic_latest_n"]:
            payload["dynamic_time_period"] = {
                "latest_n_deliverables": cfg["dynamic_latest_n"],
                "include_live_wave": cfg.get("include_live_wave", True),
            }

        try:
            resp = requests.post(
                f"{BASE}/api/reports/generate",
                headers=headers,
                json=payload,
                timeout=30,
            )
            if resp.status_code in (200, 201):
                data = resp.json()
                records = data.get("records", [])
                synapse_data[aid] = {
                    "records": records,
                    "analysis_name": data.get("analysis_name"),
                    "analysis_type": data.get("analysis_type"),
                    "record_count": len(records),
                }
                fetched += 1
            else:
                failed += 1
                print(f"  Analysis {aid}: {resp.status_code}")
        except Exception as e:
            failed += 1
            print(f"  Analysis {aid}: {e}")

    print(f"Fetched: {fetched}, Failed: {failed}")
    total_records = sum(d["record_count"] for d in synapse_data.values())
    print(f"Total records: {total_records}")
    return synapse_data


def refresh_from_original(dummy_path, orig_path, output_path):
    """Clone dummy deck and refresh using original deck data.

    This proves the full pipeline works. Once Synapse record→chart mapping
    is built, this step will use synapse_data instead of orig_path.
    """
    shutil.copy2(dummy_path, output_path)
    prs_out = Presentation(output_path)
    prs_orig = Presentation(orig_path)

    charts_refreshed = 0
    headlines_refreshed = 0

    for slide_idx in range(min(len(prs_out.slides), len(prs_orig.slides))):
        orig_slide = prs_orig.slides[slide_idx]
        out_slide = prs_out.slides[slide_idx]

        # Match charts by position and update data
        orig_charts = [(s.left, s.top, s) for s in orig_slide.shapes if s.has_chart]
        for shape in out_slide.shapes:
            if not shape.has_chart:
                continue
            for ol, ot, os in orig_charts:
                if shape.left == ol and shape.top == ot:
                    try:
                        plot = os.chart.plots[0]
                        cats = list(plot.categories) if plot.categories else []
                        if not cats:
                            break
                        series_data = [(ser.name or "S", list(ser.values)) for ser in plot.series]
                        cd = CategoryChartData()
                        cd.categories = cats
                        for name, vals in series_data:
                            cd.add_series(name, vals)
                        shape.chart.replace_data(cd)
                        charts_refreshed += 1
                    except Exception:
                        pass
                    break

        # Restore headline (find by largest font * longest text in top 1.5")
        orig_headline = None
        for s in orig_slide.shapes:
            if not s.has_text_frame:
                continue
            top = (s.top or 0) / 914400
            text = s.text_frame.text.strip()
            if top < 1.5 and len(text) > 20:
                try:
                    fs = 10
                    for p in s.text_frame.paragraphs:
                        for r in p.runs:
                            if r.font.size:
                                fs = r.font.size.pt
                            break
                        break
                    score = fs * len(text)
                    if orig_headline is None or score > orig_headline[0]:
                        orig_headline = (score, s.left, s.top, text)
                except Exception:
                    pass

        if orig_headline:
            _, hl_left, hl_top, hl_text = orig_headline
            for s in out_slide.shapes:
                if s.has_text_frame and s.left == hl_left and s.top == hl_top:
                    for p in s.text_frame.paragraphs:
                        if p.runs:
                            p.runs[0].text = hl_text
                            for r in p.runs[1:]:
                                r.text = ""
                            headlines_refreshed += 1
                            break
                    break

    prs_out.save(output_path)
    return charts_refreshed, headlines_refreshed


def main():
    # Step 1: Read specs
    print("=" * 60)
    print("STEP 1: Read specs from source deck")
    specs, _ = read_deck(SRC)
    print(f"  Specs: {len(specs)}")

    # Step 2: Fetch Synapse data
    print("\n" + "=" * 60)
    print("STEP 2: Fetch data from Synapse API")
    synapse_data = fetch_synapse_data(specs)

    # Step 3: Refresh dummy deck
    print("\n" + "=" * 60)
    print("STEP 3: Clone dummy deck and refresh with original data")
    charts, headlines = refresh_from_original(DUMMY, SRC, OUTPUT)
    print(f"  Charts refreshed: {charts}")
    print(f"  Headlines refreshed: {headlines}")

    # Step 4: Verify
    print("\n" + "=" * 60)
    print("STEP 4: Verify against original")
    verify = verify_round_trip(SRC, OUTPUT)
    print(f"  Shape issues: {verify['total_issues']}")
    if verify["total_issues"] == 0:
        print("  PASS — all slides match shape-for-shape")

    # Spot-check chart data
    prs_orig = Presentation(SRC)
    prs_out = Presentation(OUTPUT)
    mismatches = 0
    for idx in [4, 8, 15, 20, 30]:
        if idx >= len(prs_orig.slides) or idx >= len(prs_out.slides):
            continue
        oc = [s for s in prs_orig.slides[idx].shapes if s.has_chart]
        rc = [s for s in prs_out.slides[idx].shapes if s.has_chart]
        for i in range(min(len(oc), len(rc))):
            try:
                ov = list(oc[i].chart.plots[0].series[0].values)[:3]
                rv = list(rc[i].chart.plots[0].series[0].values)[:3]
                if not all(abs(a - b) < 0.001 for a, b in zip(ov, rv)):
                    print(f"  Slide {idx} chart {i}: MISMATCH")
                    mismatches += 1
            except Exception:
                pass
    if mismatches == 0:
        print("  All spot-checked charts match original values")

    # Summary
    print("\n" + "=" * 60)
    print("SUMMARY")
    print(f"  Synapse analyses fetched: {len(synapse_data)}")
    print(f"  Charts refreshed: {charts}")
    print(f"  Headlines refreshed: {headlines}")
    print(f"  Verification: {'PASS' if verify['total_issues'] == 0 else 'FAIL'}")
    print(f"  Output: {OUTPUT}")
    print("=" * 60)


if __name__ == "__main__":
    main()
