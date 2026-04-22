"""
test_intelligent_refresh.py — Intelligent slide refresh orchestrated by Claude Code.

Two-phase pipeline:
  Phase 1 (--read):  Extract slide context + fetch data → print structured summary
                     Claude Code reads this output and reasons about mappings
  Phase 2 (--refresh mapping.json):  Apply Claude Code's mapping to refresh the deck

The interpretation step happens in Claude Code itself — no API call needed.
Claude reads the slide layout, sees which labels are near which charts,
and writes the mapping JSON.

Usage:
    # Phase 1: Extract (Claude Code reads the output)
    python tests/test_intelligent_refresh.py --read

    # Phase 2: Refresh (Claude Code provides the mapping)
    python tests/test_intelligent_refresh.py --refresh path/to/mapping.json

    # Or all-in-one with a mapping file:
    python tests/test_intelligent_refresh.py --read --refresh path/to/mapping.json
"""
from __future__ import annotations

import json
import os
import shutil
import sys
from dataclasses import dataclass, field, asdict
from pathlib import Path

import pandas as pd
import requests
from dotenv import load_dotenv
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))
load_dotenv(REPO_ROOT / ".env")

SAMPLE_DIR = REPO_ROOT.parent / "Sample Decks"
SOURCE_PPTX = SAMPLE_DIR / "Repatha ATU Slide 6.pptx"
OUTPUT_PPTX = SAMPLE_DIR / "Repatha_ATU_Slide6_intelligent_refresh.pptx"

# User-provided data lineage
DATA_LINEAGE = {
    "project_id": 1428,
    "project_name": "Amgen [ATU]: Repatha",
    "reporting_plan_id": 574,
    "analysis_ids": [545991],
    "segment_ids": [14633],
    "dynamic_latest_n": 5,
}

_CHART_TYPE_MAP = {
    XL_CHART_TYPE.BAR_STACKED_100: "bar_stacked_100_horizontal",
    XL_CHART_TYPE.BAR_CLUSTERED: "bar_clustered_horizontal",
    XL_CHART_TYPE.COLUMN_CLUSTERED: "column_clustered_vertical",
    XL_CHART_TYPE.COLUMN_STACKED_100: "column_stacked_100_vertical",
    XL_CHART_TYPE.LINE_MARKERS: "line_markers_trended",
    XL_CHART_TYPE.LINE: "line_markers_trended",
    XL_CHART_TYPE.XY_SCATTER: "xy_scatter_abacus",
    XL_CHART_TYPE.DOUGHNUT: "doughnut_default",
}


# ══════════════════════════════════════════════════════════════════════════════
# Phase 1: Read — extract slide context + fetch data
# ══════════════════════════════════════════════════════════════════════════════

def phase_read(pptx_path: str, data_lineage: dict) -> dict:
    """Extract slide layout + fetch Synapse data. Returns structured context.

    Claude Code reads this output and decides how each chart/table maps to the data.
    """
    prs = Presentation(pptx_path)
    slide = prs.slides[0]

    def _in(emu):
        return round(emu / 914400, 2) if emu else 0.0

    # ── Extract all shapes ──
    shapes = []
    for shape in slide.shapes:
        left, top = _in(shape.left), _in(shape.top)
        width, height = _in(shape.width), _in(shape.height)
        base = {"name": shape.name, "left": left, "top": top,
                "width": width, "height": height}

        if shape.has_chart:
            chart = shape.chart
            plot = chart.plots[0]
            cats = [str(c) for c in plot.categories] if plot.categories else []
            s_names = [str(s.name) for s in plot.series]
            s_vals = {}
            for s in plot.series:
                try:
                    s_vals[str(s.name)] = [round(float(v), 4) if v else 0 for v in s.values][:4]
                except Exception:
                    pass
            shapes.append({**base, "type": "chart",
                           "chart_pattern": _CHART_TYPE_MAP.get(chart.chart_type, "unknown"),
                           "series_names": s_names, "categories": cats,
                           "series_values": s_vals})

        elif shape.has_table:
            tbl = shape.table
            nr, nc = len(tbl.rows), len(tbl.columns)
            headers = [tbl.cell(0, c).text.strip() for c in range(nc)] if nr > 0 else []
            sample = []
            for r in range(1, min(nr, 4)):
                sample.append([tbl.cell(r, c).text.strip() for c in range(nc)])
            shapes.append({**base, "type": "table",
                           "row_count": nr, "col_count": nc,
                           "headers": headers, "sample_rows": sample})

        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                shapes.append({**base, "type": "text", "text": text[:300]})

        # Group shape children with text
        try:
            if hasattr(shape, 'shapes'):
                for child in shape.shapes:
                    if child.has_text_frame:
                        t = child.text_frame.text.strip()
                        if t:
                            shapes.append({
                                "name": f"{shape.name}/{child.name}",
                                "left": left, "top": top,
                                "width": width, "height": height,
                                "type": "label", "text": t[:200],
                            })
        except Exception:
            pass

    shapes.sort(key=lambda s: (s["left"], s["top"]))

    # ── Fetch Synapse data ──
    token = os.environ.get("SYNAPSE_API_TOKEN", "")
    base_url = os.environ.get("SYNAPSE_API_BASE_URL", "https://synapse-api.zoomrx.com")
    headers = {
        "Authorization": token if token.startswith("Bearer") else f"Bearer {token}",
        "accept": "application/json", "Content-Type": "application/json",
    }
    payload = {
        "project_id": data_lineage["project_id"],
        "reporting_plan_id": data_lineage["reporting_plan_id"],
        "analysis_ids": data_lineage["analysis_ids"],
        "segment_ids": data_lineage.get("segment_ids", []),
        "setup_type": "DYNAMIC",
        "dynamic_time_period": {
            "latest_n_deliverables": data_lineage.get("dynamic_latest_n", 5),
            "include_live_wave": True,
        },
    }
    resp = requests.post(f"{base_url}/api/reports/generate",
                         headers=headers, json=payload, timeout=60)
    records = []
    if resp.status_code in (200, 201):
        records = resp.json().get("records", [])

    # Summarize data
    df = pd.DataFrame(records) if records else pd.DataFrame()
    data_summary = {"record_count": len(records), "columns": sorted(df.columns.tolist()) if not df.empty else []}
    if not df.empty:
        for col in df.columns:
            unique = df[col].dropna().unique()
            if len(unique) <= 20:
                data_summary[col] = sorted(str(v) for v in unique)

    return {"shapes": shapes, "data_summary": data_summary, "data_lineage": data_lineage}


# ══════════════════════════════════════════════════════════════════════════════
# Phase 2: Refresh — apply mapping to update the deck
# ══════════════════════════════════════════════════════════════════════════════

def phase_refresh(pptx_path: str, output_path: str, mapping: dict, data_lineage: dict) -> dict:
    """Apply Claude Code's mapping to refresh the deck.

    mapping format:
    {
      "charts": [
        {"chart_name": "Chart 35", "segment_filter": "CARD",
         "series_column": "x_code", "series_name_map": {"IDK": "0"},
         "row_field": "y_label", "value_field": "percentage"}
      ],
      "tables": [
        {"table_name": "Table 25", "segment_filter": "CARD",
         "columns": [
           {"header": "Product", "data_field": "y_label", "format": "string"},
           {"header": "Base", "data_field": "base", "format": "(n = {})"},
           {"header": "", "data_field": null, "format": "spacer"},
           {"header": "Easy", "data_field": "percentage", "format": "{}%",
            "filter": {"x_code": "H"}}
         ]}
      ]
    }
    """
    # Fetch data
    token = os.environ.get("SYNAPSE_API_TOKEN", "")
    base_url = os.environ.get("SYNAPSE_API_BASE_URL", "https://synapse-api.zoomrx.com")
    headers = {
        "Authorization": token if token.startswith("Bearer") else f"Bearer {token}",
        "accept": "application/json", "Content-Type": "application/json",
    }
    payload = {
        "project_id": data_lineage["project_id"],
        "reporting_plan_id": data_lineage["reporting_plan_id"],
        "analysis_ids": data_lineage["analysis_ids"],
        "segment_ids": data_lineage.get("segment_ids", []),
        "setup_type": "DYNAMIC",
        "dynamic_time_period": {
            "latest_n_deliverables": data_lineage.get("dynamic_latest_n", 5),
            "include_live_wave": True,
        },
    }
    resp = requests.post(f"{base_url}/api/reports/generate",
                         headers=headers, json=payload, timeout=60)
    if resp.status_code not in (200, 201):
        return {"error": f"Synapse API {resp.status_code}"}
    records = resp.json().get("records", [])
    df = pd.DataFrame(records)

    # Read source for series order
    src_prs = Presentation(pptx_path)
    src_chart_series = {}
    for shape in src_prs.slides[0].shapes:
        if shape.has_chart:
            src_chart_series[shape.name] = [str(s.name) for s in shape.chart.plots[0].series]

    # Clone + refresh
    shutil.copy2(pptx_path, output_path)
    prs = Presentation(output_path)
    slide = prs.slides[0]
    chart_shapes = {s.name: s for s in slide.shapes if s.has_chart}
    table_shapes = {s.name: s for s in slide.shapes if s.has_table}
    ns_c = "http://schemas.openxmlformats.org/drawingml/2006/chart"

    results = {"charts": [], "tables": []}

    # ── Charts ──
    for cm in mapping.get("charts", []):
        name = cm["chart_name"]
        shape = chart_shapes.get(name)
        if not shape:
            results["charts"].append({"name": name, "status": "not_found"})
            continue

        sub = df.copy()
        seg = cm.get("segment_filter")
        if seg and "segment_1" in sub.columns:
            sub = sub[sub["segment_1"] == seg]
        tp = cm.get("time_period")
        if tp and "time_period_name" in sub.columns:
            sub = sub[sub["time_period_name"] == tp]
        if sub.empty:
            results["charts"].append({"name": name, "status": "empty"})
            continue

        row_f = cm.get("row_field", "y_label")
        col_f = cm.get("series_column")
        val_f = cm.get("value_field", "percentage")
        name_map = cm.get("series_name_map", {})

        try:
            if col_f and col_f in sub.columns:
                pivot = sub.pivot_table(index=row_f, columns=col_f, values=val_f, aggfunc="first")
            else:
                pivot = sub.groupby(row_f, sort=False)[val_f].first().to_frame()

            categories = list(pivot.index.astype(str))
            orig_series = src_chart_series.get(name, [])

            series_data = []
            for sn in orig_series:
                data_val = name_map.get(sn, sn)
                matched = None
                for pc in pivot.columns:
                    if str(pc) == data_val or str(pc).lower() == data_val.lower():
                        matched = pc
                        break
                vals = [0.0 if pd.isna(v) else round(float(v), 4) for v in pivot[matched]] if matched is not None else [0.0] * len(categories)
                series_data.append((sn, vals))

            src_fmts = [fc.text for fc in shape.chart._chartSpace.iter(f"{{{ns_c}}}formatCode")]
            cd = CategoryChartData()
            cd.categories = categories
            for sn, vals in series_data:
                cd.add_series(sn, vals)
            shape.chart.replace_data(cd)
            for i, fc_el in enumerate(shape.chart._chartSpace.iter(f"{{{ns_c}}}formatCode")):
                if i < len(src_fmts):
                    fc_el.text = src_fmts[i]

            results["charts"].append({"name": name, "status": "ok", "segment": seg,
                                      "categories": categories[:3], "n_series": len(series_data)})
        except Exception as e:
            results["charts"].append({"name": name, "status": "error", "error": str(e)})

    # ── Tables ──
    for tm in mapping.get("tables", []):
        name = tm["table_name"]
        shape = table_shapes.get(name)
        if not shape:
            results["tables"].append({"name": name, "status": "not_found"})
            continue

        tbl = shape.table
        sub = df.copy()
        seg = tm.get("segment_filter")
        if seg and "segment_1" in sub.columns:
            sub = sub[sub["segment_1"] == seg]
        if sub.empty:
            results["tables"].append({"name": name, "status": "empty"})
            continue

        columns_spec = tm.get("columns", [])
        if not columns_spec:
            results["tables"].append({"name": name, "status": "no_columns"})
            continue

        try:
            row_key = None
            for cs in columns_spec:
                if cs.get("data_field") and cs.get("format") != "spacer":
                    row_key = cs["data_field"]
                    break
            if not row_key or row_key not in sub.columns:
                results["tables"].append({"name": name, "status": "no_row_key"})
                continue

            keys = sub[row_key].dropna().unique().tolist()
            n_rows = min(len(keys), len(tbl.rows) - 1)

            for ri in range(n_rows):
                row_data = sub[sub[row_key] == keys[ri]]
                if row_data.empty:
                    continue
                for ci, cs in enumerate(columns_spec):
                    if ci >= len(tbl.columns) or ri + 1 >= len(tbl.rows):
                        break
                    fmt = cs.get("format", "string")
                    data_f = cs.get("data_field")
                    col_filter = cs.get("filter")
                    if fmt == "spacer" or not data_f:
                        continue

                    cell_data = row_data
                    if col_filter and isinstance(col_filter, dict):
                        for fk, fv in col_filter.items():
                            if fk in cell_data.columns:
                                cell_data = cell_data[cell_data[fk] == fv]
                    if cell_data.empty or data_f not in cell_data.columns:
                        continue

                    raw = cell_data.iloc[0][data_f]
                    if fmt == "string":
                        text = str(raw)
                    elif "(n" in fmt:
                        text = f"(n = {int(float(raw))})" if raw else ""
                    elif "%" in fmt:
                        v = float(raw)
                        text = f"{v:.0%}" if abs(v) <= 1.0 else f"{int(v)}%"
                    else:
                        text = str(raw)

                    cell = tbl.cell(ri + 1, ci)
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.text = ""
                    if cell.text_frame.paragraphs and cell.text_frame.paragraphs[0].runs:
                        cell.text_frame.paragraphs[0].runs[0].text = text

            results["tables"].append({"name": name, "status": "ok", "segment": seg, "rows": n_rows})
        except Exception as e:
            results["tables"].append({"name": name, "status": "error", "error": str(e)})

    prs.save(output_path)
    return results


# ══════════════════════════════════════════════════════════════════════════════
# CLI
# ══════════════════════════════════════════════════════════════════════════════

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Intelligent slide refresh")
    parser.add_argument("--read", action="store_true", help="Phase 1: extract slide context + fetch data")
    parser.add_argument("--refresh", type=str, metavar="MAPPING_JSON", help="Phase 2: refresh with mapping file")
    parser.add_argument("--pptx", type=str, default=str(SOURCE_PPTX), help="Source PPTX path")
    parser.add_argument("--output", type=str, default=str(OUTPUT_PPTX), help="Output PPTX path")
    args = parser.parse_args()

    if args.read:
        print(json.dumps(phase_read(args.pptx, DATA_LINEAGE), indent=2, ensure_ascii=False))

    if args.refresh:
        mapping = json.loads(Path(args.refresh).read_text(encoding="utf-8"))
        results = phase_refresh(args.pptx, args.output, mapping, DATA_LINEAGE)
        print(json.dumps(results, indent=2, ensure_ascii=False))

    if not args.read and not args.refresh:
        parser.print_help()


if __name__ == "__main__":
    main()
