"""
Test Synapse record → chart data mapping using PivotConfig + MappingConfig.
Fetches data from Synapse API, pivots it, and compares with actual chart values.
"""
import json
import sys
import os
import requests
import pandas as pd
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

from dotenv import load_dotenv
load_dotenv(REPO / ".env")

from slidegen.deck_reader.tag_reader import (
    _get_shape_tags_all, _parse_custom_xml_parts,
    TAG_REPORT_CONFIG_HASH, TAG_PIVOT_CONFIG_HASH, TAG_MAPPING_CONFIG,
    XML_STORE_REPORT_CONFIG, XML_STORE_PIVOT_CONFIG,
)
from pptx import Presentation

BASE = "https://synapse-api.zoomrx.com"
TOKEN = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("SYNAPSE_API_TOKEN", "")
headers = {"Authorization": TOKEN, "accept": "application/json", "Content-Type": "application/json"}

SRC = "C:/Users/VijayGanesan/Desktop/Galen/Sample Decks/Repatha VESALIUS Weekly Pulse Study - Week 21 Final Report.pptx"


def pivot_synapse_records(records, pivot_config, mapping_config):
    """Transform flat Synapse records into chart categories + series using pivot/mapping configs.

    Returns (categories: list[str], series: list[(name, values: list[float])]) or None on failure.
    """
    if not records:
        return None

    df = pd.DataFrame(records)

    row_fields = pivot_config.get("RowFields", [])
    col_fields = pivot_config.get("ColumnFields", [])
    val_fields = pivot_config.get("ValueFields", [])

    if not row_fields or not val_fields:
        return None

    # Field name mapping: Connector PivotConfig uses internal names that
    # don't match Synapse API record columns. Known translations:
    FIELD_MAP = {
        "value": "y_label",       # Connector "value" = API "y_label" (answer/option text)
        "y_code": "y_code",       # code already matches
        "segment_1": "segment_1", # already matches
        "segment_2": "segment_2",
        "time_period_name": "time_period_name",
    }

    # Remap field names in row/col/val fields
    def remap(field):
        return FIELD_MAP.get(field, field)

    row_fields = [remap(f) for f in row_fields]
    col_fields = [remap(f) for f in col_fields]
    val_fields = [remap(f) for f in val_fields]

    # Apply filters (also remap filter column keys)
    for f in pivot_config.get("Filters", []):
        col_key = remap(f.get("ColumnKey", ""))
        val = f.get("Value")
        if col_key and val and col_key in df.columns:
            # Fuzzy match: "Overall Data" might be "Overall" in the API
            if val not in df[col_key].values:
                # Try partial match
                matches = df[df[col_key].str.contains(val.split()[0], case=False, na=False)]
                if not matches.empty:
                    df = matches
            else:
                df = df[df[col_key] == val]

    if df.empty:
        df = pd.DataFrame(records)

    # Verify required columns exist
    for f in row_fields + col_fields + val_fields:
        if f not in df.columns:
            print(f"  WARNING: field '{f}' not in records (available: {list(df.columns)})")
            return None

    # Create combined column key for multi-field pivoting
    if len(col_fields) >= 2:
        df["_col_key"] = df[col_fields[0]].astype(str)
        for cf in col_fields[1:]:
            if cf in df.columns:
                df["_col_key"] = df["_col_key"] + " - " + df[cf].astype(str)
    elif len(col_fields) == 1:
        df["_col_key"] = df[col_fields[0]].astype(str)
    else:
        df["_col_key"] = "value"

    # Pivot
    try:
        pivot = df.pivot_table(
            index=row_fields[0],
            columns="_col_key",
            values=val_fields[0],
            aggfunc="first",
        )
    except Exception as e:
        print(f"  Pivot failed: {e}")
        return None

    # Apply MappingConfig to select which columns become chart series
    selected = mapping_config.get("selectedColumns", [])
    if not selected:
        # No mapping — use all columns
        categories = list(pivot.index)
        series = [(str(col), pivot[col].tolist()) for col in pivot.columns]
        return categories, series

    # First selectedColumn = categories (the row field), rest = series
    categories = list(pivot.index)
    series = []
    for sc in selected[1:]:  # skip first (it's the row/category field)
        # Find matching pivot column (fuzzy match — alias names may differ)
        matched_col = None
        for pc_col in pivot.columns:
            if sc == str(pc_col) or sc in str(pc_col) or str(pc_col) in sc:
                matched_col = pc_col
                break
        if matched_col is not None:
            vals = pivot[matched_col].tolist()
            # Convert NaN to 0
            vals = [0.0 if pd.isna(v) else float(v) for v in vals]
            series.append((sc, vals))

    return categories, series


def main():
    prs = Presentation(SRC)
    xml_configs = _parse_custom_xml_parts(prs, pptx_path=Path(SRC))
    report_configs = xml_configs.get(XML_STORE_REPORT_CONFIG, {})
    pivot_configs = xml_configs.get(XML_STORE_PIVOT_CONFIG, {})

    # Test on all charts across a few slides
    test_slides = [4, 8, 10, 15, 20]
    total_charts = 0
    matched = 0
    mismatched = 0
    no_data = 0

    for slide_idx in test_slides:
        if slide_idx >= len(prs.slides):
            continue
        slide = prs.slides[slide_idx]

        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            tags = _get_shape_tags_all(shape)
            if TAG_REPORT_CONFIG_HASH not in tags:
                continue

            total_charts += 1
            rc = report_configs.get(tags[TAG_REPORT_CONFIG_HASH], {})
            pc = pivot_configs.get(tags.get(TAG_PIVOT_CONFIG_HASH, ""), {})
            try:
                mc = json.loads(tags.get(TAG_MAPPING_CONFIG, "{}"))
            except Exception:
                mc = {}

            # Fetch from Synapse
            # Use StaticTimePeriodIds when available — these are the exact
            # periods the chart was rendered with. DynamicTimePeriod would
            # give us the latest N which may not match the chart's current data.
            stp = rc.get("StaticTimePeriodIds", [])
            payload = {
                "project_id": rc.get("ProjectId"),
                "reporting_plan_id": rc.get("ReportingPlanId"),
                "analysis_ids": rc.get("AnalysisIds"),
                "segment_ids": rc.get("SegmentIds"),
            }
            if stp:
                payload["setup_type"] = "STATIC"
                payload["static_time_period_ids"] = stp
            else:
                dtp = rc.get("DynamicTimePeriod")
                payload["setup_type"] = "DYNAMIC"
                if dtp:
                    payload["dynamic_time_period"] = {
                        "latest_n_deliverables": dtp.get("LatestNDeliverables", 8),
                        "include_live_wave": dtp.get("IncludeLiveWave", True),
                    }

            try:
                resp = requests.post(f"{BASE}/api/reports/generate", headers=headers, json=payload, timeout=30)
                data = resp.json()
                records = data.get("records", [])
            except Exception as e:
                print(f"  Slide {slide_idx} {shape.name}: API error {e}")
                no_data += 1
                continue

            if not records:
                no_data += 1
                continue

            # Pivot using configs
            result = pivot_synapse_records(records, pc, mc)
            if result is None:
                no_data += 1
                continue

            syn_cats, syn_series = result

            # Compare with actual chart
            try:
                plot = shape.chart.plots[0]
                actual_cats = list(plot.categories) if plot.categories else []
                actual_series = [(ser.name, list(ser.values)) for ser in plot.series]
            except Exception:
                no_data += 1
                continue

            # Check category count match
            cat_match = len(syn_cats) == len(actual_cats)
            ser_match = len(syn_series) == len(actual_series)

            # Check value match (first series, first 3 values)
            val_match = False
            if syn_series and actual_series:
                sv = syn_series[0][1][:3]
                av = actual_series[0][1][:3]
                val_match = len(sv) == len(av) and all(
                    abs(a - b) < 0.02 for a, b in zip(av, sv)
                )

            if cat_match and ser_match and val_match:
                matched += 1
                print(f"  Slide {slide_idx} {shape.name}: MATCH ({len(syn_cats)} cats, {len(syn_series)} series)")
            else:
                mismatched += 1
                print(f"  Slide {slide_idx} {shape.name}: MISMATCH")
                print(f"    cats: synapse={len(syn_cats)} chart={len(actual_cats)}")
                print(f"    series: synapse={len(syn_series)} chart={len(actual_series)}")
                if syn_series and actual_series:
                    print(f"    vals[0][:3]: synapse={syn_series[0][1][:3]} chart={actual_series[0][1][:3]}")

    print(f"\n{'='*60}")
    print(f"RESULTS: {total_charts} charts tested")
    print(f"  Matched:    {matched}")
    print(f"  Mismatched: {mismatched}")
    print(f"  No data:    {no_data}")
    print(f"{'='*60}")


if __name__ == "__main__":
    main()
