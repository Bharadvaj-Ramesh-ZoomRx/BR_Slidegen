"""Extra tests for find-or-create + scoring + LLM arbitration paths."""
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))

from slidegen.headliner_full_workflow import (
    _ensure_headline_shape,
    _find_empty_headline_zone_shape,
    _find_headline_shape,
    _llm_pick_talking_header,
    _score_narrative,
)


def _make_slide(width_in=10.0):
    prs = Presentation()
    prs.slide_width = Inches(width_in)
    prs.slide_height = Inches(7.5)
    return prs.slides.add_slide(prs.slide_layouts[6])


def _add_textbox(slide, text, *, left_in, top_in, width_in, height_in=0.5,
                 name=None):
    tb = slide.shapes.add_textbox(
        Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    tb.text_frame.text = text
    if name:
        tb.name = name
    return tb


# ── _score_narrative ───────────────────────────────────────────────────


def test_score_narrative_rewards_data_claim():
    narrative = (
        "RINVOQ retained perceived edge on rapid symptom relief at 47 percent "
        "vs 33 percent for Tremfya across UC HCPs in Q1, dropping 6 points QoQ.")
    label = "Message Recall"
    assert _score_narrative(narrative) > _score_narrative(label) + 5


def test_score_penalises_label_patterns():
    label = "Interaction Details, All Products"
    narrative = (
        "Reps held steady on call quality this quarter while losing "
        "ground on access conversations.")
    assert _score_narrative(narrative) > _score_narrative(label)


def test_score_handles_empty():
    assert _score_narrative("") == -10.0


# ── _llm_pick_talking_header ───────────────────────────────────────────


def test_llm_pick_returns_index_from_response():
    candidates = ["First label.", "A real narrative claim.", "Yet another label."]
    pick = _llm_pick_talking_header(candidates, lambda p: "2")
    assert pick == 1


def test_llm_pick_handles_chatty_response():
    pick = _llm_pick_talking_header(
        ["A", "B", "C"],
        lambda p: "Looking at these, candidate 3 reads most like a claim."
    )
    assert pick == 2


def test_llm_pick_returns_none_on_unparseable():
    assert _llm_pick_talking_header(
        ["A", "B"], lambda p: "no integer here") is None
    assert _llm_pick_talking_header(
        ["A", "B"], lambda p: "5") is None


def test_llm_pick_handles_arbiter_exception():
    def boom(prompt):
        raise RuntimeError("API down")
    assert _llm_pick_talking_header(["A", "B"], boom) is None


# ── _find_headline_shape with scoring + LLM arbitration ────────────────


def test_find_headline_picks_higher_scoring_when_decisive():
    slide = _make_slide()
    _add_textbox(slide, "Interaction Details, All Products",
                 left_in=0.3, top_in=0.2, width_in=8.0)
    _add_textbox(slide,
                 "RINVOQ retained perceived edge on rapid symptom relief at "
                 "47 percent vs 33 percent for Tremfya, climbing 6 points QoQ.",
                 left_in=0.3, top_in=1.5, width_in=8.0)
    pick = _find_headline_shape(slide)
    assert pick is not None
    assert pick.text_frame.text.startswith("RINVOQ retained")


def test_find_headline_calls_llm_on_close_scores():
    slide = _make_slide()
    _add_textbox(slide,
                 "Reps held steady on call quality this period.",
                 left_in=0.3, top_in=0.2, width_in=8.0)
    _add_textbox(slide,
                 "Reps gained traction on access conversations this period.",
                 left_in=0.3, top_in=1.5, width_in=8.0)

    arbiter_calls = []
    def arbiter(prompt):
        arbiter_calls.append(prompt)
        return "1"
    pick = _find_headline_shape(slide, llm_arbiter=arbiter)
    assert len(arbiter_calls) == 1
    assert pick is not None


# ── _ensure_headline_shape — find-or-create ────────────────────────────


def test_ensure_returns_existing_when_present():
    slide = _make_slide()
    _add_textbox(slide,
                 "RINVOQ retained perceived edge on rapid symptom relief at "
                 "47 percent vs 33 percent for Tremfya across UC HCPs.",
                 left_in=0.3, top_in=0.2, width_in=8.0)
    shape, source = _ensure_headline_shape(slide)
    assert source == "existing"
    assert "RINVOQ" in shape.text_frame.text


def test_ensure_finds_empty_shape_in_zone():
    slide = _make_slide()
    empty_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.2),
                                         Inches(8.0), Inches(1.0))
    empty_box.name = "EmptyHeaderShape"
    shape, source = _ensure_headline_shape(slide)
    assert source == "empty"
    assert shape.name == "EmptyHeaderShape"
    assert shape.text_frame.text == ""


def test_ensure_creates_when_no_zone_shape_exists():
    slide = _make_slide()
    footer = slide.shapes.add_textbox(Inches(0.3), Inches(4.0),
                                      Inches(8.0), Inches(0.5))
    footer.text_frame.text = "Footer"
    n_before = len(list(slide.shapes))
    shape, source = _ensure_headline_shape(slide)
    assert source == "created"
    n_after = len(list(slide.shapes))
    assert n_after == n_before + 1
    assert shape.top < Inches(2)
    assert shape.left < Inches(1)


def test_ensure_existing_takes_priority_over_empty_zone_shape():
    slide = _make_slide()
    slide.shapes.add_textbox(Inches(0.3), Inches(0.2),
                             Inches(4.0), Inches(0.5))
    _add_textbox(slide,
                 "RINVOQ retained perceived edge on rapid symptom relief at "
                 "47 percent vs 33 percent across UC HCPs.",
                 left_in=0.3, top_in=1.0, width_in=8.0)
    shape, source = _ensure_headline_shape(slide)
    assert source == "existing"
    assert "RINVOQ" in shape.text_frame.text


def test_find_empty_headline_zone_shape_skips_text_bearing():
    slide = _make_slide()
    _add_textbox(slide, "RINVOQ retained edge.",
                 left_in=0.3, top_in=0.2, width_in=8.0)
    pick = _find_empty_headline_zone_shape(slide)
    assert pick is None


def test_refresh_headlines_skips_existing_label_style_header(tmp_path, monkeypatch):
    """When the existing talking header reads as a slide-section label
    ("Message Recall CREON"), the deck author intentionally didn't put
    a narrative claim there — leave it alone. status = skipped_label."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from slidegen.headliner_full_workflow import refresh_headlines

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Existing header is a section-style label (long enough to pass the
    # 30-char layout filter, but devoid of narrative signal — no verbs,
    # no numbers, no comparisons; matches the "Interaction Details" pattern).
    _add_textbox(slide, "Interaction Details, All Products and Segments",
                 left_in=0.3, top_in=0.2, width_in=8.0)
    cd = CategoryChartData()
    cd.categories = ["A", "B"]
    cd.add_series("S1", [1.0, 2.0])
    slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(1), Inches(2), Inches(8), Inches(4), cd)

    src_path = tmp_path / "src.pptx"; prs.save(src_path)
    ref_path = tmp_path / "ref.pptx"; prs.save(ref_path)
    out_path = tmp_path / "out.pptx"

    llm_called = []
    def fake_claude(prompt, model=None):
        llm_called.append(prompt)
        return "should not be called"
    monkeypatch.setattr(
        "slidegen.headliner_full_workflow._call_claude", fake_claude)
    monkeypatch.setenv("LLM_API_KEY", "stub")

    updates = refresh_headlines(
        source_pptx=str(src_path),
        refreshed_pptx=str(ref_path),
        spec_path=None,
        out_pptx=str(out_path),
    )
    assert updates[0].status == "skipped_label"
    # The LLM rewrite call must NOT have been made for the label header
    # (only ones we'd expect are arbiter calls — not relevant here since
    # there's only one candidate that passes the filter, no arbitration).
    assert not any("PowerPoint" in p for p in llm_called)


def test_refresh_headlines_writes_when_existing_header_is_narrative(tmp_path, monkeypatch):
    """Counter-test: an existing narrative-style header still gets
    rewritten (skipped_label only fires for label-style text)."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from slidegen.headliner_full_workflow import refresh_headlines

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(slide,
                 "RINVOQ retained perceived edge on rapid symptom relief at "
                 "47 percent vs 33 percent for Tremfya across UC HCPs.",
                 left_in=0.3, top_in=0.2, width_in=8.0)
    cd = CategoryChartData()
    cd.categories = ["A", "B"]
    cd.add_series("S1", [1.0, 2.0])
    slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(1), Inches(2), Inches(8), Inches(4), cd)

    src_path = tmp_path / "src.pptx"; prs.save(src_path)
    ref_path = tmp_path / "ref.pptx"; prs.save(ref_path)
    out_path = tmp_path / "out.pptx"

    monkeypatch.setattr(
        "slidegen.headliner_full_workflow._call_claude",
        lambda prompt, model=None: "Rewritten claim from Claude."
    )
    monkeypatch.setenv("LLM_API_KEY", "stub")

    updates = refresh_headlines(
        source_pptx=str(src_path),
        refreshed_pptx=str(ref_path),
        spec_path=None,
        out_pptx=str(out_path),
    )
    assert updates[0].status == "updated"
