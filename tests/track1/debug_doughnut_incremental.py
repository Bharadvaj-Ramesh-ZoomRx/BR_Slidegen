"""
Incremental-construction debug for the doughnut Repair bug.

Line chart was fixed (dLblPos="inEnd" invalid for line; clamped to "t").
Doughnut still broken — different root cause. This script builds a doughnut
chart step-by-step to isolate.

Outputs: experiments/deck_analysis/outputs/generated/debug_doughnut_v*.pptx
"""
from __future__ import annotations
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

OUT_DIR = REPO / "experiments" / "deck_analysis" / "outputs" / "generated"
OUT_DIR.mkdir(parents=True, exist_ok=True)


def _new_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.500)
    return prs


def build_v01_baseline():
    """Raw python-pptx doughnut — no modifications at all."""
    prs = _new_prs()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Rybrevant+Lazcluze", "Tagrisso", "Other EGFRi", "Chemo+IO"]
    data.add_series("Allocation", [0.30, 0.45, 0.15, 0.10])
    slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT,
        Inches(2), Inches(1), Inches(9), Inches(5),
        data,
    )
    path = OUT_DIR / "debug_doughnut_v01_baseline.pptx"
    prs.save(str(path))
    return path


def build_v02_slice_colors():
    """Baseline + set_pie_slice_colors (per-slice dPt elements)."""
    prs = _new_prs()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Rybrevant+Lazcluze", "Tagrisso", "Other EGFRi", "Chemo+IO"]
    data.add_series("Allocation", [0.30, 0.45, 0.15, 0.10])
    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT,
        Inches(2), Inches(1), Inches(9), Inches(5),
        data,
    )
    chart = cf.chart

    from slidegen.pptx_utils.lxml_helpers import set_pie_slice_colors, set_donut_hole_size
    colors = [
        RGBColor(0xF7, 0x58, 0x24),  # Rybrevant orange
        RGBColor(0x70, 0x30, 0xA0),  # Tagrisso purple
        RGBColor(0xBF, 0xBF, 0xBF),  # Other grey
        RGBColor(0x4F, 0x81, 0xBD),  # Chemo+IO blue
    ]
    set_pie_slice_colors(chart, colors)
    set_donut_hole_size(chart, 50)

    path = OUT_DIR / "debug_doughnut_v02_slices.pptx"
    prs.save(str(path))
    return path


def build_v03_chrome():
    """v02 + all chrome: legend hide, title off, data labels, element reorder."""
    prs = _new_prs()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Rybrevant+Lazcluze", "Tagrisso", "Other EGFRi", "Chemo+IO"]
    data.add_series("Allocation", [0.30, 0.45, 0.15, 0.10])
    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT,
        Inches(2), Inches(1), Inches(9), Inches(5),
        data,
    )
    chart = cf.chart

    from slidegen.pptx_utils.lxml_helpers import set_pie_slice_colors, set_donut_hole_size
    from slidegen.pptx_utils.lxml_helpers import set_datalabel_format, set_datalabel_pos_inside_end
    colors = [
        RGBColor(0xF7, 0x58, 0x24),
        RGBColor(0x70, 0x30, 0xA0),
        RGBColor(0xBF, 0xBF, 0xBF),
        RGBColor(0x4F, 0x81, 0xBD),
    ]
    set_pie_slice_colors(chart, colors)
    set_donut_hole_size(chart, 50)

    # Chrome — same as slide_creator's _apply_chart_chrome + _apply_data_labels
    chart.has_title = False

    # Legend: keep but hide (same as line fix)
    from slidegen.slide_creator import _hide_legend_but_keep_element
    chart.has_legend = True
    _hide_legend_but_keep_element(chart)

    # Data labels
    for series in chart.series:
        try:
            series.data_labels.show_value = True
            set_datalabel_format(series, "0%")
            set_datalabel_pos_inside_end(series)
        except Exception:
            pass

    # Element reorder
    from slidegen.slide_creator import _enforce_ser_child_order
    _enforce_ser_child_order(chart)

    path = OUT_DIR / "debug_doughnut_v03_chrome.pptx"
    prs.save(str(path))
    return path


def build_v04_full():
    """v03 + chrome text boxes (headline/footer/section)."""
    prs = _new_prs()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Rybrevant+Lazcluze", "Tagrisso", "Other EGFRi", "Chemo+IO"]
    data.add_series("Allocation", [0.30, 0.45, 0.15, 0.10])
    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT,
        Inches(2), Inches(1), Inches(9), Inches(5),
        data,
    )
    chart = cf.chart

    from slidegen.pptx_utils.lxml_helpers import set_pie_slice_colors, set_donut_hole_size
    from slidegen.pptx_utils.lxml_helpers import set_datalabel_format, set_datalabel_pos_inside_end
    colors = [
        RGBColor(0xF7, 0x58, 0x24),
        RGBColor(0x70, 0x30, 0xA0),
        RGBColor(0xBF, 0xBF, 0xBF),
        RGBColor(0x4F, 0x81, 0xBD),
    ]
    set_pie_slice_colors(chart, colors)
    set_donut_hole_size(chart, 50)

    chart.has_title = False
    from slidegen.slide_creator import _hide_legend_but_keep_element
    chart.has_legend = True
    _hide_legend_but_keep_element(chart)

    for series in chart.series:
        try:
            series.data_labels.show_value = True
            set_datalabel_format(series, "0%")
            set_datalabel_pos_inside_end(series)
        except Exception:
            pass

    from slidegen.slide_creator import _enforce_ser_child_order
    _enforce_ser_child_order(chart)

    # Chrome text boxes
    from slidegen.pptx_utils.layout import slide_header, slide_footer, section_header_bar
    section_header_bar(slide, "Patient Allocation", top=1.40)
    slide_header(slide, "Rybrevant+Lazcluze gains 5pp share QoQ, closing the gap with Tagrisso")
    slide_footer(slide, "Source: Rybrevant PET Q1 2026, n=312 NSCLC prescribers.")

    path = OUT_DIR / "debug_doughnut_v04_full.pptx"
    prs.save(str(path))
    return path


def main() -> int:
    print("Generating 4 incremental doughnut debug files:")
    print()
    for builder, desc in [
        (build_v01_baseline, "v01 -- raw baseline (no modifications)"),
        (build_v02_slice_colors, "v02 -- + set_pie_slice_colors + hole_size"),
        (build_v03_chrome, "v03 -- + legend hide + title off + data labels + reorder"),
        (build_v04_full, "v04 -- + headline/footer/section text boxes"),
    ]:
        path = builder()
        print(f"  [{desc:55s}] {path.name}  ({path.stat().st_size} bytes)")
    print()
    print("Open each in PowerPoint and report first one that triggers Repair.")
    print("  v01 broken -> raw python-pptx doughnut is the issue")
    print("  v01 OK, v02 broken -> set_pie_slice_colors / dPt elements are the issue")
    print("  v02 OK, v03 broken -> chrome operations (legend hide or data labels) are the issue")
    print("  v03 OK, v04 broken -> chrome text boxes are the issue")
    print("  All clean -> render_spec_to_file does something debug doesn't replicate")
    return 0


if __name__ == "__main__":
    sys.exit(main())
