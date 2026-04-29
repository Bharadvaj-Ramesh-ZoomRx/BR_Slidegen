"""Step 2 — No silent chart corruption after refresh.

Asserts the property that the existing refresh-execution eval misses:

    For every chart where the source deck had at least one non-None series
    value, the refreshed deck must also have at least one non-None series
    value for that same chart.

What this catches: the mapper's alignment-failure fallback at
synapse_chart_mapper.py:913 writes `[None] * len(categories)` when it
can't align the API series names to the source chart's series. The chart
shape stays in place but renders as empty in PowerPoint. Today these
slides are reported as `refresh_status: "ok"` because the API call
succeeded; this eval makes the corruption auditable.

Surfaced 2026-04-29 during Step 7 work. ~40 charts across 18 slides
(14 ATU + 4 CREON) currently fail this assertion.

Expected to FAIL until the mapper is fixed to either preserve source
values on alignment failure or emit `alignment_failed` status. Tracked
in memory as `project_step2_chart_corruption.md`.

Run:
    pytest tests/evals/end_to_end/test_step2_no_chart_corruption.py -v
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation

REPO_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO_ROOT))

from tests.evals.fixtures import FIXTURE_DECKS, REFRESHED_DECKS  # noqa: E402


def _chart_has_values(chart) -> bool:
    for plot in chart.plots:
        for s in plot.series:
            if any(v is not None for v in s.values):
                return True
    return False


def _chart_all_none(chart) -> bool:
    for plot in chart.plots:
        for s in plot.series:
            if any(v is not None for v in s.values):
                return False
    return True


@pytest.mark.parametrize("deck_key", list(REFRESHED_DECKS.keys()))
def test_no_all_none_corruption(deck_key: str):
    """Refreshed chart must have ≥1 non-None value where source had values."""
    src_path = FIXTURE_DECKS.get(deck_key)
    ref_path = REFRESHED_DECKS.get(deck_key)

    if src_path is None or not src_path.exists():
        pytest.skip(f"Source deck not on disk: {src_path}")
    if ref_path is None or not ref_path.exists():
        pytest.skip(f"Refreshed deck not on disk: {ref_path}")

    src = Presentation(str(src_path))
    ref = Presentation(str(ref_path))

    corrupted = []
    for s_idx, (s_slide, r_slide) in enumerate(zip(src.slides, ref.slides)):
        src_charts = [sh for sh in s_slide.shapes if sh.has_chart]
        ref_charts = [sh for sh in r_slide.shapes if sh.has_chart]
        if not src_charts or not ref_charts:
            continue
        # Pair charts by index — best-effort, since names may differ
        for sc, rc in zip(src_charts, ref_charts):
            if _chart_has_values(sc.chart) and _chart_all_none(rc.chart):
                corrupted.append((s_idx + 1, sc.name))

    assert not corrupted, (
        f"[{deck_key}] {len(corrupted)} charts have all-None refreshed values "
        f"where source had values:\n"
        + "\n".join(f"  slide {s}  chart {n!r}" for s, n in corrupted[:30])
        + (f"\n  ... +{len(corrupted) - 30} more" if len(corrupted) > 30 else "")
    )
