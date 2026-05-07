"""Step 7 — Headline rewrite contract eval.

Codifies the contract between the refresh layer (Step 2) and the headline
rewrite layer (Step 7). Verifies that for every connected slide, the headline
outcome matches what the refresh outcome + headline-finder logic dictate —
without making any LLM API calls (dry_run mode).

The contract (per design + user 2026-04-29):

    refresh succeeded + values changed + slide had narrative   -> updated
    refresh succeeded + values unchanged                       -> unchanged
    refresh succeeded + values changed + no narrative          -> no_headline
    refresh corrupted (all-None refreshed values)              -> unchanged (preserved)
    refresh execution succeeded but no chart on slide          -> no_chart

Two separate assertions:

  1. Status consistency — `refresh_headlines(dry_run=True)` produces the
     same status that an independent re-derivation from the underlying
     primitives (_values_eq, _all_none, _find_headline_shape) would produce.
     Catches drift between the orchestration and the rules.

  2. Deck writes match status — after a real run (when *_with_headlines.pptx
     exists), the title shape's text in the with_headlines deck differs from
     the refreshed deck if-and-only-if the slide's status is "updated".

Run:
    pytest tests/evals/end_to_end/test_step7_headlines.py -v
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO_ROOT))

from pptx import Presentation  # noqa: E402

from slidegen.headliner_full_workflow import (  # noqa: E402
    _all_none,
    _find_headline_shape,
    _largest_chart,
    _values_eq,
    refresh_headlines,
)
from tests.evals.end_to_end.compare_decks import _extract_chart_data  # noqa: E402
from tests.evals.fixtures import FIXTURE_DECKS, REFRESHED_DECKS, REPO_ROOT as FIX_REPO  # noqa: E402

OUTPUT_DIR = FIX_REPO / "output" / "step2_test_connected"


def _spec_path(deck_key: str) -> Path:
    return OUTPUT_DIR / f"{deck_key}_full_spec.json"


def _with_headlines_path(deck_key: str) -> Path:
    return OUTPUT_DIR / f"{deck_key}_with_headlines.pptx"


def _expected_status(s_slide, r_slide) -> str:
    """Re-derive expected headline status from the primitives.

    This is the contract spelled out as a pure function. If
    refresh_headlines() produces a status that disagrees with this
    function, the orchestration has drifted from the rules.
    """
    src_chart = _largest_chart(s_slide)
    ref_chart = _largest_chart(r_slide)
    if src_chart is None or ref_chart is None:
        return "no_chart"
    try:
        _, src_series = _extract_chart_data(src_chart)
        _, ref_series = _extract_chart_data(ref_chart)
    except Exception:
        return "no_chart"
    if _values_eq(src_series, ref_series):
        return "unchanged"
    if _all_none(ref_series) and not _all_none(src_series):
        return "unchanged"
    if _find_headline_shape(r_slide) is None:
        return "no_headline"
    return "updated"


@pytest.mark.parametrize("deck_key", list(REFRESHED_DECKS.keys()))
def test_status_consistency(deck_key: str):
    """Dry-run statuses must match independent rule re-derivation per slide."""
    src_path = FIXTURE_DECKS.get(deck_key)
    ref_path = REFRESHED_DECKS.get(deck_key)
    spec_path = _spec_path(deck_key)

    if src_path is None or not src_path.exists():
        pytest.skip(f"Source deck not on disk: {src_path}")
    if ref_path is None or not ref_path.exists():
        pytest.skip(f"Refreshed deck not on disk: {ref_path}")
    if not spec_path.exists():
        pytest.skip(f"Spec not on disk: {spec_path}")

    updates = refresh_headlines(
        source_pptx=str(src_path),
        refreshed_pptx=str(ref_path),
        spec_path=str(spec_path),
        out_pptx=str(OUTPUT_DIR / f"{deck_key}_eval_unused.pptx"),
        dry_run=True,
    )
    by_idx = {u.slide_idx: u for u in updates}

    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    spec_slides = {s["slide_index"] for s in spec.get("slides", [])}

    src = Presentation(str(src_path))
    ref = Presentation(str(ref_path))

    failures = []
    for s_idx, (s_slide, r_slide) in enumerate(zip(src.slides, ref.slides)):
        if s_idx not in spec_slides:
            continue
        actual = by_idx.get(s_idx)
        if actual is None:
            failures.append((s_idx + 1, "no status emitted"))
            continue
        expected = _expected_status(s_slide, r_slide)
        if actual.status != expected:
            failures.append(
                (s_idx + 1, f"expected={expected!r} got={actual.status!r}")
            )

    assert not failures, (
        f"Step 7 contract violations on {deck_key} ({len(failures)} slides):\n"
        + "\n".join(f"  slide {sn}: {msg}" for sn, msg in failures)
    )


@pytest.mark.parametrize("deck_key", list(REFRESHED_DECKS.keys()))
def test_deck_writes_match_status(deck_key: str):
    """After a real run, with_headlines deck must reflect the contract.

    For each slide with a real run output:
      - status='updated'   => title text differs from refreshed deck
      - status in {'unchanged', 'no_headline', 'no_chart'} => title text matches

    Skips gracefully if no real-run output exists yet.
    """
    ref_path = REFRESHED_DECKS.get(deck_key)
    with_headlines = _with_headlines_path(deck_key)
    status_json = OUTPUT_DIR / f"{deck_key}_with_headlines_status.json"
    spec_path = _spec_path(deck_key)

    if not (ref_path and ref_path.exists()):
        pytest.skip(f"Refreshed deck not on disk: {ref_path}")
    if not with_headlines.exists():
        pytest.skip(
            f"with_headlines deck not produced yet: {with_headlines}\n"
            "Generate by running: python -m slidegen.headliner_full_workflow "
            f"{deck_key}"
        )
    if not status_json.exists():
        pytest.skip(f"Status sidecar not on disk: {status_json}")
    if not spec_path.exists():
        pytest.skip(f"Spec not on disk: {spec_path}")

    status_data = json.loads(status_json.read_text(encoding="utf-8"))
    by_idx = {u["slide_idx"]: u for u in status_data.get("updates", [])}
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    spec_slides = {s["slide_index"] for s in spec.get("slides", [])}

    ref = Presentation(str(ref_path))
    hdl = Presentation(str(with_headlines))

    failures = []
    for s_idx, (r_slide, h_slide) in enumerate(zip(ref.slides, hdl.slides)):
        if s_idx not in spec_slides:
            continue
        update = by_idx.get(s_idx)
        if update is None:
            continue
        status = update.get("status")
        if status not in {"updated", "unchanged", "no_headline", "no_chart"}:
            continue

        r_shape = _find_headline_shape(r_slide)
        h_shape = _find_headline_shape(h_slide)
        if r_shape is None or h_shape is None:
            continue
        r_text = (r_shape.text_frame.text or "").strip()
        h_text = (h_shape.text_frame.text or "").strip()

        if status == "updated" and r_text == h_text:
            failures.append(
                (s_idx + 1, "marked 'updated' but title text unchanged")
            )
        elif status in {"unchanged", "no_headline", "no_chart"} and r_text != h_text:
            failures.append(
                (s_idx + 1,
                 f"marked {status!r} but title text changed:\n"
                 f"      ref:  {r_text[:80]!r}\n"
                 f"      hdl:  {h_text[:80]!r}")
            )

    assert not failures, (
        f"Step 7 deck writes don't match status on {deck_key} "
        f"({len(failures)} slides):\n"
        + "\n".join(f"  slide {sn}: {msg}" for sn, msg in failures)
    )
