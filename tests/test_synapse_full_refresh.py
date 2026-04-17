"""
Full end-to-end Synapse refresh:
1. Read Connector tags from source deck
2. Fetch fresh data from Synapse API for each chart
3. Pivot + map using PivotConfig + MappingConfig
4. Clone the DUMMY deck, write Synapse data into charts
5. Compare output chart values against original deck values
"""
import json
import shutil
import sys
import os
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

from dotenv import load_dotenv
load_dotenv(REPO / ".env")

from pptx import Presentation
from pptx.chart.data import CategoryChartData

from slidegen.deck_reader.tag_reader import (
    _get_shape_tags_all, _parse_custom_xml_parts,
    TAG_REPORT_CONFIG_HASH, TAG_PIVOT_CONFIG_HASH, TAG_MAPPING_CONFIG,
    XML_STORE_REPORT_CONFIG, XML_STORE_PIVOT_CONFIG,
)
from slidegen.synapse_chart_mapper import (
    refresh_chart_from_synapse, ChartRefreshData,
)
from slidegen.slide_refresher import verify_round_trip

BASE = "https://synapse-api.zoomrx.com"
TOKEN = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("SYNAPSE_API_TOKEN", "")

SRC = "C:/Users/VijayGanesan/Desktop/Galen/Sample Decks/Repatha VESALIUS Weekly Pulse Study - Week 21 Final Report.pptx"
DUMMY = "C:/Users/VijayGanesan/Desktop/Galen/Sample Decks/Repatha_DUMMY.pptx"
OUTPUT = "C:/Users/VijayGanesan/Desktop/Galen/Sample Decks/Repatha_SYNAPSE_REAL.pptx"


def main():
    print("=" * 60)
    print("FULL SYNAPSE REFRESH — real data from API")
    print("=" * 60)

    # Parse Connector tags
    prs_src = Presentation(SRC)
    xml_configs = _parse_custom_xml_parts(prs_src, pptx_path=Path(SRC))
    report_configs = xml_configs.get(XML_STORE_REPORT_CONFIG, {})
    pivot_configs = xml_configs.get(XML_STORE_PIVOT_CONFIG, {})

    # Clone DUMMY deck
    shutil.copy2(DUMMY, OUTPUT)
    prs = Presentation(OUTPUT)

    charts_refreshed = 0
    charts_failed = 0
    charts_skipped = 0
    charts_matched = 0

    # Also open original for comparison
    prs_orig = Presentation(SRC)

    for slide_idx in range(len(prs.slides)):
        if slide_idx >= len(prs_src.slides):
            break

        src_slide = prs_src.slides[slide_idx]
        out_slide = prs.slides[slide_idx]
        orig_slide = prs_orig.slides[slide_idx]

        # Build shape lookup by position for the output deck
        out_charts = {(s.left, s.top): s for s in out_slide.shapes if s.has_chart}
        orig_charts = {(s.left, s.top): s for s in orig_slide.shapes if s.has_chart}

        for shape in src_slide.shapes:
            if not shape.has_chart:
                continue

            tags = _get_shape_tags_all(shape)
            if TAG_REPORT_CONFIG_HASH not in tags:
                charts_skipped += 1
                continue

            rc = report_configs.get(tags[TAG_REPORT_CONFIG_HASH], {})
            pc = pivot_configs.get(tags.get(TAG_PIVOT_CONFIG_HASH, ""), {})
            try:
                mc = json.loads(tags.get(TAG_MAPPING_CONFIG, "{}"))
            except Exception:
                mc = {}

            if not rc.get("AnalysisIds"):
                charts_skipped += 1
                continue

            # Fetch + pivot + map from Synapse
            result = refresh_chart_from_synapse(
                shape, rc, pc, mc, BASE, TOKEN
            )

            if not result.success or not result.categories or not result.series:
                charts_failed += 1
                if result.error:
                    pass  # silently skip — some charts have complex configs
                continue

            # Find matching shape in output deck by position
            key = (shape.left, shape.top)
            out_shape = out_charts.get(key)
            if out_shape is None:
                charts_failed += 1
                continue

            # Write Synapse data into the chart
            try:
                cd = CategoryChartData()
                cd.categories = result.categories
                for name, vals in result.series:
                    cd.add_series(name, vals)
                out_shape.chart.replace_data(cd)
                charts_refreshed += 1

                # Compare with original chart values
                orig_shape = orig_charts.get(key)
                if orig_shape:
                    try:
                        orig_plot = orig_shape.chart.plots[0]
                        orig_vals = list(orig_plot.series[0].values)
                        syn_vals = result.series[0][1]
                        # Check if the last N values match (charts may show subset)
                        n = min(len(orig_vals), len(syn_vals))
                        if n > 0:
                            # Compare last n values (Synapse may return more periods)
                            ov = orig_vals[-n:]
                            sv = syn_vals[-n:]
                            if all(abs(a - b) < 0.02 for a, b in zip(ov, sv)):
                                charts_matched += 1
                    except Exception:
                        pass

            except Exception as e:
                charts_failed += 1

    # Restore headlines from original (headline-writer integration TBD)
    headlines_restored = 0
    for slide_idx in range(min(len(prs.slides), len(prs_orig.slides))):
        orig_slide = prs_orig.slides[slide_idx]
        out_slide = prs.slides[slide_idx]

        best = None
        best_score = 0
        for s in orig_slide.shapes:
            if not s.has_text_frame:
                continue
            top = (s.top or 0) / 914400
            text = s.text_frame.text.strip()
            if top < 1.5 and len(text) > 20:
                try:
                    fs = 10
                    for p in s.text_frame.paragraphs:
                        for r in p.runs:
                            if r.font.size:
                                fs = r.font.size.pt
                            break
                        break
                    score = fs * len(text)
                    if score > best_score:
                        best_score = score
                        best = (s.left, s.top, text)
                except Exception:
                    pass

        if best:
            hl_left, hl_top, hl_text = best
            for s in out_slide.shapes:
                if s.has_text_frame and s.left == hl_left and s.top == hl_top:
                    for p in s.text_frame.paragraphs:
                        if p.runs:
                            p.runs[0].text = hl_text
                            for r in p.runs[1:]:
                                r.text = ""
                            headlines_restored += 1
                            break
                    break

    prs.save(OUTPUT)

    # Verify shape structure
    verify = verify_round_trip(SRC, OUTPUT)

    print(f"\nRESULTS:")
    print(f"  Charts with Synapse data: {charts_refreshed}")
    print(f"  Charts matched original:  {charts_matched}/{charts_refreshed}")
    print(f"  Charts failed:            {charts_failed}")
    print(f"  Charts skipped (no tags): {charts_skipped}")
    print(f"  Headlines restored:       {headlines_restored}")
    print(f"  Shape verification:       {verify['total_issues']} issues")
    print(f"\n  Output: {OUTPUT}")
    print("=" * 60)


if __name__ == "__main__":
    main()
