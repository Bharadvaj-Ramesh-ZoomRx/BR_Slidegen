"""Refresh workflow tracker.

Regenerates refresh_workflow_tracker.xlsx. Tracks the end-to-end wave refresh
work and maps who (Vijay / Bharadvaj / pending) has done what across the
End-of-April PRD milestones.

Edit the DATA blocks below and re-run:
    python create_tracker.py
"""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── Colors ────────────────────────────────────────────────────────────────
C_HEADER = 'FF1F3864'
C_TIER1  = 'FF2E75B6'
C_TIER2  = 'FF70AD47'
C_WARN   = 'FFFFC000'
C_PASS   = 'FF92D050'
C_FAIL   = 'FFFF0000'
C_INFO   = 'FFBDD7EE'
C_ALT    = 'FFF2F2F2'
C_VIJAY  = 'FFDEEBF7'  # light blue  — Vijay
C_BHARAD = 'FFE2EFDA'  # light green — Bharadvaj
C_PEND   = 'FFFCE4D6'  # light orange — pending

thin = Side(border_style='thin', color='FFBFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def hcell(ws, row, col, text, bg=C_HEADER, fc='FFFFFFFF', align='center'):
    c = ws.cell(row=row, column=col, value=text)
    c.fill = PatternFill('solid', fgColor=bg)
    c.font = Font(color=fc, bold=True, size=11)
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    c.border = BORDER
    return c

def dcell(ws, row, col, text, bg=None, bold=False, align='left', fc='FF000000'):
    c = ws.cell(row=row, column=col, value=text)
    if bg:
        c.fill = PatternFill('solid', fgColor=bg)
    c.font = Font(bold=bold, color=fc, size=10)
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    c.border = BORDER
    return c

def col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def owner_color(owner):
    o = str(owner).upper()
    if 'VIJAY' in o:     return C_VIJAY
    if 'BHARADVAJ' in o: return C_BHARAD
    if 'PENDING' in o:   return C_PEND
    return None

def status_fill(cell, status):
    s = str(status).upper()
    if s.startswith('PASS') or s == 'DONE':      cell.fill = PatternFill('solid', fgColor=C_PASS)
    elif s.startswith('FAIL'):                   cell.fill = PatternFill('solid', fgColor=C_FAIL)
    elif s.startswith('INFO'):                   cell.fill = PatternFill('solid', fgColor=C_INFO)
    elif s in ('PENDING', 'BLOCKED', 'PARTIAL'): cell.fill = PatternFill('solid', fgColor=C_WARN)

TODAY = '2026-04-21'

# ──────────────────────────────────────────────────────────────────────────
# SHEET 1 — Dashboard: high-level component status
# ──────────────────────────────────────────────────────────────────────────
ws1 = wb.active
ws1.title = 'Dashboard'
ws1.sheet_view.showGridLines = False

hcell(ws1, 1, 1, 'REFRESH WORKFLOW — PROGRESS TRACKER', align='center')
ws1.merge_cells('A1:E1')
ws1.row_dimensions[1].height = 30

for col, h in enumerate(['Component', 'Status', 'Owner', 'Last Tested', 'Notes'], 1):
    hcell(ws1, 3, col, h)

dashboard = [
    ('deck-reader — Tier 1 tag extraction',         'DONE',    'Vijay + Bharadvaj', TODAY, 'Extraction works; 55/55 ATU tagged slides complete lineage. Vijay built, Bharadvaj audited.'),
    ('deck-reader — Tier 2 inference',              'PARTIAL', 'Vijay',             '2026-04-20', 'Prototype in data_inference.py (commit 973d843). Tested on Repatha Slide 6. Not production-ready.'),
    ('Synapse token / API auth (MSAL)',             'DONE',    'Vijay + Bharadvaj', TODAY,     'Vijay built synapse_auth.py (commit 2b88c92). Bharadvaj installed sk_ API key in .env; verified HTTP 200 on /api/projects.'),
    ('synapse-read (fetch via Synapse API)',        'DONE',    'Vijay',             '2026-04-17', 'Vijay built fetch_synapse_records. Tested in UAT pipeline.'),
    ('synapse_chart_mapper (records → chart data)', 'DONE',    'Vijay',             '2026-04-20', 'Full pivot-faithful rewrite. 446 lines. Commit 973d843.'),
    ('slide-updater (in-place chart refresh)',      'DONE',    'Vijay',             '2026-04-20', 'replace_data() flow working. 25/27 verified on UAT.'),
    ('headline-writer (regenerate on data change)', 'PENDING', 'Pending',           '',        'Not built yet. May milestone. Vijay currently restores original headlines.'),
    ('deck-assembler (final PPTX output)',          'DONE',    'Vijay',             '2026-04-20', 'In-place save works. No separate assembler needed for refresh path.'),
    ('End-to-end refresh — UAT deck',               'DONE',    'Vijay',             '2026-04-20', '25/27 slides visually correct.'),
    ('End-to-end refresh — real PET client deck',   'PENDING', 'Bharadvaj',         '',        'Item B: run pipeline on ATU Q1\'26 or JJ RYB Q4\'25. Bharadvaj next task.'),
    ('Evals harness — tests/evals/ scaffolded',     'DONE',    'Bharadvaj',         TODAY,     'Folder structure + fixtures.py + first eval (tag_counts) + golden file + mutation test.'),
    ('Evals — Eval #1 tag counts per slide',        'DONE',    'Bharadvaj',         TODAY,     'Positive + mutation test both pass. Short-circuit bug found and fixed.'),
    ('Evals — refresh eval per PRD §6.9',           'PENDING', 'Bharadvaj',         '',        'The PRD-named end-of-Apr deliverable. To be built after Item B run.'),
    ('Visual regression harness',                   'PENDING', 'Pradeep/Vijay',     '',        'PRD target #1 for end-of-Apr. Not Bharadvaj lane.'),
]

for r, (comp, status, owner, tested, notes) in enumerate(dashboard, 4):
    dcell(ws1, r, 1, comp)
    sc = ws1.cell(row=r, column=2, value=status)
    sc.alignment = Alignment(horizontal='center', vertical='center')
    sc.font = Font(bold=True, size=10)
    sc.border = BORDER
    status_fill(sc, status)
    oc = dcell(ws1, r, 3, owner, align='center')
    bg = owner_color(owner)
    if bg: oc.fill = PatternFill('solid', fgColor=bg)
    dcell(ws1, r, 4, tested, align='center')
    dcell(ws1, r, 5, notes)

col_widths(ws1, [48, 12, 22, 14, 75])
ws1.freeze_panes = 'A4'

# ──────────────────────────────────────────────────────────────────────────
# SHEET 2 — End-of-April Milestone: 4 PRD targets × sub-steps × owner
# ──────────────────────────────────────────────────────────────────────────
ws2 = wb.create_sheet('End-of-April Milestone')
ws2.sheet_view.showGridLines = False
ws2.row_dimensions[1].height = 32

hcell(ws2, 1, 1, 'END-OF-APRIL MILESTONE — PRD §9.3 breakdown', align='center')
ws2.merge_cells('A1:F1')

for col, h in enumerate(['Target', 'Sub-step', 'Owner', 'Status', 'Done Date', 'Source / Notes'], 1):
    hcell(ws2, 3, col, h)

milestone = [
    # Target 1
    ('T1: Visual regression on top-6 chart patterns', '905-deck grounding analysis',             'Vijay',      'DONE',    '2026-04-17', 'Commits 0757f03, f9a743a, a44dc9a. 905 decks analyzed.'),
    ('',                                              'pptx_utils regrounding',                  'Vijay',      'DONE',    '2026-04-17', 'Commit f9a743a.'),
    ('',                                              'Visual regression harness in tests/',     'Pradeep',    'PENDING', '',           'Not Bharadvaj lane. Biggest lift of the 4 targets.'),
    ('',                                              'Top-6 pattern × brand goldens',           'Pradeep',    'PENDING', '',           'Depends on harness.'),

    # Target 2
    ('T2: Line/doughnut Repair bug closed',           'Fix PowerPoint Repair bugs',              'Vijay',      'DONE',    '2026-04-16', 'Commit 59e1412.'),

    # Target 3
    ('T3: Connector-tag integration on real PET deck','Fix Tier 1 tag extraction from real decks','Vijay',     'DONE',    '2026-04-16', 'Commit b56f441.'),
    ('',                                              'Build 3-stage refresh test pipeline',     'Vijay',      'DONE',    '2026-04-17', 'Commit 56bb2b4 (test_spec_refresh_pipeline.py).'),
    ('',                                              'SlideSpec v1.2 spec-as-config',           'Vijay',      'DONE',    '2026-04-20', 'Commit 973d843. 25/27 verified on UAT deck.'),
    ('',                                              'Run pipeline on ATU Q1\'26 deck (Item B)','Bharadvaj',  'PENDING', '',           'Next task. ATU deck already deck-read; 197/199 shapes tagged.'),
    ('',                                              'Analyze pass rate + failures',            'Bharadvaj',  'PENDING', '',           'After B run completes.'),

    # Target 4
    ('T4: Evals harness + ≥1 refresh eval on JJ RYB', 'tests/evals/ folder scaffolded',          'Bharadvaj',  'DONE',    TODAY,        'fixtures.py, generate_golden.py, __init__ files.'),
    ('',                                              'Eval #1 tag counts (regression)',         'Bharadvaj',  'DONE',    TODAY,        'Positive + mutation test passing. File: tests/evals/deck_reader/test_tag_counts.py.'),
    ('',                                              'Refresh eval — end-to-end on client deck','Bharadvaj',  'PENDING', '',           'Build after Item B succeeds. Ports Vijay pipeline into tests/evals/end_to_end/.'),
    ('',                                              'Mutation-test the refresh eval',          'Bharadvaj',  'PENDING', '',           'Same workflow as Eval #1.'),
]

for r, row in enumerate(milestone, 4):
    bg = None if r % 2 == 0 else C_ALT
    target, substep, owner, status, donedate, notes = row
    tc = dcell(ws2, r, 1, target, bg=bg, bold=bool(target))
    if target:  # target header row
        tc.fill = PatternFill('solid', fgColor='FFE7E6E6')
        tc.font = Font(bold=True, size=10)
    dcell(ws2, r, 2, substep, bg=bg)
    oc = dcell(ws2, r, 3, owner, bg=bg, align='center')
    oc_bg = owner_color(owner)
    if oc_bg: oc.fill = PatternFill('solid', fgColor=oc_bg)
    sc = ws2.cell(row=r, column=4, value=status)
    sc.alignment = Alignment(horizontal='center', vertical='center')
    sc.font = Font(bold=True, size=10)
    sc.border = BORDER
    status_fill(sc, status)
    dcell(ws2, r, 5, donedate, bg=bg, align='center')
    dcell(ws2, r, 6, notes, bg=bg)

col_widths(ws2, [40, 42, 12, 12, 12, 60])
ws2.freeze_panes = 'A4'

# ──────────────────────────────────────────────────────────────────────────
# SHEET 3 — Activity Log: chronological record of work done
# ──────────────────────────────────────────────────────────────────────────
ws3 = wb.create_sheet('Activity Log')
ws3.row_dimensions[1].height = 35

for col, h in enumerate(['Date', 'Step', 'Owner', 'Activity', 'Deck / Target', 'What Was Done',
                         'Result', 'Issues / Notes', 'Source (commit / run)'], 1):
    hcell(ws3, 1, col, h)

activity = [
    # ── Vijay contributions (from git) ──
    ('2026-04-16', 'Vijay',  'Vijay',     'Fix PowerPoint Repair bugs (line + doughnut)',
     'pet-deck skill', 'Debugged and fixed line + doughnut chart Repair dialog bugs',
     'PASS', 'Hits PRD Target #2', 'git commit 59e1412'),
    ('2026-04-16', 'Vijay',  'Vijay',     'Fix Tier 1 tag extraction from real Synapse decks',
     'deck-reader', 'Connector tags now read correctly from real Synapse-connected decks',
     'PASS', 'Foundation for all refresh work', 'git commit b56f441'),
    ('2026-04-17', 'Vijay',  'Vijay',     'Add human-readable name fields to DataLineage',
     'deck-reader', 'project_name, reporting_plan_name etc. extracted from tags',
     'PASS', '', 'git commit ea93fc8'),
    ('2026-04-17', 'Vijay',  'Vijay',     'Build Synapse chart mapper',
     'synapse_chart_mapper', 'PivotConfig + MappingConfig → chart data',
     'PASS', 'First version', 'git commit 2c9067c'),
    ('2026-04-17', 'Vijay',  'Vijay',     'Build 3-stage refresh test pipeline',
     'test_spec_refresh_pipeline.py', 'Source → dummy → Synapse fetch → refresh → verify',
     'PASS', 'Test scaffold only at this commit', 'git commit 56bb2b4'),
    ('2026-04-17', 'Vijay',  'Vijay',     'MSAL + .env token auth with auto-expiry',
     'synapse_auth.py', 'Device code flow, cache, expiry checks',
     'PASS', 'Auth infra complete', 'git commit 2b88c92'),
    ('2026-04-20', 'Vijay',  'Vijay',     'SlideSpec v1.2 — Connector-faithful refresh',
     'UAT deck (27 slides)', 'Full rewrite of synapse_chart_mapper; spec-as-config schema',
     'PASS 25/27', '2/27 slides had issues — not yet root-caused',
     'git commit 973d843'),

    # ── Bharadvaj contributions ──
    ('2026-04-20', 'B1',     'Bharadvaj', 'Run deck-reader CLI on ATU Q1\'26 deck',
     'ATU Q1\'26 (99 slides)', 'Generated 99 SlideSpec JSONs via python -m slidegen.deck_reader',
     'INFO', 'Old CLI path: data_mapping was None on all components', 'local run'),
    ('2026-04-20', 'B2',     'Bharadvaj', 'Tier 1 spec quality audit',
     'ATU Q1\'26', 'Inspected data_lineage completeness across 55 tagged slides',
     'PASS 55/55', 'All have project_id, reporting_plan_id, analysis_ids, survey_id populated',
     'local analysis'),
    ('2026-04-21', 'B3',     'Bharadvaj', 'Re-run deck-reader via generate_config_specs() API',
     'ATU Q1\'26', 'Confirmed raw_pivot_config + raw_mapping_config populated on chart components',
     'PASS', 'CLI path and library path produce different output — flagged for Vijay',
     'local run'),
    ('2026-04-21', 'B4',     'Bharadvaj', 'Scaffold tests/evals/ folder + fixtures',
     'evals infrastructure', 'Created folder tree, fixtures.py, generate_golden.py',
     'PASS', '', 'local — files committed'),
    ('2026-04-21', 'B5',     'Bharadvaj', 'Build Eval #1 — tag count regression',
     'ATU Q1\'26 golden', 'Captured golden tag-count snapshot; wrote pytest diff test',
     'PASS (positive)', 'Mutation test found short-circuit bug — fixed',
     'tests/evals/deck_reader/test_tag_counts.py'),
    ('2026-04-21', 'B6',     'Bharadvaj', 'Validate Eval #1 via mutation test',
     'ATU golden (mutated)', 'Corrupted 3 values; verified all 3 reported clearly; restored',
     'PASS', 'Eval #1 now marked COVERED on evals tracker',
     'local validation'),
    ('2026-04-21', 'B7',     'Bharadvaj', 'Build evals coverage tracker',
     'evals_coverage_tracker.xlsx', '40 evals catalogued across 9 pipeline stages',
     'PASS', '', 'local file'),

    # ── Pending Item B sub-steps ──
    ('2026-04-21', 'B8',     'Bharadvaj', 'Install Synapse token in .env',
     'synapse_auth', 'Wrote SYNAPSE_API_TOKEN=sk_... into .env; verified get_synapse_token() returns it and /api/projects returns HTTP 200',
     'PASS', 'sk_ API key works as Bearer token; Vijay\'s JWT-decode exception handler passes it through unchanged', 'local run'),
    ('(pending)',  'B9',     'Bharadvaj', 'Parameterize test_spec_refresh_pipeline for new decks',
     'test scaffold', 'Add CLI arg to accept deck path',
     'PENDING', '',                                             ''),
    ('(pending)',  'B10',    'Bharadvaj', 'Run Stage 1 on ATU Q1\'26 (dummy deck)',
     'ATU Q1\'26', 'Clone + wipe chart data + dummy headlines',
     'PENDING', 'No Synapse needed for this stage',             ''),
    ('(pending)',  'B11',    'Bharadvaj', 'Run Stage 2 on ATU Q1\'26 (Synapse refresh)',
     'ATU Q1\'26', 'Fetch records + pivot + replace chart data',
     'PENDING', 'First real-client Synapse API call',           ''),
    ('(pending)',  'B12',    'Bharadvaj', 'Run Stage 3 on ATU Q1\'26 (verify)',
     'ATU Q1\'26', 'Shape-by-shape compare against source',
     'PENDING', 'Produces pass rate',                           ''),
    ('(pending)',  'B13',    'Bharadvaj', 'Analyze failing slides + document',
     'ATU Q1\'26', 'Root-cause each failure, log pattern',
     'PENDING', 'Feedback loop into Vijay pipeline',            ''),
    ('(pending)',  'B14',    'Bharadvaj', 'Verdict — is pipeline production-worthy for real decks?',
     'writeup', 'One-paragraph readout for leadership',
     'PENDING', '',                                             ''),
    ('(pending)',  'T4',     'Bharadvaj', 'Port refresh test into tests/evals/end_to_end/',
     'Eval framework', 'Wrap Vijay pipeline as pytest eval + golden',
     'PENDING', 'Hits PRD Target #4 (end-of-Apr)',              ''),
    ('(pending)',  'T4-mut', 'Bharadvaj', 'Mutation-test the refresh eval',
     'Eval framework', 'Corrupt inputs, verify eval catches',
     'PENDING', 'Same workflow as Eval #1',                     ''),
]

for r, row in enumerate(activity, 2):
    bg = None if r % 2 == 0 else C_ALT
    for c, val in enumerate(row, 1):
        cell = dcell(ws3, r, c, val, bg=bg)
        if c == 3:  # Owner
            ob = owner_color(val)
            if ob: cell.fill = PatternFill('solid', fgColor=ob)
        if c == 7:  # Result
            status_fill(cell, val)
            cell.font = Font(bold=True, size=10)

col_widths(ws3, [12, 8, 12, 38, 26, 48, 16, 50, 30])
ws3.freeze_panes = 'D2'

# ──────────────────────────────────────────────────────────────────────────
# SHEET 4 — Legend
# ──────────────────────────────────────────────────────────────────────────
ws4 = wb.create_sheet('Legend')
ws4.sheet_view.showGridLines = False

hcell(ws4, 1, 1, 'HOW TO READ THIS TRACKER', align='center')
ws4.merge_cells('A1:B1')
ws4.row_dimensions[1].height = 28

legend = [
    ('', ''),
    ('Owner colors', ''),
    ('  Vijay',     'Light blue — work done or to be done by Vijay (rendering / refresh engine / pipeline infra)'),
    ('  Bharadvaj', 'Light green — work done or to be done by Bharadvaj (wave-refresh-workflow + evals)'),
    ('  Pending',   'Light orange — step with no owner assigned yet or explicitly deferred'),
    ('  Pradeep',   'Shown without color — rendering fidelity lane, separate from refresh'),
    ('', ''),
    ('Status colors', ''),
    ('  DONE / PASS',  'Green — fully complete and validated'),
    ('  PENDING / BLOCKED / PARTIAL', 'Amber — started but not complete, or blocked on dependency'),
    ('  FAIL',         'Red — attempted and did not work'),
    ('  INFO',         'Light blue — informational (not pass/fail)'),
    ('', ''),
    ('How this tracker updates', ''),
    ('  When Vijay ships a commit', 'Add a row to Activity Log under his name, mark relevant Milestone sub-step DONE, update Dashboard if a component status changes'),
    ('  When Bharadvaj does work',  'Same pattern — new row, update milestone + dashboard'),
    ('  When something is blocked', 'Mark PENDING on the milestone view with the blocker in the Notes column'),
    ('  Pulling from git',          'Run git log upstream/vijay-slidegen --since="2 days ago" to find new commits; log them in Activity Log'),
    ('', ''),
    ('Sheet guide', ''),
    ('  Dashboard',                    'One-line status per pipeline component. Start here.'),
    ('  End-of-April Milestone',       'Maps the 4 PRD §9.3 end-of-Apr targets into sub-steps, each with an owner'),
    ('  Activity Log',                 'Chronological list of actual work done — commits, local runs, evals built'),
    ('  Legend',                       'This page'),
]

for r, (k, v) in enumerate(legend, 2):
    bold = k.strip() in ('Owner colors', 'Status colors', 'How this tracker updates', 'Sheet guide')
    c1 = dcell(ws4, r, 1, k, bold=bold)
    if bold: c1.fill = PatternFill('solid', fgColor='FFE7E6E6')
    dcell(ws4, r, 2, v)

col_widths(ws4, [32, 90])

# ──────────────────────────────────────────────────────────────────────────
wb.save('refresh_workflow_tracker.xlsx')
print('Saved: refresh_workflow_tracker.xlsx')
