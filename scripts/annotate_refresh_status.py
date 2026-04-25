"""Annotate a refreshed deck with a per-slide status badge.

Reads the refreshed pptx + the eval reports (roundtrip + format) +
the connector spec, then drops a small tag in the top-right of each
slide so a viewer can immediately tell:

    GREEN  TAGGED · all components match  (Step 2 + Step 3 both clean)
    AMBER  TAGGED · partial match         (some chart/table values drift)
    RED    TAGGED · no components match   (refresh failed end-to-end)
    GREY   UNTAGGED · not tested          (slide has no Connector tags)

Usage:
    python scripts/annotate_refresh_status.py <deck_key>

Reads:
    output/<deck_key>_refresh/<deck_key>_refreshed.pptx
    output/<deck_key>_refresh/<deck_key>_full_spec.json
    output/<deck_key>_refresh/<deck_key>_roundtrip_*.json   (latest)
    output/<deck_key>_refresh/<deck_key>_format_*.json      (latest)

Writes:
    output/<deck_key>_refresh/<deck_key>_refreshed_annotated.pptx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

# Status palette
GREEN = RGBColor(0x2E, 0x7D, 0x32)   # all-pass
AMBER = RGBColor(0xF9, 0xA8, 0x25)   # partial
RED = RGBColor(0xC6, 0x28, 0x28)     # all-fail
GREY = RGBColor(0x9E, 0x9E, 0x9E)    # untagged
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _latest(work_dir: Path, prefix: str) -> Path | None:
    matches = sorted(work_dir.glob(f"{prefix}*.json"))
    return matches[-1] if matches else None


def _build_slide_status(deck_key: str, work_dir: Path) -> dict[int, dict]:
    """Build slide_idx -> status dict from spec + roundtrip + format reports."""
    spec_path = work_dir / f"{deck_key}_full_spec.json"
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    tagged_idx = {s["slide_index"] for s in spec["slides"]}

    roundtrip_path = _latest(work_dir, f"{deck_key}_roundtrip_")
    format_path = _latest(work_dir, f"{deck_key}_format_")

    rt = json.loads(roundtrip_path.read_text(encoding="utf-8")) if roundtrip_path else {"components": []}
    ft = json.loads(format_path.read_text(encoding="utf-8")) if format_path else {"components": []}

    # Aggregate by slide
    rt_by_slide: dict[int, list] = {}
    for c in rt.get("components", []):
        rt_by_slide.setdefault(c["slide_idx"], []).append(c)
    ft_by_slide: dict[int, list] = {}
    for c in ft.get("components", []):
        ft_by_slide.setdefault(c["slide_idx"], []).append(c)

    status_by_slide: dict[int, dict] = {}
    all_slide_idx = set(tagged_idx) | set(rt_by_slide.keys())
    for sidx in all_slide_idx:
        rt_comps = rt_by_slide.get(sidx, [])
        ft_comps = ft_by_slide.get(sidx, [])

        # Step 2 — values_match across all components (tables count too)
        s2_total = len(rt_comps)
        s2_pass = sum(1 for c in rt_comps if c.get("values_match") is True)

        # Step 3 — colors + chart_type across charts
        s3_total = len(ft_comps)
        s3_pass = sum(
            1 for c in ft_comps
            if c.get("chart_type_match") is True and c.get("series_colors_match") is True
        )

        is_tagged = sidx in tagged_idx
        status_by_slide[sidx] = {
            "tagged": is_tagged,
            "s2_total": s2_total,
            "s2_pass": s2_pass,
            "s3_total": s3_total,
            "s3_pass": s3_pass,
        }
    return status_by_slide


def _classify(s: dict) -> tuple[RGBColor, str]:
    """Pick a badge color + label text from a slide status dict."""
    if not s["tagged"]:
        return GREY, "UNTAGGED · not tested"

    s2t, s2p = s["s2_total"], s["s2_pass"]
    s3t, s3p = s["s3_total"], s["s3_pass"]

    if s2t == 0:
        # Tagged but nothing comparable (shouldn't really happen)
        return GREY, "TAGGED · no comparables"

    s2_full = s2p == s2t
    s2_zero = s2p == 0
    s3_full = (s3t == 0) or (s3p == s3t)

    if s2_full and s3_full:
        color = GREEN
        label = f"TAGGED · {s2p}/{s2t} match"
    elif s2_zero:
        color = RED
        label = f"TAGGED · 0/{s2t} match"
    else:
        color = AMBER
        label = f"TAGGED · {s2p}/{s2t} match"

    if s3t > 0 and s3p < s3t:
        label += f" · fmt {s3p}/{s3t}"
    return color, label


def _add_badge(slide, color: RGBColor, label: str, slide_w_emu: int) -> None:
    """Drop a small colored rectangle with white text in the top-right corner."""
    width = Inches(2.6)
    height = Inches(0.28)
    left = slide_w_emu - width - Inches(0.15).emu
    top = Inches(0.1)

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.fill.background()
    box.shadow.inherit = False

    tf = box.text_frame
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = WHITE


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)

    deck_key = sys.argv[1]
    work_dir = REPO_ROOT / "output" / f"{deck_key}_refresh"
    in_pptx = work_dir / f"{deck_key}_refreshed.pptx"
    out_pptx = work_dir / f"{deck_key}_refreshed_annotated.pptx"

    if not in_pptx.exists():
        print(f"ERROR: refreshed deck not found: {in_pptx}")
        sys.exit(2)

    status = _build_slide_status(deck_key, work_dir)

    pres = Presentation(str(in_pptx))
    slide_w = pres.slide_width

    counts = {"GREEN": 0, "AMBER": 0, "RED": 0, "GREY": 0}
    for sidx, slide in enumerate(pres.slides):
        s = status.get(sidx, {"tagged": False, "s2_total": 0, "s2_pass": 0,
                              "s3_total": 0, "s3_pass": 0})
        color, label = _classify(s)
        bucket = ("GREEN" if color == GREEN else
                  "AMBER" if color == AMBER else
                  "RED" if color == RED else "GREY")
        counts[bucket] += 1
        _add_badge(slide, color, label, slide_w)

    pres.save(str(out_pptx))
    total = sum(counts.values())
    print(f"=== {deck_key} ===")
    print(f"Annotated {total} slides:")
    for k in ("GREEN", "AMBER", "RED", "GREY"):
        print(f"  {k}: {counts[k]}")
    print(f"\nWrote: {out_pptx}")


if __name__ == "__main__":
    main()
