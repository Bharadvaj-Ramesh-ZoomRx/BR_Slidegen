"""Recompute per-slide headline status without making any API calls.

Runs the same finder + value-comparison logic as headline_refresh.refresh_headlines
in dry-run mode and writes a JSON sidecar that the annotator can consume.

Usage:
    python scripts/compute_headline_status.py atu_q1_26
    python scripts/compute_headline_status.py creon_pet_w33
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation

from slidegen.headline_refresh import (
    HeadlineUpdate,
    _all_none,
    _find_headline_shape,
    _largest_chart,
    _values_eq,
)
from tests.evals.end_to_end.compare_decks import _extract_chart_data

REPO = Path(__file__).resolve().parents[1]

DECKS = {
    "atu_q1_26": {
        "src": REPO / "projects" / "J&J Rybrevant PET" / "Template" / "ZoomRx_UC_ATU_Report_Q1_'26.pptx",
        "ref": REPO / "output" / "step2_test_connected" / "atu_q1_26.pptx",
        "spec": REPO / "output" / "step2_test_connected" / "atu_q1_26_full_spec.json",
        "out": REPO / "output" / "step2_test_connected" / "atu_q1_26_with_headlines_status.json",
    },
    "creon_pet_w33": {
        "src": REPO / "projects" / "J&J Rybrevant PET" / "Template" / "CREON Share of Voice Study - W33.pptx",
        "ref": REPO / "output" / "step2_test_connected" / "creon_pet_w33.pptx",
        "spec": REPO / "output" / "step2_test_connected" / "creon_pet_w33_full_spec.json",
        "out": REPO / "output" / "step2_test_connected" / "creon_pet_w33_with_headlines_status.json",
    },
}


def compute(deck_key: str) -> None:
    paths = DECKS[deck_key]
    spec = json.loads(paths["spec"].read_text(encoding="utf-8"))
    spec_slides = {s["slide_index"] for s in spec.get("slides", [])}

    src = Presentation(str(paths["src"]))
    ref = Presentation(str(paths["ref"]))

    updates = []
    for s_idx, (s_slide, r_slide) in enumerate(zip(src.slides, ref.slides)):
        if s_idx not in spec_slides:
            updates.append({"slide_idx": s_idx, "status": "not_in_spec",
                            "reason": "slide is not connected (not in spec)"})
            continue

        src_chart = _largest_chart(s_slide)
        ref_chart = _largest_chart(r_slide)
        if src_chart is None or ref_chart is None:
            updates.append({"slide_idx": s_idx, "status": "no_chart",
                            "reason": "no chart found"})
            continue
        try:
            _, src_series = _extract_chart_data(src_chart)
            _, ref_series = _extract_chart_data(ref_chart)
        except Exception as exc:
            updates.append({"slide_idx": s_idx, "status": "extract_error",
                            "reason": f"chart extract error: {exc}"})
            continue
        if _values_eq(src_series, ref_series):
            updates.append({"slide_idx": s_idx, "status": "unchanged",
                            "reason": "values unchanged after refresh"})
            continue
        if _all_none(ref_series) and not _all_none(src_series):
            updates.append({"slide_idx": s_idx, "status": "unchanged",
                            "reason": "refreshed values all-None (corrupted) — preserved"})
            continue

        # Values changed — does the slide have a narrative headline?
        # Run finder against the refreshed deck (same as refresh_headlines does)
        headline_shape = _find_headline_shape(r_slide)
        if headline_shape is None:
            updates.append({"slide_idx": s_idx, "status": "no_headline",
                            "reason": "values changed but slide had no narrative headline (per-rule skip)"})
            continue

        old_headline = headline_shape.text_frame.text.strip()
        updates.append({
            "slide_idx": s_idx,
            "status": "updated",
            "reason": "values changed and slide has narrative headline — qualifies for rewrite",
            "old_headline": old_headline,
        })

    paths["out"].write_text(
        json.dumps({"deck_key": deck_key, "updates": updates}, indent=2),
        encoding="utf-8",
    )

    from collections import Counter
    counts = Counter(u["status"] for u in updates)
    print(f"=== {deck_key} ===")
    for k, n in counts.most_common():
        print(f"  {k}: {n}")
    print(f"\nWrote: {paths['out'].relative_to(REPO)}")


if __name__ == "__main__":
    if len(sys.argv) != 2 or sys.argv[1] not in DECKS:
        print("Usage: python scripts/compute_headline_status.py {atu_q1_26|creon_pet_w33}")
        sys.exit(2)
    compute(sys.argv[1])
