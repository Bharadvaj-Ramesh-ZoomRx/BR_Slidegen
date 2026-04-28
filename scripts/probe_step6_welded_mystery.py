"""Probe a few CREON charts the verifier flagged as 'shifted' (new wave labels
visible) but Step 6 classified as 'welded' (values identical to source).

Goal: find out whether the discrepancy is
  (a) a Step 6 pairing bug (_find_match pairs the wrong shapes)
  (b) values genuinely identical to source despite labels changing (e.g.,
      stable metric across waves)
  (c) the verifier overcounting (false-positive on label match)
"""
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation
from tests.evals.end_to_end.compare_decks import _find_match, _pos, _extract_chart_data

REPO = Path(__file__).resolve().parents[1]
SRC = REPO / "projects" / "J&J Rybrevant PET" / "Template" / "CREON Share of Voice Study - W33.pptx"
REF = REPO / "output" / "_step6" / "creon_pet_w33" / "refreshed_shift.pptx"

src = Presentation(str(SRC))
ref = Presentation(str(REF))

# Look at slide 10 first (verifier hit, Step 6 may say welded — main chart on
# this slide is dynamic_latest_n=6 with month-quarter labels).
TARGET_SLIDES = [10, 11, 12, 13, 14]

for slide_idx in TARGET_SLIDES:
    if slide_idx >= len(src.slides) or slide_idx >= len(ref.slides):
        continue
    s = src.slides[slide_idx]
    r = ref.slides[slide_idx]
    src_charts = [sh for sh in s.shapes if sh.has_chart]
    ref_charts = [sh for sh in r.shapes if sh.has_chart]
    print(f"\n=== slide {slide_idx} | {len(src_charts)} src charts / {len(ref_charts)} ref charts ===")
    for i, src_shape in enumerate(src_charts):
        sl, st = _pos(src_shape)
        ref_shape = _find_match(src_shape, ref_charts)
        print(f"\n  [{i}] {src_shape.name!r} pos=({sl},{st})")
        if ref_shape is None:
            print("    NO REF MATCH")
            continue
        rsl, rst = _pos(ref_shape)
        print(f"      ref pos=({rsl},{rst}) name={ref_shape.name!r}")
        try:
            sc, ss = _extract_chart_data(src_shape)
            rc, rs = _extract_chart_data(ref_shape)
        except Exception as e:
            print(f"    extract err: {e}")
            continue
        print(f"    src cats ({len(sc)}): {sc[:6]}{'...' if len(sc)>6 else ''}")
        print(f"    ref cats ({len(rc)}): {rc[:6]}{'...' if len(rc)>6 else ''}")
        print(f"    src series names: {[n for n,_ in ss[:4]]}")
        print(f"    ref series names: {[n for n,_ in rs[:4]]}")
        # Are series names identical?
        src_names = [n for n,_ in ss]
        ref_names = [n for n,_ in rs]
        names_eq = src_names == ref_names
        # Are values identical positionally (first series only)
        if ss and rs and len(ss[0][1]) == len(rs[0][1]):
            sv = [round(v,4) if v is not None else None for v in ss[0][1]]
            rv = [round(v,4) if v is not None else None for v in rs[0][1]]
            vals_eq = sv == rv
            print(f"    series0 src vals: {sv[:6]}{'...' if len(sv)>6 else ''}")
            print(f"    series0 ref vals: {rv[:6]}{'...' if len(rv)>6 else ''}")
        else:
            vals_eq = False
            print(f"    cannot compare series0 values (len mismatch or empty)")
        print(f"    => names_eq={names_eq}  series0_vals_eq={vals_eq}")
