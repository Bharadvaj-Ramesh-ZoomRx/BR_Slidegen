"""Scan refreshed decks for charts where the refresh wrote all-None values.
This is a Step 2 / refresh-pipeline issue surfaced via Step 7's data inspection.
"""
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation

REPO = Path(__file__).resolve().parents[1]
BASE = REPO / "output" / "step2_test_connected"

DECKS = {
    "atu_q1_26": {
        "src": REPO / "projects" / "J&J Rybrevant PET" / "Template" / "ZoomRx_UC_ATU_Report_Q1_'26.pptx",
        "ref": BASE / "atu_q1_26.pptx",
    },
    "creon_pet_w33": {
        "src": REPO / "projects" / "J&J Rybrevant PET" / "Template" / "CREON Share of Voice Study - W33.pptx",
        "ref": BASE / "creon_pet_w33.pptx",
    },
}


def chart_all_none(chart) -> bool:
    for plot in chart.plots:
        for s in plot.series:
            if any(v is not None for v in s.values):
                return False
    return True


def chart_has_data(chart) -> bool:
    for plot in chart.plots:
        for s in plot.series:
            if any(v is not None for v in s.values):
                return True
    return False


for deck_key, paths in DECKS.items():
    src = Presentation(str(paths["src"]))
    ref = Presentation(str(paths["ref"]))
    print(f"\n========== {deck_key} ==========")
    corrupted = []
    for s_idx, (s_slide, r_slide) in enumerate(zip(src.slides, ref.slides)):
        src_charts = [sh for sh in s_slide.shapes if sh.has_chart]
        ref_charts = [sh for sh in r_slide.shapes if sh.has_chart]
        if not src_charts or not ref_charts:
            continue
        for sc, rc in zip(src_charts, ref_charts):
            if chart_has_data(sc.chart) and chart_all_none(rc.chart):
                corrupted.append((s_idx + 1, sc.name))
    print(f"Charts with all-None refreshed values: {len(corrupted)}")
    for snum, name in corrupted[:30]:
        print(f"  slide {snum:3d}  chart {name!r}")
    if len(corrupted) > 30:
        print(f"  ... +{len(corrupted) - 30} more")
