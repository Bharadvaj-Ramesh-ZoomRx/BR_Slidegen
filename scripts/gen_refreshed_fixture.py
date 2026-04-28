"""Generate the Step 2 "test connected" deck via the dual-mode refresh engine.

Pipeline:
  source PPTX -> Connector specs -> full spec.json -> refresh_deck_from_spec()
                                                   -> output/step2_test_connected/<deck_key>.pptx

Usage:
    python scripts/gen_refreshed_fixture.py <deck_key>

Where <deck_key> is a key in tests/evals/fixtures.py::FIXTURE_DECKS.
"""
from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)

    deck_key = sys.argv[1]

    from tests.evals.fixtures import FIXTURE_DECKS
    source_pptx = FIXTURE_DECKS.get(deck_key)
    if source_pptx is None:
        print(f"ERROR: unknown deck key '{deck_key}'. Known keys: {list(FIXTURE_DECKS.keys())}")
        sys.exit(2)
    if not source_pptx.exists():
        print(f"ERROR: source deck not on disk: {source_pptx}")
        sys.exit(2)

    work_dir = REPO_ROOT / "output" / "step2_test_connected"
    work_dir.mkdir(parents=True, exist_ok=True)

    connector_specs_path = work_dir / f"{deck_key}_connector_specs.json"
    full_spec_path = work_dir / f"{deck_key}_full_spec.json"
    refreshed_pptx = work_dir / f"{deck_key}.pptx"

    # Step 1 — generate Connector specs via deck-reader
    print(f"[1/4] Generating Connector specs for {deck_key}...")
    from slidegen.deck_reader.tag_reader import generate_config_specs
    from slidegen.slide_spec.schema import dump_spec
    from pptx import Presentation

    specs, summary = generate_config_specs(str(source_pptx))
    connector_specs = [json.loads(dump_spec(s)) for s in specs]
    connector_specs_path.write_text(
        json.dumps(connector_specs, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    print(f"     wrote {len(connector_specs)} specs, {summary.tagged_shapes} tagged shapes")

    # Step 2 — build the full spec (data_sources + slides) using Vijay's helpers
    print(f"[2/4] Building full spec.json...")
    from tests.build_full_spec import build_data_sources, build_connected_slide

    data_sources = build_data_sources(connector_specs)
    print(f"     {len(data_sources)} unique data sources")

    # Map slide_index -> data_source key. Must match the keying logic in
    # build_full_spec._lineage_ds_key so each slide points to the right
    # bucket (encodes analysis ids, dynamic_latest_n, and segment_ids).
    from tests.build_full_spec import _lineage_ds_key
    slide_to_ds = {}
    for spec in connector_specs:
        si = spec["slide_index"]
        lin = spec.get("data_lineage", {})
        key = _lineage_ds_key(lin)
        if key:
            slide_to_ds[si] = key

    # Build slide entries and populate shape names from the PPTX
    prs = Presentation(str(source_pptx))
    slide_entries = []

    for spec in connector_specs:
        si = spec["slide_index"]
        ds_key = slide_to_ds.get(si)
        if not ds_key:
            continue
        slide_entry = build_connected_slide(spec, ds_key)

        # Fill in shape names by position matching (Vijay's convention)
        if si < len(prs.slides):
            pptx_slide = prs.slides[si]
            chart_shapes = sorted(
                [s for s in pptx_slide.shapes if s.has_chart],
                key=lambda x: (x.left or 0, x.top or 0),
            )
            table_shapes = sorted(
                [s for s in pptx_slide.shapes if s.has_table],
                key=lambda x: (x.left or 0, x.top or 0),
            )
            for comp in slide_entry["components"]:
                pos = comp.get("position", {})
                cl = pos.get("left", 0)
                ct = pos.get("top", 0)
                candidates = chart_shapes if comp["type"] == "chart" else table_shapes
                for shape in candidates:
                    sl = round(shape.left / 914400, 2) if shape.left else 0
                    st = round(shape.top / 914400, 2) if shape.top else 0
                    if abs(sl - cl) < 0.3 and abs(st - ct) < 0.3:
                        comp["name"] = shape.name
                        break

        slide_entries.append(slide_entry)

    full_spec = {
        "source_deck": source_pptx.name,
        "data_sources": data_sources,
        "slides": slide_entries,
    }
    full_spec_path.write_text(
        json.dumps(full_spec, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    n_raw_configs = sum(
        1 for s in slide_entries for c in s.get("components", [])
        if c.get("raw_pivot_config")
    )
    print(f"     wrote {full_spec_path.name}  slides={len(slide_entries)}  components with raw configs={n_raw_configs}")

    # Step 3 — run dual-mode refresh
    print(f"[3/4] Running dual-mode refresh...")
    from slidegen.intelligent_refresh import refresh_deck_from_spec

    result = refresh_deck_from_spec(
        spec_path=str(full_spec_path),
        pptx_path=str(source_pptx),
        output_path=str(refreshed_pptx),
    )

    # Save per-chart refresh status sidecar so evals can distinguish
    # "mapper succeeded but data didn't move" from "mapper failed".
    refresh_status_path = work_dir / f"{deck_key}_refresh_status.json"
    refresh_status_path.write_text(
        json.dumps(result, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )
    print(f"     refresh status sidecar -> {refresh_status_path.relative_to(REPO_ROOT)}")

    # Step 4 — summary
    print(f"\n[4/4] Refresh complete")
    if isinstance(result, dict):
        for k, v in result.items():
            if isinstance(v, (int, str, bool)):
                print(f"     {k}: {v}")

    print(f"\n[OK] Refreshed deck at: {refreshed_pptx.relative_to(REPO_ROOT)}")
    print(f"     size: {refreshed_pptx.stat().st_size / 1024:.0f} KB")
    print(f"\nNext: python -m tests.evals.end_to_end.generate_golden")


if __name__ == "__main__":
    main()
