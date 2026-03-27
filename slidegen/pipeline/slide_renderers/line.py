"""
line.py — Line chart and trended slide renderers.

Slide types:
  - trended_scorecard   (template slide 12 — multi-panel line chart grid)
  - trended_activity    (template slide 13 — side-by-side line + stacked column panels)
"""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ._shared import (
    SLIDE_W, CHART_TOP_STD, FOOTER_TOP,
    C_WHITE, C_GREY, C_FTGREY, C_LBGREY, C_HDRGREY,
    FONT_TEXT,
    _resolve_template, _slide_chrome, _make_legend,
    _pptx_table, _style_tbl_cell, _cell_bottom_border,
    textbox, solidrect, horiz_line,
    add_line_chart, add_stacked_column_chart,
    set_chart_plot_area,
    ProjectConfig, AskConfig, parse_color,
    logger,
)


# ── Trended Scorecard ────────────────────────────────────────────────────────
# Template slide 12: 5 small line charts stacked vertically with metric labels
# on the left. Each chart shows QoQ trend lines for 2-3 series.

def render_trended_scorecard(slide, config: ProjectConfig, ask: AskConfig,
                              data: dict, *, namer=None):
    """Render a trended scorecard with multiple mini line charts in a vertical grid.

    Data format (mock rows): list of dicts, each representing one chart panel:
        {
            "metric": "Share of Voice",
            "axis_label": "% of Interactions",
            "sample_label": "s",
            "series": [
                {"name": "J&J", "color": "#F95924", "values": [0.52, 0.56, ...]},
                {"name": "AZ",  "color": "#7638A4", "values": [0.48, 0.44, ...]},
            ]
        }

    ask.extra keys:
        categories: ["Q4'24", "Q1'25", "Q2'25", "Q3'25", "Q4'25"]
        panels: (data is read from data[ask.data_key])
        scale_as_pct: true  (multiply values by 100 for display)
        marker_size: 5
        line_width: 2.25
        label_font_size: 9
    """
    _slide_chrome(slide, config, ask)
    extra = ask.extra or {}

    rows = data.get(ask.data_key, [])
    if not rows:
        logger.warning("trended_scorecard: no data for key=%s", ask.data_key)
        return

    categories = extra.get("categories", [])
    scale_pct = extra.get("scale_as_pct", True)
    marker_size = extra.get("marker_size", 5)
    line_width = extra.get("line_width", 2.25)
    label_fsize = extra.get("label_font_size", 9)
    num_fmt = extra.get("num_fmt", '0"%"')
    font = config.font_body or FONT_TEXT

    n_panels = len(rows)

    # Layout dimensions
    left_margin = 0.30
    label_w = 2.10           # metric label column width
    chart_left = left_margin + label_w + 0.05
    chart_right_margin = 0.30
    chart_w = SLIDE_W - chart_left - chart_right_margin

    # Scorecard header row (category labels across top)
    header_top = CHART_TOP_STD - 0.02
    header_h = 0.30

    # Category header labels
    n_cats = len(categories)
    if n_cats > 0:
        cat_w = chart_w / n_cats
        for ci, cat in enumerate(categories):
            textbox(slide, cat,
                    chart_left + ci * cat_w, header_top,
                    cat_w, header_h,
                    fsize=7.5, bold=True, color=C_FTGREY,
                    align=PP_ALIGN.CENTER, font=font)

    # Panels — each gets a metric label + line chart
    panel_top = header_top + header_h + 0.02
    total_h = FOOTER_TOP - panel_top - 0.40  # leave room for legend
    panel_h = total_h / n_panels
    chart_h = panel_h - 0.04  # small gap between panels

    for pi, panel in enumerate(rows):
        y = panel_top + pi * panel_h
        metric = panel.get("metric", f"Metric {pi+1}")
        axis_label = panel.get("axis_label", "")
        series_data = panel.get("series", [])

        # Grey background band for alternating rows
        if pi % 2 == 0:
            solidrect(slide, left_margin, y, SLIDE_W - 2 * left_margin, chart_h,
                      C_LBGREY)

        # Metric label (left column) — bold metric name + smaller axis label, vertically centered
        shape = slide.shapes.add_textbox(
            Inches(left_margin + 0.05), Inches(y),
            Inches(label_w - 0.10), Inches(chart_h))
        tf = shape.text_frame
        tf.word_wrap = True
        # Vertical centering
        from pptx.oxml.ns import qn as _qn
        bodyPr = tf._txBody.find(_qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", "ctr")

        # Line 1: metric name (bold, larger)
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.LEFT
        run1 = p1.add_run()
        run1.text = metric
        run1.font.size = Pt(9)
        run1.font.bold = True
        run1.font.color.rgb = C_FTGREY
        run1.font.name = font

        # Line 2: axis label (smaller, not bold)
        if axis_label:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            p2.space_before = Pt(1)
            run2 = p2.add_run()
            run2.text = axis_label
            run2.font.size = Pt(7)
            run2.font.bold = False
            run2.font.color.rgb = C_GREY
            run2.font.name = font

        # Build series for chart
        if not series_data or not categories:
            continue

        series_list = []
        colors = []
        for s in series_data:
            vals = s.get("values", [])
            if scale_pct:
                vals = [v * 100 if v <= 1.0 else v for v in vals]
            series_list.append((s.get("name", ""), vals))
            colors.append(parse_color(s.get("color", "#999999")))

        # Determine scale from data
        all_vals = [v for _, vs in series_list for v in vs if v is not None]
        if all_vals:
            data_min = min(all_vals)
            data_max = max(all_vals)
            scale_min = max(0, data_min - 15)
            scale_max = min(100, data_max + 15)
        else:
            scale_min, scale_max = 0, 100

        panel_fmt = panel.get("num_fmt", num_fmt)

        cf, ch = add_line_chart(
            slide, categories, series_list,
            left=chart_left, top=y, width=chart_w, height=chart_h,
            colors=colors, marker_size=marker_size,
            line_width_pt=line_width,
            show_labels=True, label_fsize=label_fsize,
            num_fmt=panel_fmt,
            scale_min=scale_min, scale_max=scale_max,
            hide_axes=True, font_name=font,
        )
        # Expand plot area to fill chart frame
        set_chart_plot_area(ch, x=0.02, y=0.10, w=0.96, h=0.80)

    # Legend at bottom
    if rows:
        first_series = rows[0].get("series", [])
        legend_items = [(parse_color(s.get("color", "#999999")), s.get("name", ""))
                        for s in first_series]
        legend_top = FOOTER_TOP - 0.25
        _make_legend(slide, legend_items, left=chart_left, top=legend_top,
                     font_name=font,
                     center_over=(chart_left, chart_w))


# ── Trended Activity ─────────────────────────────────────────────────────────
# Template slide 13: Side-by-side panels (Reach / SOV / Frequency)
# Mix of line charts and stacked column charts.

def render_trended_activity(slide, config: ProjectConfig, ask: AskConfig,
                             data: dict, *, namer=None):
    """Render side-by-side trended panels with line and/or stacked column charts.

    Data format (mock rows): list of dicts, each representing one panel:
        {
            "metric": "Reach",
            "chart_type": "line",          # "line" or "stacked_column"
            "axis_label": "% of HCPs",
            "sample_labels": ["Q3'25\\n(n = 134)", "Q4'25\\n(n = 149)"],
            "series": [
                {"name": "J&J", "color": "#F95924", "values": [0.72, 0.69]},
                {"name": "AZ",  "color": "#7638A4", "values": [0.75, 0.72]},
            ]
        }

    ask.extra keys:
        categories: ["Q3'25", "Q4'25"]
        insight_text: "..."    (optional insight callout below charts)
    """
    _slide_chrome(slide, config, ask)
    extra = ask.extra or {}

    rows = data.get(ask.data_key, [])
    if not rows:
        logger.warning("trended_activity: no data for key=%s", ask.data_key)
        return

    categories = extra.get("categories", [])
    scale_pct = extra.get("scale_as_pct", True)
    font = config.font_body or FONT_TEXT

    n_panels = len(rows)

    # Layout — divide slide width into equal panels
    left_margin = 0.40
    right_margin = 0.40
    panel_gap = 0.30
    total_w = SLIDE_W - left_margin - right_margin - (n_panels - 1) * panel_gap
    panel_w = total_w / n_panels

    # Vertical layout
    metric_label_top = CHART_TOP_STD + 0.05
    metric_label_h = 0.30
    chart_top = metric_label_top + metric_label_h + 0.10
    chart_bottom = FOOTER_TOP - 0.60  # room for sample labels + legend
    chart_h = chart_bottom - chart_top

    # Collect legend items from first panel
    legend_items = []

    for pi, panel in enumerate(rows):
        px = left_margin + pi * (panel_w + panel_gap)
        metric = panel.get("metric", f"Panel {pi+1}")
        chart_type = panel.get("chart_type", "line")
        axis_label = panel.get("axis_label", "")
        sample_labels = panel.get("sample_labels", [])
        series_data = panel.get("series", [])

        # Metric header
        textbox(slide, metric,
                px, metric_label_top, panel_w, metric_label_h,
                fsize=10, bold=True, color=C_FTGREY,
                align=PP_ALIGN.CENTER, font=font)

        # Axis label below metric
        if axis_label:
            textbox(slide, axis_label,
                    px, metric_label_top + 0.22, panel_w, 0.20,
                    fsize=7, bold=False, color=C_GREY,
                    align=PP_ALIGN.CENTER, font=font)

        if not series_data or not categories:
            continue

        # Build series
        series_list = []
        colors = []
        for s in series_data:
            vals = s.get("values", [])
            if scale_pct:
                vals = [v * 100 if v <= 1.0 else v for v in vals]
            series_list.append((s.get("name", ""), vals))
            colors.append(parse_color(s.get("color", "#999999")))

        # Collect legend items from first panel only
        if pi == 0:
            legend_items = [(parse_color(s.get("color", "#999999")), s.get("name", ""))
                            for s in series_data]

        # Determine scale
        all_vals = [v for _, vs in series_list for v in vs if v is not None]
        if all_vals:
            data_min = min(all_vals)
            data_max = max(all_vals)
            scale_min = max(0, data_min - 15)
            scale_max = min(100, data_max + 15)
        else:
            scale_min, scale_max = 0, 100

        # Per-panel num_fmt override (e.g. frequency uses "0.0" not "0%")
        panel_fmt = panel.get("num_fmt", '0"%"')

        if chart_type == "stacked_column":
            cf, ch = add_stacked_column_chart(
                slide, categories, series_list,
                left=px, top=chart_top, width=panel_w, height=chart_h,
                colors=colors, label_fsize=9, num_fmt=panel_fmt,
                gap=100, font_name=font,
            )
            set_chart_plot_area(ch, x=0.05, y=0.02, w=0.90, h=0.85)
        else:
            cf, ch = add_line_chart(
                slide, categories, series_list,
                left=px, top=chart_top, width=panel_w, height=chart_h,
                colors=colors, marker_size=6, line_width_pt=2.25,
                show_labels=True, label_fsize=9, num_fmt=panel_fmt,
                scale_min=scale_min, scale_max=scale_max,
                hide_axes=True, font_name=font,
            )
            set_chart_plot_area(ch, x=0.05, y=0.08, w=0.90, h=0.80)

        # Sample labels below chart
        if sample_labels:
            sl_top = chart_top + chart_h + 0.02
            sl_w = panel_w / len(sample_labels)
            for si, sl in enumerate(sample_labels):
                textbox(slide, sl,
                        px + si * sl_w, sl_top, sl_w, 0.30,
                        fsize=6.5, bold=False, color=C_GREY,
                        align=PP_ALIGN.CENTER, font=font)

    # Legend centered below all panels
    if legend_items:
        legend_top = FOOTER_TOP - 0.25
        block_left = left_margin
        block_w = SLIDE_W - left_margin - right_margin
        _make_legend(slide, legend_items, left=block_left, top=legend_top,
                     font_name=font, center_over=(block_left, block_w))

    # Optional insight callout
    insight = extra.get("insight_text")
    if insight:
        insight_top = chart_top + chart_h + 0.35
        insight_left = left_margin
        insight_w = SLIDE_W - 2 * left_margin
        textbox(slide, insight,
                insight_left, insight_top, insight_w, 0.30,
                fsize=7.5, bold=False, color=C_FTGREY, font=font)
