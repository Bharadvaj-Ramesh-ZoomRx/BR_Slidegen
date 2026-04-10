"""
qualitative.py — Qualitative theme analysis slide renderer (client-ready).

Renders a dedicated qualitative slide with:
  Left panel:  Native PPT bar chart (BAR_CLUSTERED) + TABLE for theme labels
               + sub-title with question text + sample size
               + x-axis label ("% of respondents")
  Right panel: 4 grouped quote boxes with line separator + attribution above quote

Data source: qualitative_data.json themes from validated_analysis.md
  ask.extra.themes: [{"theme": str, "pct": float, "count": int}, ...]
  ask.extra.quotes: [{"text": str, "attribution": str}, ...]
  ask.extra.qual_source: str (e.g., "Q1.53A (n=68)")
  ask.extra.qual_subtitle: str (optional, e.g., "Prescribing Conversion – Rationale (Unaided)")
  ask.extra.qual_sample: str (optional, e.g., "s = 102")
"""

from __future__ import annotations

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu

from ._shared import (
    CHART_TOP_STD, SLIDE_W, FOOTER_TOP,
    _resolve_template, _slide_chrome,
    FONT_HDR, FONT_BODY,
    C_GREEN, C_GREY, C_LBGREY, C_HDRGREY, C_WHITE, C_RED,
    textbox, solidrect,
    ProjectConfig, AskConfig,
    logger,
)
from slidegen.pptx_utils.lxml_helpers import (
    hide_axis, invert_cat_axis, set_plot_area_gap,
    suppress_para_bullets, suppress_cat_axis_bullets,
    hide_cat_labels, set_chart_plot_area,
    set_series_color, set_series_no_border,
    set_val_axis_scale, _get_or_add,
)
from slidegen.pptx_utils.charts import enable_data_labels
from pptx.oxml.ns import qn


# ── Layout constants ─────────────────────────────────────────────────────

_MARGIN = 0.30
_PANEL_GAP = 0.25
_PANEL_TOP = 1.85
_PANEL_BOTTOM = 6.50

# Two-panel split: left ~55% (chart area), right ~45% (quotes)
_LEFT_PANEL_L = _MARGIN
_LEFT_PANEL_W = (SLIDE_W - 2 * _MARGIN - _PANEL_GAP) * 0.55
_RIGHT_PANEL_L = _LEFT_PANEL_L + _LEFT_PANEL_W + _PANEL_GAP
_RIGHT_PANEL_W = SLIDE_W - _RIGHT_PANEL_L - _MARGIN

# Sub-title table
_SUBTITLE_TOP = _PANEL_TOP - 0.22
_SUBTITLE_H = 0.45

# Theme label table + chart layout
_LABEL_TABLE_L = _LEFT_PANEL_L
_LABEL_TABLE_W = 2.20
_CHART_L = _LABEL_TABLE_L + _LABEL_TABLE_W + 0.08
_CHART_W = _LEFT_PANEL_W - _LABEL_TABLE_W - 0.08
_CHART_TOP = _PANEL_TOP + _SUBTITLE_H + 0.20
_CHART_BOTTOM = _PANEL_BOTTOM - 0.35  # room for x-axis label
_CHART_H_MAX = _CHART_BOTTOM - _CHART_TOP

# X-axis label
_XAXIS_LABEL_TOP = _CHART_BOTTOM + 0.02

# Quote group layout
_QUOTE_TOP = _PANEL_TOP + 0.10
_QUOTE_BOTTOM = _PANEL_BOTTOM
_QUOTE_GAP = 0.10
_QUOTE_MAX = 4

# Colors
_BAR_COLOR = RGBColor(0xF7, 0x58, 0x24)  # RYB orange (default)
_QUOTE_TEXT_COLOR = RGBColor(0x33, 0x33, 0x33)
_ATTR_COLOR = RGBColor(0x00, 0x00, 0x00)
_LINE_COLOR = RGBColor(0x00, 0x00, 0x00)
_SUBTITLE_BG = RGBColor(0xF0, 0xF0, 0xF0)


def render_qual_theme_analysis(slide, config: ProjectConfig, ask: AskConfig,
                                data: dict, *, namer=None):
    """Qualitative theme analysis: native PPT bar chart + label table + grouped quotes."""
    _slide_chrome(slide, config, ask)

    extra = ask.extra or {}
    themes = extra.get("themes", [])
    quotes = extra.get("quotes", [])
    qual_source = extra.get("qual_source", "")
    qual_subtitle = extra.get("qual_subtitle", "")
    qual_sample = extra.get("qual_sample", "")

    if not themes:
        if isinstance(data, dict):
            themes = data.get("themes", [])
            quotes = data.get("quotes", quotes)
            qual_source = data.get("qual_source", qual_source)

    if not themes:
        textbox(slide, "No qualitative theme data available",
                2.0, 3.0, 9.0, 0.5, fsize=12, color=C_GREY)
        return

    # Resolve bar color from brand config
    bar_color = config.primary.color_current if config.primary else _BAR_COLOR

    # ── Sub-title table (question text + sample size) ────────────────
    subtitle_top = _PANEL_TOP - 0.10
    if qual_subtitle or qual_sample:
        sub_text = qual_subtitle
        if qual_sample:
            sub_text += f"\n({qual_sample})" if sub_text else f"({qual_sample})"

        tbl_shape = slide.shapes.add_table(
            1, 1,
            Inches(_CHART_L), Inches(subtitle_top),
            Inches(_CHART_W), Inches(_SUBTITLE_H))
        tbl = tbl_shape.table
        tbl.rows[0].height = Inches(_SUBTITLE_H)
        tbl.columns[0].width = Inches(_CHART_W)
        cell = tbl.cell(0, 0)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _SUBTITLE_BG
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        suppress_para_bullets(p._p)
        run = p.add_run()
        run.text = sub_text
        run.font.size = Pt(10)
        run.font.bold = False
        run.font.color.rgb = C_GREY
        run.font.name = config.font_body

    # ── Compute chart area ───────────────────────────────────────────
    chart_top = subtitle_top + _SUBTITLE_H + 0.12 if (qual_subtitle or qual_sample) else _PANEL_TOP + 0.10
    chart_h = min(_CHART_H_MAX, _CHART_BOTTOM - chart_top)
    n_themes = len(themes)

    # ── Left panel: Theme label TABLE ────────────────────────────────
    row_h = chart_h / n_themes

    tbl_shape = slide.shapes.add_table(
        n_themes, 1,
        Inches(_LABEL_TABLE_L), Inches(chart_top),
        Inches(_LABEL_TABLE_W), Inches(chart_h))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(_LABEL_TABLE_W)

    for i, theme in enumerate(themes):
        tbl.rows[i].height = Emu(int(row_h * 914400))
        cell = tbl.cell(i, 0)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_WHITE
        # Remove cell margins for tighter alignment
        cell.margin_left = Inches(0.05)
        cell.margin_right = Inches(0.05)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)

        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        suppress_para_bullets(p._p)
        run = p.add_run()
        run.text = theme.get("theme", f"Theme {i+1}")
        run.font.size = Pt(10)
        run.font.bold = False
        run.font.color.rgb = RGBColor(0x50, 0x50, 0x50)
        run.font.name = config.font_body

        # Vertical center
        from slidegen.pptx_utils.lxml_helpers import cell_vcenter
        cell_vcenter(cell)

    # Remove table borders
    _remove_table_borders(tbl_shape)

    # ── Left panel: Native PPT BAR_CLUSTERED chart ───────────────────
    # Categories are theme names (hidden — using external label table)
    # Values are percentages, reversed order for PPT bottom-to-top rendering
    cat_names = [t.get("theme", "") for t in themes]
    pct_values = [t.get("pct", 0) for t in themes]

    # PPT bar charts render bottom-to-top, so reverse to match top-to-bottom label table
    chart_data = CategoryChartData()
    chart_data.categories = list(reversed(cat_names))
    chart_data.add_series("Themes", tuple(reversed(pct_values)))

    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(_CHART_L), Inches(chart_top),
        Inches(_CHART_W), Inches(chart_h),
        chart_data)
    ch = cf.chart
    ch.has_legend = False
    ch.has_title = False

    # Style the series
    series = ch.series[0]
    set_series_color(series, bar_color)
    set_series_no_border(series)

    # Data labels: show percentage with "%" suffix, white text inside bars
    enable_data_labels(series, C_WHITE, fsize=10, pos="inEnd",
                       num_fmt='0"%"', font_name=config.font_body)

    # Axes: hide value axis, hide category labels (external table), set gap
    hide_axis(ch, "val")
    ch.category_axis.has_major_gridlines = False
    hide_cat_labels(ch)
    set_plot_area_gap(ch, 80)

    # Expand plot area to fill the chart frame (minimize internal padding)
    set_chart_plot_area(ch, x=0.02, y=0.02, w=0.96, h=0.96)

    # Set value axis max to ensure bars don't overflow
    max_pct = max(pct_values) if pct_values else 100
    val_axis_max = min(100, max_pct * 1.3)  # 30% headroom
    set_val_axis_scale(ch, 0, val_axis_max)

    # ── X-axis label ─────────────────────────────────────────────────
    xaxis_top = chart_top + chart_h + 0.02
    textbox(slide, "% of respondents",
            _CHART_L, xaxis_top, _CHART_W, 0.25,
            fsize=12, color=RGBColor(0x50, 0x50, 0x50),
            align=PP_ALIGN.CENTER, font=config.font_body)

    # ── Right panel: Grouped quote boxes ─────────────────────────────
    if not quotes:
        if namer:
            namer.name_remaining(slide)
        return

    n_quotes = min(len(quotes), _QUOTE_MAX)
    available_h = _QUOTE_BOTTOM - _QUOTE_TOP
    quote_h = min(1.20, (available_h - (n_quotes - 1) * _QUOTE_GAP) / n_quotes)

    y = _QUOTE_TOP
    for i, quote in enumerate(quotes[:n_quotes]):
        q_text = quote.get("text", "")
        attribution = quote.get("attribution", "")

        _render_quote_group(
            slide, _RIGHT_PANEL_L, y, _RIGHT_PANEL_W, quote_h,
            q_text, attribution, config.font_body, bar_color)

        y += quote_h + _QUOTE_GAP
        if y > _QUOTE_BOTTOM:
            break

    # ── Namer ─────────────────────────────────────────────────────────
    if namer:
        namer.name_remaining(slide)


def _render_quote_group(slide, left: float, top: float, width: float, height: float,
                         quote_text: str, attribution: str, font_name: str,
                         accent_color: RGBColor):
    """Render a single quote group matching slide 36 manual style:
    - Attribution label (bold italic) above a thin black horizontal line
    - Quote text (italic) below the line
    """
    attr_h = 0.28
    line_y = top + attr_h
    quote_text_top = line_y + 0.06
    quote_text_h = height - attr_h - 0.06

    # Left accent bar (thin colored stripe)
    accent_w = 0.04
    solidrect(slide, left, top, accent_w, height, fill=accent_color)

    text_left = left + accent_w + 0.10
    text_w = width - accent_w - 0.15

    # Attribution label above line
    if attribution:
        textbox(slide, attribution,
                text_left, top, text_w, attr_h,
                fsize=10, bold=True, italic=True,
                color=_ATTR_COLOR, font=font_name)

    # Horizontal separator line
    from slidegen.pptx_utils.shapes import horiz_line
    horiz_line(slide, text_left, line_y, text_w,
               color=_LINE_COLOR, width_pt=0.75)

    # Quote text below line
    if quote_text:
        q_display = f"\u201c{quote_text}\u201d"
        textbox(slide, q_display,
                text_left, quote_text_top, text_w, quote_text_h,
                fsize=11, italic=True,
                color=_QUOTE_TEXT_COLOR, font=font_name)


def _remove_table_borders(tbl_shape):
    """Remove all cell borders from a table shape for a clean label-only look."""
    tbl = tbl_shape.table
    for row in tbl.rows:
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.find(qn("a:tcPr"))
            if tcPr is None:
                tcPr = tc.makeelement(qn("a:tcPr"), {})
                tc.append(tcPr)
            for border_name in ["lnL", "lnR", "lnT", "lnB"]:
                ln = tcPr.find(qn(f"a:{border_name}"))
                if ln is not None:
                    tcPr.remove(ln)
                ln = tc.makeelement(qn(f"a:{border_name}"), {"w": "0"})
                noFill = tc.makeelement(qn("a:noFill"), {})
                ln.append(noFill)
                tcPr.append(ln)
