"""
bar_dual.py — Dual-bar slide renderers (dual_bar_with_delta, dual_bar_qoq).
"""

from __future__ import annotations

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from ._shared import (
    # constants
    CHART_TOP_STD, CHART_TOP_DUAL, LEGEND_GAP, SLIDE_W,
    DUAL_HEADER_ROW_TOP, DUAL_HEADER_ROW_H, DUAL_CHART_TOP,
    DUAL_MR_LEFT, DUAL_MR_WIDTH, DUAL_MR_DELTA_LEFT,
    DUAL_ME_LEFT, DUAL_ME_WIDTH, DUAL_ME_DELTA_LEFT,
    DUAL_DELTA_WIDTH, DUAL_SEPARATOR_X, DUAL_MAX_CHART_HEIGHT,
    DUAL_T_CAT_LEFT, DUAL_T_CAT_W, DUAL_T_MAIN_LEFT,
    DUAL_T_MR_COL_W, DUAL_T_ME_COL_W,
    DUAL_T_HDR_H, DUAL_T_ROW_H,
    DUAL_T_MR_CHART_L, DUAL_T_MR_CHART_W, DUAL_T_MR_DELTA_L, DUAL_T_MR_DELTA_W,
    DUAL_T_ME_CHART_L, DUAL_T_ME_CHART_W, DUAL_T_ME_DELTA_L, DUAL_T_ME_DELTA_W,
    DUAL_T_CALLOUT_L, DUAL_T_CALLOUT_W,
    DELTA_COL_NARROW,
    MAX_DUAL_CHART_HEIGHT, MIN_CHART_HEIGHT,
    ROW_SCALE_FACTOR,
    BAR_GAP_STD, CLUSTERED_OVERLAP,
    HEADER_ROW_HEIGHT_IN,
    LABEL_MAX_DUAL,
    # helpers
    _slide_chrome, _get_brand_colors, _sort_data, _make_legend,
    _pptx_table, _style_tbl_cell, _cell_bottom_border, _cap_chart_h,
    # pptx_utils
    C_GREEN, C_WHITE, C_GREY, C_FTGREY, C_LBGREY, C_RED,
    PP_ALIGN,
    textbox, dashed_separator, callout_box,
    hide_axis, set_series_color, set_plot_area_gap, set_overlap,
    set_series_no_border, invert_cat_axis, hide_cat_labels, suppress_cat_axis_bullets,
    set_chart_plot_area, set_val_axis_scale,
    enable_data_labels,
    add_single_bar_chart, add_delta_table,
    chart_header_row,
    # data_loaders
    delta,
    # project_config types
    ProjectConfig, AskConfig, parse_color,
)


def render_dual_bar_with_delta(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Two side-by-side bar charts (e.g. MR + ME) each with delta columns.

    Uses Archetype 3 layout: MR wider (7.30") with cat labels,
    ME narrower (4.40") without cat labels, red header row above charts.
    """
    _slide_chrome(slide, config, ask)

    rows = data.get(ask.data_key, [])
    rows = _sort_data(rows, ask.sort_by, ask.sort_desc)
    if not rows:
        textbox(slide, "Data not available", 2, 3, 8, 1, fsize=14, color=C_RED)
        return

    color_current, color_prior = _get_brand_colors(config, ask)
    font = config.font_body
    extra = ask.extra

    left_cfg = extra.get("left", {})
    right_cfg = extra.get("right", {})
    left_prefix = left_cfg.get("field_prefix", "mr")
    right_prefix = right_cfg.get("field_prefix", "me")
    callouts = extra.get("callouts", [])
    narrow = bool(callouts)

    labels = [r.get("short") or r.get("desc", "") for r in rows]
    n = len(labels)
    chart_top = DUAL_CHART_TOP
    has_ax = bool(extra.get("left", {}).get("axis_label") or extra.get("right", {}).get("axis_label"))
    chart_h = _cap_chart_h(min(DUAL_MAX_CHART_HEIGHT, n * 0.42), chart_top, has_ax)
    row_h = chart_h / max(n, 1)

    # ── Common data extraction ────────────────────────────────────────
    cat_header = extra.get("category_header", "")
    left_label = left_cfg.get("label", "Left (%)")
    right_label = right_cfg.get("label", "Right (%)")
    left_delta_header = left_cfg.get("delta_header", "Δ")
    right_delta_header = right_cfg.get("delta_header", "Δ")
    left_axis_label = left_cfg.get("axis_label", "")
    right_axis_label = right_cfg.get("axis_label", "")

    # Left data
    left_current = [r.get(f"{left_prefix}_current") or 0 for r in rows]
    left_prior = [r.get(f"{left_prefix}_prior") for r in rows]

    # Right data
    right_current = [r.get(f"{right_prefix}_current") or 0 for r in rows]
    right_prior = [r.get(f"{right_prefix}_prior") for r in rows]

    # Shared axis scale
    all_vals = left_current + right_current
    axis_max = max(all_vals) if all_vals else 100
    axis_max = min(100, ((int(axis_max) // 10) + 1) * 10)

    left_deltas = [delta(c, p) if p is not None else None
                   for c, p in zip(left_current, left_prior)]
    right_deltas = [delta(c, p) if p is not None else None
                    for c, p in zip(right_current, right_prior)]

    if narrow:
        # ── NARROW MODE: template-matched pptx table layout ──────────────
        table_top = DUAL_CHART_TOP           # table starts here (below section bar)
        hdr_h = DUAL_T_HDR_H                 # header row height
        chart_top_n = table_top + hdr_h      # charts/data start below header
        has_ax_label = bool(left_axis_label or right_axis_label)
        chart_h_n = _cap_chart_h(n * DUAL_T_ROW_H, chart_top_n, has_ax_label)
        row_h = chart_h_n / max(n, 1)        # recalculate per-row height

        row_heights = [hdr_h] + [row_h] * n

        # 1. Category (Tag) pptx table
        _, cat_tbl = _pptx_table(
            slide, [DUAL_T_CAT_W], row_heights, DUAL_T_CAT_LEFT, table_top)
        _style_tbl_cell(cat_tbl.cell(0, 0), cat_header,
                        bg=C_RED, fg=C_WHITE, fsize=8, bold=True,
                        align=PP_ALIGN.CENTER, font=font)
        for i, label in enumerate(labels):
            cell = cat_tbl.cell(i + 1, 0)
            _style_tbl_cell(cell, label,
                            bg=C_LBGREY if i % 2 == 0 else C_WHITE,
                            fg=C_GREY, fsize=7.5, align=PP_ALIGN.RIGHT, font=font,
                            ml=0.04, mr=0.06)
            cell.text_frame.word_wrap = True

        # 2. Main header table (Recall | Effectiveness columns)
        _, main_tbl = _pptx_table(
            slide, [DUAL_T_MR_COL_W, DUAL_T_ME_COL_W], row_heights,
            DUAL_T_MAIN_LEFT, table_top)
        _style_tbl_cell(main_tbl.cell(0, 0), left_label,
                        bg=C_RED, fg=C_WHITE, fsize=8, bold=True,
                        align=PP_ALIGN.CENTER, font=font)
        _style_tbl_cell(main_tbl.cell(0, 1), right_label,
                        bg=C_RED, fg=C_WHITE, fsize=8, bold=True,
                        align=PP_ALIGN.CENTER, font=font)
        for i in range(n):
            for ci in range(2):
                main_tbl.cell(i + 1, ci).fill.background()

        # 3. Recall chart (overlaid, no category labels)
        cf1, ch1 = add_single_bar_chart(
            slide, labels, left_current,
            left=DUAL_T_MR_CHART_L, top=chart_top_n,
            width=DUAL_T_MR_CHART_W, height=chart_h_n,
            fill_color=color_current, font_name=font,
        )
        ch1.has_title = False
        hide_cat_labels(ch1)
        set_val_axis_scale(ch1, 0, axis_max)
        set_chart_plot_area(ch1, x=0.0, y=0.0, w=1.0, h=1.0)

        # 4. Recall delta (no header — header is in main_tbl row 0)
        add_delta_table(
            slide, left_deltas,
            left=DUAL_T_MR_DELTA_L, top=chart_top_n,
            width=DUAL_T_MR_DELTA_W, row_height=row_h,
            header_text=left_delta_header, font_name=font,
            show_header=False,
        )

        # 5. Effectiveness chart (overlaid, no category labels)
        cd2 = CategoryChartData()
        cd2.categories = labels
        cd2.add_series(config.period_current, right_current)
        cf2 = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            Inches(DUAL_T_ME_CHART_L), Inches(chart_top_n),
            Inches(DUAL_T_ME_CHART_W), Inches(chart_h_n), cd2)
        ch2 = cf2.chart
        ch2.has_legend = False
        ch2.has_title = False
        s2 = ch2.series[0]
        set_series_color(s2, color_current)
        set_series_no_border(s2)
        enable_data_labels(s2, C_WHITE, pos="inEnd", font_name=font)
        hide_axis(ch2, "val")
        hide_cat_labels(ch2)
        invert_cat_axis(ch2)
        set_plot_area_gap(ch2, BAR_GAP_STD)
        set_val_axis_scale(ch2, 0, axis_max)
        set_chart_plot_area(ch2, x=0.02, y=0.0, w=0.98, h=1.0)

        # 6. Effectiveness delta (no header)
        add_delta_table(
            slide, right_deltas,
            left=DUAL_T_ME_DELTA_L, top=chart_top_n,
            width=DUAL_T_ME_DELTA_W, row_height=row_h,
            header_text=right_delta_header, font_name=font,
            show_header=False,
        )

        # 7. Callout boxes (stacked vertically)
        n_cb = len(callouts)
        cb_area_h = chart_h_n
        box_h = min(1.10, cb_area_h / max(n_cb, 1) - 0.10)
        spacing = (cb_area_h - n_cb * box_h) / max(n_cb + 1, 1)
        for i, cb in enumerate(callouts):
            cb_top = chart_top_n + spacing * (i + 1) + box_h * i
            cb_color = parse_color(cb.get("color", "")) if cb.get("color") else None
            callout_box(slide,
                        left=DUAL_T_CALLOUT_L, top=cb_top,
                        width=DUAL_T_CALLOUT_W, height=box_h,
                        text=cb.get("text", ""), border_color=cb_color,
                        fsize=cb.get("fsize", 8))

        # 8. Axis sub-labels
        ax_y = chart_top_n + chart_h_n + 0.04
        if left_axis_label:
            textbox(slide, left_axis_label,
                    DUAL_T_MR_CHART_L, ax_y,
                    DUAL_T_MR_CHART_W + DUAL_T_MR_DELTA_W, 0.20,
                    fsize=7, color=C_FTGREY, align=PP_ALIGN.CENTER, font=font)
        if right_axis_label:
            textbox(slide, right_axis_label,
                    DUAL_T_ME_CHART_L, ax_y,
                    DUAL_T_ME_CHART_W + DUAL_T_ME_DELTA_W, 0.20,
                    fsize=7, color=C_FTGREY, align=PP_ALIGN.CENTER, font=font)

        # 9. Legend
        ly = ax_y + (0.24 if left_axis_label or right_axis_label else LEGEND_GAP)
        sample = config.sample_sizes.get(ask.brand or "primary")
        n_label = f" (n={sample.current})" if sample else ""
        _make_legend(slide, [
            (color_current, f"{config.period_current}{n_label}"),
            (C_GREEN, "Positive Δ"),
            (C_RED, "Negative Δ"),
        ], 0, ly, font, center_over=(0, SLIDE_W))

    else:
        # ── WIDE MODE: existing shape-primitive layout ────────────────────
        delta_row_h = (chart_h - HEADER_ROW_HEIGHT_IN) / max(n, 1)

        header_columns = [
            {"label": cat_header, "left": DUAL_MR_LEFT, "width": 2.0,
             "align": PP_ALIGN.LEFT},
            {"label": left_label,
             "left": DUAL_MR_LEFT + 2.0,
             "width": DUAL_MR_WIDTH + DUAL_DELTA_WIDTH - 2.0},
            {"label": right_label,
             "left": DUAL_ME_LEFT,
             "width": DUAL_ME_WIDTH + DUAL_DELTA_WIDTH},
        ]
        chart_header_row(slide, header_columns, top=DUAL_HEADER_ROW_TOP,
                         height=DUAL_HEADER_ROW_H, font=font)

        # Left chart (with category labels)
        cf1, ch1 = add_single_bar_chart(
            slide, labels, left_current,
            left=DUAL_MR_LEFT, top=chart_top, width=DUAL_MR_WIDTH, height=chart_h,
            fill_color=color_current, font_name=font,
        )
        ch1.has_title = False
        set_val_axis_scale(ch1, 0, axis_max)
        set_chart_plot_area(ch1, y=0.04, h=0.90)

        add_delta_table(
            slide, left_deltas,
            left=DUAL_MR_DELTA_LEFT, top=chart_top, width=DUAL_DELTA_WIDTH,
            row_height=delta_row_h,
            header_text=left_delta_header, font_name=font,
        )

        sep_h = chart_h + DUAL_HEADER_ROW_H
        dashed_separator(slide, DUAL_SEPARATOR_X, DUAL_HEADER_ROW_TOP,
                         sep_h, color=C_FTGREY, width_pt=0.5, vertical=True)

        cd2 = CategoryChartData()
        cd2.categories = labels
        cd2.add_series(config.period_current, right_current)
        cf2 = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            Inches(DUAL_ME_LEFT), Inches(chart_top),
            Inches(DUAL_ME_WIDTH), Inches(chart_h), cd2)
        ch2 = cf2.chart
        ch2.has_legend = False
        ch2.has_title = False
        s2 = ch2.series[0]
        set_series_color(s2, color_current)
        set_series_no_border(s2)
        enable_data_labels(s2, C_WHITE, pos="inEnd", font_name=font)
        hide_axis(ch2, "val")
        hide_cat_labels(ch2)
        invert_cat_axis(ch2)
        set_plot_area_gap(ch2, BAR_GAP_STD)
        set_val_axis_scale(ch2, 0, axis_max)
        set_chart_plot_area(ch2, x=0.02, y=0.04, w=0.94, h=0.90)

        add_delta_table(
            slide, right_deltas,
            left=DUAL_ME_DELTA_LEFT, top=chart_top, width=DUAL_DELTA_WIDTH,
            row_height=delta_row_h,
            header_text=right_delta_header, font_name=font,
        )

        ax_y = chart_top + chart_h + 0.04
        if left_axis_label:
            textbox(slide, left_axis_label,
                    DUAL_MR_LEFT, ax_y, DUAL_MR_WIDTH + DUAL_DELTA_WIDTH, 0.20,
                    fsize=7, color=C_FTGREY, align=PP_ALIGN.CENTER, font=font)
        if right_axis_label:
            textbox(slide, right_axis_label,
                    DUAL_ME_LEFT, ax_y, DUAL_ME_WIDTH + DUAL_DELTA_WIDTH, 0.20,
                    fsize=7, color=C_FTGREY, align=PP_ALIGN.CENTER, font=font)

        ly = chart_top + chart_h + (0.28 if left_axis_label or right_axis_label else LEGEND_GAP)
        sample = config.sample_sizes.get(ask.brand or "primary")
        n_label = f" (n={sample.current})" if sample else ""
        _make_legend(slide, [
            (color_current, f"{config.period_current}{n_label}"),
            (C_GREEN, "Positive Δ"),
            (C_RED, "Negative Δ"),
        ], 0, ly, font, center_over=(0, SLIDE_W))


def render_dual_bar_qoq(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Two side-by-side clustered bar charts (Q4 vs Q3) each with delta columns."""
    _slide_chrome(slide, config, ask)

    rows = data.get(ask.data_key, [])
    rows = _sort_data(rows, ask.sort_by, ask.sort_desc)
    if not rows:
        textbox(slide, "Data not available", 2, 3, 8, 1, fsize=14, color=C_RED)
        return

    color_current, color_prior = _get_brand_colors(config, ask)
    font = config.font_body
    extra = ask.extra

    left_cfg = extra.get("left", {})
    right_cfg = extra.get("right", {})
    left_prefix = left_cfg.get("field_prefix", "believ")
    right_prefix = right_cfg.get("field_prefix", "me")

    labels = [r.get("short") or r.get("desc", "") for r in rows]
    n = len(labels)
    chart_top = CHART_TOP_DUAL
    chart_h = max(MIN_CHART_HEIGHT, min(MAX_DUAL_CHART_HEIGHT, n * 0.42))
    row_h = chart_h / max(n, 1)

    # Left data
    left_current = [r.get(f"{left_prefix}_current") or 0 for r in rows]
    left_prior = [r.get(f"{left_prefix}_prior") or 0 for r in rows]

    # Right data
    right_current = [r.get(f"{right_prefix}_current") or 0 for r in rows]
    right_prior = [r.get(f"{right_prefix}_prior") or 0 for r in rows]

    # Left chart label
    left_label = left_cfg.get("label", "Left (%)")
    textbox(slide, left_label, 0.30, 1.72, 3.0, 0.25,
            fsize=9, bold=True, color=C_GREY, font=config.font_display)

    # Left clustered bar (Q4 vs Q3)
    cd1 = CategoryChartData()
    cd1.categories = labels
    cd1.add_series(config.period_current, left_current)
    cd1.add_series(config.period_prior, left_prior)

    cf1 = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(0.30), Inches(chart_top), Inches(4.8), Inches(chart_h), cd1)
    ch1 = cf1.chart
    ch1.has_legend = False
    set_series_color(ch1.series[0], color_current)
    set_series_no_border(ch1.series[0])
    set_series_color(ch1.series[1], color_prior)
    set_series_no_border(ch1.series[1])
    enable_data_labels(ch1.series[0], color_current, fsize=7, font_name=font)
    enable_data_labels(ch1.series[1], color_prior, fsize=7, font_name=font)
    hide_axis(ch1, "val")
    ch1.category_axis.tick_labels.font.size = Pt(7)
    ch1.category_axis.tick_labels.font.name = font
    suppress_cat_axis_bullets(ch1)
    invert_cat_axis(ch1)
    set_plot_area_gap(ch1, BAR_GAP_STD)
    set_overlap(ch1, CLUSTERED_OVERLAP)

    # Left delta
    left_deltas = [delta(c, p) for c, p in zip(left_current, left_prior)]
    left_delta_header = left_cfg.get("delta_header",
                                     left_prefix[0].upper() + " Δ")
    add_delta_table(
        slide, left_deltas,
        left=5.15, top=chart_top, width=DELTA_COL_NARROW, row_height=row_h * ROW_SCALE_FACTOR,
        header_text=left_delta_header, font_name=font,
    )

    # Right chart label
    right_label = right_cfg.get("label", "Right (%)")
    textbox(slide, right_label, 5.90, 1.72, 3.5, 0.25,
            fsize=9, bold=True, color=C_GREY, font=config.font_display)

    # Right clustered bar (Q4 vs Q3)
    cd2 = CategoryChartData()
    cd2.categories = labels
    cd2.add_series(config.period_current, right_current)
    cd2.add_series(config.period_prior, right_prior)

    cf2 = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(5.90), Inches(chart_top), Inches(4.8), Inches(chart_h), cd2)
    ch2 = cf2.chart
    ch2.has_legend = False
    set_series_color(ch2.series[0], color_current)
    set_series_no_border(ch2.series[0])
    set_series_color(ch2.series[1], color_prior)
    set_series_no_border(ch2.series[1])
    enable_data_labels(ch2.series[0], color_current, fsize=7, font_name=font)
    enable_data_labels(ch2.series[1], color_prior, fsize=7, font_name=font)
    hide_axis(ch2, "val")
    hide_cat_labels(ch2)
    invert_cat_axis(ch2)
    set_plot_area_gap(ch2, BAR_GAP_STD)
    set_overlap(ch2, CLUSTERED_OVERLAP)

    # Right delta
    right_deltas = [delta(c, p) for c, p in zip(right_current, right_prior)]
    right_delta_header = right_cfg.get("delta_header",
                                       right_prefix.upper() + " Δ")
    add_delta_table(
        slide, right_deltas,
        left=10.75, top=chart_top, width=DELTA_COL_NARROW, row_height=row_h * ROW_SCALE_FACTOR,
        header_text=right_delta_header, font_name=font,
    )

    # Legend
    ly = chart_top + chart_h + LEGEND_GAP
    _make_legend(slide, [
        (color_current, config.period_current),
        (color_prior, config.period_prior),
    ], 0, ly, font, center_over=(0, SLIDE_W))


