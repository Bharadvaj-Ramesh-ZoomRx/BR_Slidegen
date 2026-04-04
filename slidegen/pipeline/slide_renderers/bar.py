"""
bar.py — Single-bar, QoQ bar, and two-section bar slide renderers.
"""

from __future__ import annotations

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

from ._shared import (
    # constants
    CHART_TOP_STD, LEGEND_GAP, SLIDE_W, CHART_DELTA_GAP,
    SINGLE_BAR_WIDTH, DELTA_COL_WIDTH, DELTA_COL_NARROW,
    MAX_CHART_HEIGHT, MAX_QOQ_HEIGHT, MIN_CHART_HEIGHT,
    ROW_SCALE_FACTOR,
    BAR_GAP_STD, CLUSTERED_OVERLAP,
    LABEL_MAX_SINGLE, LABEL_MAX_DUAL,
    ELEMENT_GAP, SIDE_MARGIN, FOOTER_BUFFER, ROW_H_MIN, ROW_H_MAX, HDR_H_STD,
    # helpers
    _resolve_template, _slide_chrome, _get_brand_colors, _sort_data, _make_legend, _compute_row_h,
    _cap_chart_h, _auto_label_width, _vcenter_top, _prepare_rows,
    _alt_row_bg, _no_data_placeholder, _layout_blocks, _build_label_table,
    FONT_HDR, FONT_BODY,
    # pptx_utils
    C_GREEN, C_GREY, C_LBGREY, C_HDRGREY, C_WHITE, C_RED,
    FOOTER_TOP, PP_ALIGN,
    textbox,
    hide_axis, hide_cat_labels, set_series_color, set_plot_area_gap, set_overlap,
    set_series_no_border, invert_cat_axis, suppress_cat_axis_bullets,
    set_chart_plot_area,
    enable_data_labels,
    add_single_bar_chart, add_delta_table,
    _pptx_table, _style_tbl_cell,
    # data_loaders
    delta,
    # project_config types
    ProjectConfig, AskConfig,
)


def render_single_bar_with_delta(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Single horizontal bar chart with label table and QoQ delta column."""
    _slide_chrome(slide, config, ask)

    rows = _prepare_rows(slide, data, ask)
    if rows is None:
        return

    color_current, color_prior = _get_brand_colors(config, ask)
    font = config.font_body

    labels = [r.get("short") or r.get("desc", "") for r in rows]
    current_vals = [r.get("current") or 0 for r in rows]
    prior_vals = [r.get("prior") for r in rows]

    n = len(labels)

    # Layout: [Label table] [Bar chart (no cat labels)] [Delta col]
    delta_w = DELTA_COL_WIDTH
    label_w = _auto_label_width(labels)
    chart_w = SLIDE_W - SIDE_MARGIN * 2 - label_w - delta_w - ELEMENT_GAP * 2
    hdr_h = HDR_H_STD
    max_body_h = FOOTER_TOP - CHART_TOP_STD - FOOTER_BUFFER
    row_h = _compute_row_h(n, max_body_h, min_h=ROW_H_MIN, max_h=ROW_H_MAX)
    body_h = n * row_h
    content_h = hdr_h + body_h
    chart_top = _vcenter_top(content_h)

    # Center horizontally
    block_w = label_w + ELEMENT_GAP + chart_w + ELEMENT_GAP + delta_w
    _, (label_l, chart_l, delta_l) = _layout_blocks(SLIDE_W, label_w, chart_w, delta_w)

    # 1. Label table
    _build_label_table(slide, labels, label_w, row_h, label_l, chart_top, font=font)

    # 2. Bar chart (no category labels)
    cf, ch = add_single_bar_chart(
        slide, [f"R{i}" for i in range(n)], current_vals,
        left=chart_l, top=chart_top + hdr_h, width=chart_w, height=body_h,
        fill_color=color_current, font_name=font,
    )
    hide_cat_labels(ch)
    set_chart_plot_area(ch, x=0.0, y=0.0, w=1.0, h=1.0)

    # 3. Delta table
    deltas = [delta(c, p) if p is not None else None
              for c, p in zip(current_vals, prior_vals)]
    add_delta_table(
        slide, deltas,
        left=delta_l, top=chart_top, width=delta_w,
        row_height=row_h,
        header_text="QoQ Δ", font_name=font,
    )

    # 4. Legend
    ly = chart_top + hdr_h + body_h + LEGEND_GAP
    sample = config.sample_sizes.get(ask.brand or "primary")
    n_label = f" (n={sample.current})" if sample else ""
    _make_legend(slide, [
        (color_current, f"{config.period_current}{n_label}"),
        (C_GREEN, "Positive Δ"),
        (C_RED, "Negative Δ"),
    ], 0, ly, font, center_over=(label_l, block_w))


def render_qoq_bar_with_delta(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Clustered Q4 vs Q3 bar chart with delta column."""
    _slide_chrome(slide, config, ask)

    rows = _prepare_rows(slide, data, ask)
    if rows is None:
        return

    color_current, color_prior = _get_brand_colors(config, ask)
    font = config.font_body

    labels = [r.get("desc", "")[:LABEL_MAX_DUAL] for r in rows]
    current_vals = [r.get("current") or 0 for r in rows]
    prior_vals = [r.get("prior") or 0 for r in rows]

    n = len(labels)
    chart_h = max(MIN_CHART_HEIGHT, min(MAX_QOQ_HEIGHT, n * 0.50))
    row_h = chart_h / max(n, 1)
    chart_top = _vcenter_top(chart_h)

    # Center the chart + delta block horizontally
    qoq_chart_w = 8.0
    block_w = qoq_chart_w + CHART_DELTA_GAP + DELTA_COL_WIDTH
    _, (chart_left, delta_left) = _layout_blocks(SLIDE_W, qoq_chart_w, DELTA_COL_WIDTH, gap=CHART_DELTA_GAP)

    # Clustered bar
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series(config.period_current, current_vals)
    cd.add_series(config.period_prior, prior_vals)

    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(chart_left), Inches(chart_top), Inches(qoq_chart_w), Inches(chart_h), cd)
    ch = cf.chart
    ch.has_legend = False

    set_series_color(ch.series[0], color_current)
    set_series_no_border(ch.series[0])
    enable_data_labels(ch.series[0], color_current, font_name=font)

    set_series_color(ch.series[1], color_prior)
    set_series_no_border(ch.series[1])
    enable_data_labels(ch.series[1], color_prior, font_name=font)

    hide_axis(ch, "val")
    ch.category_axis.tick_labels.font.size = Pt(8)
    ch.category_axis.tick_labels.font.name = font
    suppress_cat_axis_bullets(ch)
    invert_cat_axis(ch)
    set_plot_area_gap(ch, BAR_GAP_STD)
    set_overlap(ch, CLUSTERED_OVERLAP)
    set_chart_plot_area(ch, x=0.0, y=0.0, w=1.0, h=1.0)

    # Delta
    deltas = [delta(c, p) for c, p in zip(current_vals, prior_vals)]
    add_delta_table(
        slide, deltas,
        left=delta_left, top=chart_top, width=DELTA_COL_WIDTH, row_height=row_h * ROW_SCALE_FACTOR,
        header_text="QoQ Δ", font_name=font,
    )

    # Manual legend below chart
    _make_legend(slide, [
        (color_current, config.period_current),
        (color_prior, config.period_prior),
    ], 0, chart_top + chart_h + LEGEND_GAP, font,
        center_over=(chart_left, block_w))


def render_two_section_bar(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Two vertically stacked single-bar sections (e.g. RYB Rx + TAG Rx)."""
    _slide_chrome(slide, config, ask)

    extra = ask.extra
    font = config.font_body

    # Estimate total content height for vertical centering
    _total_h = 0
    for _, sc in [("top", extra.get("top", {})), ("bottom", extra.get("bottom", {}))]:
        sr = data.get(sc.get("data_key", ""), [])
        _total_h += (max(1.5, len(sr) * 0.55) if sr else 1.7) + 0.50 + 0.15
    chart_top = _vcenter_top(_total_h, has_legend=False)

    sections = [("top", extra.get("top", {})), ("bottom", extra.get("bottom", {}))]
    y_offset = chart_top

    for pos, sec_cfg in sections:
        sec_data_key = sec_cfg.get("data_key", "")
        sec_rows = data.get(sec_data_key, [])
        sec_label = _resolve_template(sec_cfg.get("label", ""), config)
        sec_brand = sec_cfg.get("brand", "primary")
        brand = config.brands.get(sec_brand, config.primary)

        # Center the chart + delta block
        _block_w = SINGLE_BAR_WIDTH + CHART_DELTA_GAP + DELTA_COL_NARROW
        _chart_left = (SLIDE_W - _block_w) / 2
        _delta_left = _chart_left + SINGLE_BAR_WIDTH + CHART_DELTA_GAP

        # Section label
        textbox(slide, sec_label, _chart_left, y_offset - 0.15, SINGLE_BAR_WIDTH, 0.25,
                fsize=9, bold=True, color=brand.color_current, font=config.font_display)

        if sec_rows:
            labels = [r.get("desc", "")[:LABEL_MAX_DUAL] for r in sec_rows]
            current_vals = [r.get("current") or 0 for r in sec_rows]
            prior_vals = [r.get("prior") for r in sec_rows]
            n = len(labels)
            ch_h = max(1.5, n * 0.55)

            cd = CategoryChartData()
            cd.categories = labels
            cd.add_series(config.period_current, current_vals)

            cf = slide.shapes.add_chart(
                XL_CHART_TYPE.BAR_CLUSTERED,
                Inches(_chart_left), Inches(y_offset), Inches(SINGLE_BAR_WIDTH), Inches(ch_h), cd)
            ch = cf.chart
            ch.has_legend = False
            s = ch.series[0]
            set_series_color(s, brand.color_current)
            set_series_no_border(s)
            enable_data_labels(s, brand.color_current, font_name=font)
            hide_axis(ch, "val")
            ch.category_axis.tick_labels.font.size = Pt(8)
            ch.category_axis.tick_labels.font.name = font
            suppress_cat_axis_bullets(ch)
            invert_cat_axis(ch)
            set_plot_area_gap(ch, BAR_GAP_STD)

            # Delta
            d = [delta(c, p) if p is not None else None
                 for c, p in zip(current_vals, prior_vals)]
            add_delta_table(
                slide, d,
                left=_delta_left, top=y_offset, width=DELTA_COL_NARROW,
                row_height=ch_h / max(n, 1) * ROW_SCALE_FACTOR,
                header_text="Δ", font_name=font,
            )

            y_offset += ch_h + 0.50
        else:
            y_offset += 1.7

