"""
slide_renderers.py — Reusable slide type renderers.

Each renderer takes (slide, config, ask, data, *, namer=None) and builds
the slide content generically from the ask definition + extracted data.

Slide types:
  - cover
  - executive_summary
  - single_bar_with_delta
  - dual_bar_with_delta
  - dual_bar_qoq
  - clustered_compare
  - qoq_bar_with_delta
  - two_section_bar
  - stacked_order
"""

from __future__ import annotations
import logging

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from lxml import etree
from pptx.oxml.ns import qn

from slidegen.pptx_utils import (
    C_GREEN, C_WHITE, C_GREY, C_FTGREY, C_LBGREY, C_HDRGREY, C_RED,
    textbox, solidrect, slide_header, slide_footer,
    section_header_bar, section_breadcrumb, chart_header_row,
    callout_box, dashed_separator, hide_axis, set_series_color,
    set_plot_area_gap, set_overlap, set_series_no_border,
    invert_cat_axis, hide_cat_labels, set_data_label_color,
    set_chart_plot_area, set_val_axis_scale,
    _get_or_add,
    enable_data_labels, delete_data_label,
    add_single_bar_chart, add_clustered_bar_chart,
    add_delta_table, add_value_table,
    cover_slide,
)
from slidegen.pipeline.data_loaders import delta
from slidegen.pipeline.project_config import ProjectConfig, AskConfig, parse_color

logger = logging.getLogger(__name__)


# ── Layout constants ─────────────────────────────────────────────────────────
# Y-axis positions (inches)
SECTION_BAR_TOP = 1.40
CHART_TOP_STD = 1.85       # standard chart top below section bar
LEGEND_GAP = 0.15           # gap below chart before legend
FOOTER_TOP = 6.93           # footer Y position (spec)

# Chart dimensions (inches) — generic
SINGLE_BAR_WIDTH = 9.0
CLUSTERED_BAR_WIDTH = 8.5
DELTA_COL_WIDTH = 0.65
DELTA_COL_NARROW = 0.55

# Dual bar (Archetype 3 — MR+ME) positions from spec
DUAL_HEADER_ROW_TOP = 1.40   # red header row Y (just below section bar)
DUAL_HEADER_ROW_H = 0.30     # red header row height
DUAL_CHART_TOP = 1.68        # chart top (below header row)
DUAL_MR_LEFT = 0.20          # left chart X
DUAL_MR_WIDTH = 7.30         # left chart width (wider — has cat labels)
DUAL_MR_DELTA_LEFT = 7.54    # left delta column X
DUAL_ME_LEFT = 8.26          # right chart X
DUAL_ME_WIDTH = 4.40         # right chart width (narrower — no cat labels)
DUAL_ME_DELTA_LEFT = 12.70   # right delta column X
DUAL_DELTA_WIDTH = 0.62      # delta column width
DUAL_SEPARATOR_X = 8.16      # vertical dashed separator X
DUAL_MAX_CHART_HEIGHT = 5.05  # max chart height (fills to footer area)

# Chart sizing constraints
MAX_CHART_HEIGHT = 4.5
MAX_CLUSTERED_HEIGHT = 4.8
MAX_QOQ_HEIGHT = 4.0
MIN_CHART_HEIGHT = 3.0       # minimum chart height to avoid excessive white space
ROW_SCALE_FACTOR = 0.85     # delta table row height as fraction of chart row

# Compat aliases for renderers not yet updated to Archetype 3 layout
CHART_TOP_DUAL = 1.80                # used by render_dual_bar_qoq
DUAL_BAR_WIDTH = 5.0                 # used by render_dual_bar_qoq
MAX_DUAL_CHART_HEIGHT = 4.2          # used by render_dual_bar_qoq

# Bar chart gap/overlap
BAR_GAP_STD = 80
CLUSTERED_OVERLAP = -10
CLUSTERED_GAP = 65

# Label truncation lengths
LABEL_MAX_SINGLE = 40
LABEL_MAX_DUAL = 35
LABEL_MAX_CLUSTERED = 45
LABEL_MAX_STACKED = 40

# Small segment threshold for stacked bars (%)
STACKED_HIDE_THRESHOLD = 3


# ── Template resolution ───────────────────────────────────────────────────────

def _resolve_template(text: str, config: ProjectConfig) -> str:
    """Replace {{...}} placeholders in headline/section text."""
    if not text:
        return text
    text = text.replace("{{primary.name}}", config.primary.name)
    text = text.replace("{{primary.full_name}}", config.primary.full_name)
    text = text.replace("{{competitor.name}}", config.competitor.name)
    text = text.replace("{{competitor.full_name}}", config.competitor.full_name)
    text = text.replace("{{period_current}}", config.period_current)
    text = text.replace("{{period_prior}}", config.period_prior)
    text = text.replace("{{client}}", config.client)
    return text


# ── Shared layout helpers ─────────────────────────────────────────────────────

def _slide_chrome(slide, config: ProjectConfig, ask: AskConfig):
    """Add header, section bar, breadcrumb, and footer to a slide."""
    headline = _resolve_template(ask.headline, config)
    slide_header(slide, headline, font=config.font_display)

    section = _resolve_template(ask.section, config)
    if section:
        section_header_bar(slide, section, top=SECTION_BAR_TOP,
                           icon_path=config.section_icon_path or None,
                           font=config.font_body)

    # Breadcrumb (top-right, e.g. "Appendix" or "Key Findings")
    breadcrumb = ask.extra.get("breadcrumb", "") if ask.extra else ""
    if breadcrumb:
        section_breadcrumb(slide, _resolve_template(breadcrumb, config))

    source = _resolve_template(ask.source_text, config)
    if source:
        slide_footer(slide, source, font=config.font_body)


def _get_brand_colors(config: ProjectConfig, ask: AskConfig):
    """Return (color_current, color_prior) for the ask's brand."""
    brand_key = ask.brand if ask.brand else "primary"
    brand = config.brands.get(brand_key, config.primary)
    return brand.color_current, brand.color_prior


def _sort_data(rows: list[dict], sort_by: str | None, sort_desc: bool = True) -> list[dict]:
    """Sort rows by a field, handling None values."""
    if not sort_by:
        return rows
    return sorted(rows, key=lambda x: x.get(sort_by) or 0, reverse=sort_desc)


def _make_legend(slide, items: list[tuple], left: float, top: float, font_name: str = None):
    """Add a manual legend row. items: [(color, label), ...]"""
    x = left
    for color, label in items:
        solidrect(slide, x, top, 0.18, 0.12, color)
        textbox(slide, label, x + 0.25, top - 0.02, 1.8, 0.18,
                fsize=7, color=C_GREY, font=font_name)
        x += 2.0


# ══════════════════════════════════════════════════════════════════════════════
# RENDERERS
# ══════════════════════════════════════════════════════════════════════════════

def render_cover(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Title/cover slide."""
    headline = _resolve_template(ask.headline, config)
    extra = ask.extra
    cover_slide(
        slide,
        headline,
        extra.get("subtitle", ""),
        extra.get("date", ""),
        _resolve_template(extra.get("client", ""), config),
    )


def render_executive_summary(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Bullet-list executive summary. Insights are in ask.extra['insights']."""
    headline = _resolve_template(ask.headline, config)
    slide_header(slide, headline, font=config.font_display)

    color_current = config.primary.color_current
    insights = ask.extra.get("insights", [])

    y = 1.50
    for ins in insights:
        solidrect(slide, 0.40, y + 0.06, 0.08, 0.08, color_current)
        textbox(slide, ins, 0.60, y, 12.40, 0.50,
                fsize=9, color=C_GREY, font=config.font_body)
        y += 0.62

    source = _resolve_template(ask.source_text, config)
    if source:
        slide_footer(slide, source, font=config.font_body)


def render_single_bar_with_delta(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Single horizontal bar chart with a QoQ delta column."""
    _slide_chrome(slide, config, ask)

    rows = data.get(ask.data_key, [])
    rows = _sort_data(rows, ask.sort_by, ask.sort_desc)
    if not rows:
        textbox(slide, "Data not available", 2, 3, 8, 1, fsize=14, color=C_RED)
        return

    color_current, color_prior = _get_brand_colors(config, ask)
    font = config.font_body

    labels = [r.get("short", r.get("desc", ""))[:LABEL_MAX_SINGLE] for r in rows]
    current_vals = [r.get("current") or 0 for r in rows]
    prior_vals = [r.get("prior") for r in rows]

    n = len(labels)
    chart_top = CHART_TOP_STD
    chart_h = max(MIN_CHART_HEIGHT, min(MAX_CHART_HEIGHT, n * 0.40))
    row_h = chart_h / max(n, 1)

    # Bar chart
    add_single_bar_chart(
        slide, labels, current_vals,
        left=0.30, top=chart_top, width=SINGLE_BAR_WIDTH, height=chart_h,
        fill_color=color_current, font_name=font,
    )

    # Delta table
    deltas = [delta(c, p) if p is not None else None
              for c, p in zip(current_vals, prior_vals)]
    add_delta_table(
        slide, deltas,
        left=9.40, top=chart_top, width=DELTA_COL_WIDTH,
        row_height=row_h * ROW_SCALE_FACTOR,
        header_text="QoQ Δ", font_name=font,
    )

    # Legend
    ly = chart_top + chart_h + LEGEND_GAP
    sample = config.sample_sizes.get(ask.brand or "primary")
    n_label = f" (n={sample.current})" if sample else ""
    _make_legend(slide, [
        (color_current, f"{config.period_current}{n_label}"),
        (C_GREEN, "Positive Δ"),
        (C_RED, "Negative Δ"),
    ], 2.0, ly, font)


def render_dual_bar_with_delta(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Two side-by-side bar charts (e.g. MR + ME) each with delta columns.

    Uses Archetype 3 layout: MR wider (7.30") with cat labels,
    ME narrower (4.40") without cat labels, red header row above charts.
    """
    _slide_chrome(slide, config, ask)

    rows = data.get(ask.data_key, [])
    rows = _sort_data(rows, ask.sort_by, ask.sort_desc)
    if not rows:
        textbox(slide, "Data not available", 2, 3, 8, 1, fsize=14, color=C_RED)
        return

    color_current, color_prior = _get_brand_colors(config, ask)
    font = config.font_body
    extra = ask.extra

    left_cfg = extra.get("left", {})
    right_cfg = extra.get("right", {})
    left_prefix = left_cfg.get("field_prefix", "mr")
    right_prefix = right_cfg.get("field_prefix", "me")

    labels = [r.get("short", r.get("desc", ""))[:LABEL_MAX_DUAL] for r in rows]
    n = len(labels)
    chart_top = DUAL_CHART_TOP
    chart_h = min(DUAL_MAX_CHART_HEIGHT, n * 0.42)
    row_h = chart_h / max(n, 1)

    # Left data
    left_current = [r.get(f"{left_prefix}_current") or 0 for r in rows]
    left_prior = [r.get(f"{left_prefix}_prior") for r in rows]

    # Right data
    right_current = [r.get(f"{right_prefix}_current") or 0 for r in rows]
    right_prior = [r.get(f"{right_prefix}_prior") for r in rows]

    # ── Red header row (config-driven column titles) ─────────────────────
    cat_header = extra.get("category_header", "")
    left_label = left_cfg.get("label", "Left (%)")
    right_label = right_cfg.get("label", "Right (%)")

    left_delta_header = left_cfg.get("delta_header", "Δ")
    right_delta_header = right_cfg.get("delta_header", "Δ")

    header_columns = [
        {"label": cat_header, "left": DUAL_MR_LEFT, "width": 2.0,
         "align": PP_ALIGN.LEFT},
        {"label": left_label,
         "left": DUAL_MR_LEFT + 2.0,
         "width": DUAL_MR_WIDTH + DUAL_DELTA_WIDTH - 2.0},
        {"label": right_label,
         "left": DUAL_ME_LEFT,
         "width": DUAL_ME_WIDTH + DUAL_DELTA_WIDTH},
    ]
    chart_header_row(slide, header_columns, top=DUAL_HEADER_ROW_TOP,
                     height=DUAL_HEADER_ROW_H, font=font)

    # ── Shared axis scale (so same % = same visual bar width) ────────────
    all_vals = left_current + right_current
    axis_max = max(all_vals) if all_vals else 100
    # Round up to next 10 for clean axis bounds
    axis_max = min(100, ((int(axis_max) // 10) + 1) * 10)

    # Delta row height: total delta table height = chart height
    # (header roughly matches chart's auto top-padding)
    from slidegen.pptx_utils.tables import HEADER_ROW_HEIGHT_IN
    delta_row_h = (chart_h - HEADER_ROW_HEIGHT_IN) / max(n, 1)

    # ── Left chart (MR — wider, with category labels) ────────────────────
    cf1, ch1 = add_single_bar_chart(
        slide, labels, left_current,
        left=DUAL_MR_LEFT, top=chart_top, width=DUAL_MR_WIDTH, height=chart_h,
        fill_color=color_current, font_name=font,
    )
    ch1.has_title = False
    set_val_axis_scale(ch1, 0, axis_max)

    # Left delta column
    left_deltas = [delta(c, p) if p is not None else None
                   for c, p in zip(left_current, left_prior)]
    add_delta_table(
        slide, left_deltas,
        left=DUAL_MR_DELTA_LEFT, top=chart_top, width=DUAL_DELTA_WIDTH,
        row_height=delta_row_h,
        header_text=left_delta_header, font_name=font,
    )

    # ── Vertical dashed separator ────────────────────────────────────────
    sep_h = chart_h + DUAL_HEADER_ROW_H
    dashed_separator(slide, DUAL_SEPARATOR_X, DUAL_HEADER_ROW_TOP,
                     sep_h, color=C_FTGREY, width_pt=0.5, vertical=True)

    # ── Right chart (ME — narrower, no category labels) ──────────────────
    cd2 = CategoryChartData()
    cd2.categories = labels
    cd2.add_series(config.period_current, right_current)

    cf2 = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(DUAL_ME_LEFT), Inches(chart_top),
        Inches(DUAL_ME_WIDTH), Inches(chart_h), cd2)
    ch2 = cf2.chart
    ch2.has_legend = False
    ch2.has_title = False
    s2 = ch2.series[0]
    set_series_color(s2, color_current)
    set_series_no_border(s2)
    enable_data_labels(s2, color_current, font_name=font)
    hide_axis(ch2, "val")
    hide_cat_labels(ch2)
    invert_cat_axis(ch2)
    set_plot_area_gap(ch2, BAR_GAP_STD)
    set_val_axis_scale(ch2, 0, axis_max)

    # Right delta column
    right_deltas = [delta(c, p) if p is not None else None
                    for c, p in zip(right_current, right_prior)]
    add_delta_table(
        slide, right_deltas,
        left=DUAL_ME_DELTA_LEFT, top=chart_top, width=DUAL_DELTA_WIDTH,
        row_height=delta_row_h,
        header_text=right_delta_header, font_name=font,
    )

    # ── Legend ────────────────────────────────────────────────────────────
    ly = chart_top + chart_h + LEGEND_GAP
    sample = config.sample_sizes.get(ask.brand or "primary")
    n_label = f" (n={sample.current})" if sample else ""
    _make_legend(slide, [
        (color_current, f"{config.period_current}{n_label}"),
        (C_GREEN, "Positive Δ"),
        (C_RED, "Negative Δ"),
    ], 2.0, ly, font)


def render_dual_bar_qoq(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Two side-by-side clustered bar charts (Q4 vs Q3) each with delta columns."""
    _slide_chrome(slide, config, ask)

    rows = data.get(ask.data_key, [])
    rows = _sort_data(rows, ask.sort_by, ask.sort_desc)
    if not rows:
        textbox(slide, "Data not available", 2, 3, 8, 1, fsize=14, color=C_RED)
        return

    color_current, color_prior = _get_brand_colors(config, ask)
    font = config.font_body
    extra = ask.extra

    left_cfg = extra.get("left", {})
    right_cfg = extra.get("right", {})
    left_prefix = left_cfg.get("field_prefix", "believ")
    right_prefix = right_cfg.get("field_prefix", "me")

    labels = [r.get("short", r.get("desc", ""))[:LABEL_MAX_DUAL] for r in rows]
    n = len(labels)
    chart_top = CHART_TOP_DUAL
    chart_h = max(MIN_CHART_HEIGHT, min(MAX_DUAL_CHART_HEIGHT, n * 0.42))
    row_h = chart_h / max(n, 1)

    # Left data
    left_current = [r.get(f"{left_prefix}_current") or 0 for r in rows]
    left_prior = [r.get(f"{left_prefix}_prior") or 0 for r in rows]

    # Right data
    right_current = [r.get(f"{right_prefix}_current") or 0 for r in rows]
    right_prior = [r.get(f"{right_prefix}_prior") or 0 for r in rows]

    # Left chart label
    left_label = left_cfg.get("label", "Left (%)")
    textbox(slide, left_label, 0.30, 1.72, 3.0, 0.25,
            fsize=9, bold=True, color=C_GREY, font=config.font_display)

    # Left clustered bar (Q4 vs Q3)
    cd1 = CategoryChartData()
    cd1.categories = labels
    cd1.add_series(config.period_current, left_current)
    cd1.add_series(config.period_prior, left_prior)

    cf1 = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(0.30), Inches(chart_top), Inches(4.8), Inches(chart_h), cd1)
    ch1 = cf1.chart
    ch1.has_legend = False
    set_series_color(ch1.series[0], color_current)
    set_series_no_border(ch1.series[0])
    set_series_color(ch1.series[1], color_prior)
    set_series_no_border(ch1.series[1])
    enable_data_labels(ch1.series[0], color_current, fsize=7, font_name=font)
    enable_data_labels(ch1.series[1], color_prior, fsize=7, font_name=font)
    hide_axis(ch1, "val")
    ch1.category_axis.tick_labels.font.size = Pt(7)
    ch1.category_axis.tick_labels.font.name = font
    invert_cat_axis(ch1)
    set_plot_area_gap(ch1, BAR_GAP_STD)
    set_overlap(ch1, CLUSTERED_OVERLAP)

    # Left delta
    left_deltas = [delta(c, p) for c, p in zip(left_current, left_prior)]
    left_delta_header = left_cfg.get("delta_header",
                                     left_prefix[0].upper() + " Δ")
    add_delta_table(
        slide, left_deltas,
        left=5.15, top=chart_top, width=DELTA_COL_NARROW, row_height=row_h * ROW_SCALE_FACTOR,
        header_text=left_delta_header, font_name=font,
    )

    # Right chart label
    right_label = right_cfg.get("label", "Right (%)")
    textbox(slide, right_label, 5.90, 1.72, 3.5, 0.25,
            fsize=9, bold=True, color=C_GREY, font=config.font_display)

    # Right clustered bar (Q4 vs Q3)
    cd2 = CategoryChartData()
    cd2.categories = labels
    cd2.add_series(config.period_current, right_current)
    cd2.add_series(config.period_prior, right_prior)

    cf2 = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(5.90), Inches(chart_top), Inches(4.8), Inches(chart_h), cd2)
    ch2 = cf2.chart
    ch2.has_legend = False
    set_series_color(ch2.series[0], color_current)
    set_series_no_border(ch2.series[0])
    set_series_color(ch2.series[1], color_prior)
    set_series_no_border(ch2.series[1])
    enable_data_labels(ch2.series[0], color_current, fsize=7, font_name=font)
    enable_data_labels(ch2.series[1], color_prior, fsize=7, font_name=font)
    hide_axis(ch2, "val")
    hide_cat_labels(ch2)
    invert_cat_axis(ch2)
    set_plot_area_gap(ch2, BAR_GAP_STD)
    set_overlap(ch2, CLUSTERED_OVERLAP)

    # Right delta
    right_deltas = [delta(c, p) for c, p in zip(right_current, right_prior)]
    right_delta_header = right_cfg.get("delta_header",
                                       right_prefix.upper() + " Δ")
    add_delta_table(
        slide, right_deltas,
        left=10.75, top=chart_top, width=DELTA_COL_NARROW, row_height=row_h * ROW_SCALE_FACTOR,
        header_text=right_delta_header, font_name=font,
    )

    # Legend
    ly = chart_top + chart_h + LEGEND_GAP
    _make_legend(slide, [
        (color_current, config.period_current),
        (color_prior, config.period_prior),
    ], 2.5, ly, font)


def render_clustered_compare(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Clustered horizontal bar comparing two groups + delta columns."""
    _slide_chrome(slide, config, ask)

    extra = ask.extra
    series_cfgs = extra.get("series", [])

    # Support combined data keys (CTA merges ryb_cta + tag_cta)
    rows = data.get(ask.data_key, [])

    # If data_key not found, check for primary_key + comp_key pattern
    if not rows and "primary_key" in extra:
        logger.warning(
            "clustered_compare '%s': data_key '%s' not found, merging "
            "primary_key '%s' + comp_key '%s' instead",
            ask.id, ask.data_key, extra.get("primary_key"), extra.get("comp_key"),
        )
        primary_rows = data.get(extra["primary_key"], [])
        comp_rows = data.get(extra["comp_key"], [])
        # Merge: align by desc/label
        merged = []
        for pr in primary_rows:
            label = pr.get("desc", "")
            cr = next((c for c in comp_rows if c.get("desc") == label), {})
            merged.append({
                "desc": label,
                "primary_current": pr.get("current"),
                "primary_prior": pr.get("prior"),
                "comp_current": cr.get("current"),
                "comp_prior": cr.get("prior"),
            })
        rows = merged

    rows = _sort_data(rows, ask.sort_by, ask.sort_desc)
    if not rows:
        textbox(slide, "Data not available", 2, 3, 8, 1, fsize=14, color=C_RED)
        return

    font = config.font_body

    # Determine series fields and colors
    if len(series_cfgs) >= 2:
        s1_field = series_cfgs[0].get("field", "primary")
        s2_field = series_cfgs[1].get("field", "comp")
        s1_label = _resolve_template(series_cfgs[0].get("label", "Series 1"), config)
        s2_label = _resolve_template(series_cfgs[1].get("label", "Series 2"), config)
        s1_color = parse_color(series_cfgs[0]["color"]) if "color" in series_cfgs[0] else config.primary.color_current
        s2_color = parse_color(series_cfgs[1]["color"]) if "color" in series_cfgs[1] else config.competitor.color_current
    else:
        s1_field, s2_field = "primary", "comp"
        s1_label, s2_label = config.primary.name, config.competitor.name
        s1_color, s2_color = config.primary.color_current, config.competitor.color_current

    labels = [r.get("desc", "")[:LABEL_MAX_CLUSTERED] for r in rows]

    # Extract values — support both flat (hi_current, other_current) and
    # prefixed (primary_current, comp_current) naming
    s1_current = []
    s1_prior = []
    s2_current = []
    s2_prior = []
    for r in rows:
        s1_current.append(r.get(f"{s1_field}_current") or r.get(s1_field) or 0)
        s1_prior.append(r.get(f"{s1_field}_prior"))
        s2_current.append(r.get(f"{s2_field}_current") or r.get(s2_field) or 0)
        s2_prior.append(r.get(f"{s2_field}_prior"))

    n = len(labels)
    chart_top = CHART_TOP_STD
    chart_h = max(MIN_CHART_HEIGHT, min(MAX_CLUSTERED_HEIGHT, n * 0.32))
    row_h = chart_h / max(n, 1)

    # Clustered bar
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series(s1_label, s1_current)
    cd.add_series(s2_label, s2_current)

    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(0.30), Inches(chart_top), Inches(CLUSTERED_BAR_WIDTH), Inches(chart_h), cd)
    ch = cf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(8)
    ch.legend.font.name = font

    set_series_color(ch.series[0], s1_color)
    set_series_no_border(ch.series[0])
    enable_data_labels(ch.series[0], s1_color, fsize=7, font_name=font)

    set_series_color(ch.series[1], s2_color)
    set_series_no_border(ch.series[1])
    enable_data_labels(ch.series[1], s2_color, fsize=7, font_name=font)

    hide_axis(ch, "val")
    ch.category_axis.tick_labels.font.size = Pt(6.5)
    ch.category_axis.tick_labels.font.name = font
    invert_cat_axis(ch)
    set_plot_area_gap(ch, CLUSTERED_GAP)
    set_overlap(ch, -15)

    # Delta columns
    has_prior_1 = any(p is not None for p in s1_prior)
    has_prior_2 = any(p is not None for p in s2_prior)

    if has_prior_1 and has_prior_2:
        # Two delta columns (QoQ for each series)
        d1 = [delta(c, p) if p is not None else None for c, p in zip(s1_current, s1_prior)]
        d2 = [delta(c, p) if p is not None else None for c, p in zip(s2_current, s2_prior)]
        add_delta_table(slide, d1, left=9.10, top=chart_top, width=DELTA_COL_WIDTH,
                        row_height=row_h * ROW_SCALE_FACTOR,
                        header_text=s1_label.split("(")[0].strip() + " Δ",
                        font_name=font)
        add_delta_table(slide, d2, left=9.85, top=chart_top, width=DELTA_COL_WIDTH,
                        row_height=row_h * ROW_SCALE_FACTOR,
                        header_text=s2_label.split("(")[0].strip() + " Δ",
                        font_name=font)
    else:
        # Gap column (difference between series)
        gaps = [round((c1 or 0) - (c2 or 0), 1) for c1, c2 in zip(s1_current, s2_current)]
        header = extra.get("gap_header", "Gap (pp)")
        add_delta_table(slide, gaps, left=9.10, top=chart_top, width=DELTA_COL_WIDTH,
                        row_height=row_h * ROW_SCALE_FACTOR, header_text=header,
                        font_name=font)


def render_qoq_bar_with_delta(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Clustered Q4 vs Q3 bar chart with delta column."""
    _slide_chrome(slide, config, ask)

    rows = data.get(ask.data_key, [])
    rows = _sort_data(rows, ask.sort_by, ask.sort_desc)
    if not rows:
        textbox(slide, "Data not available", 2, 3, 8, 1, fsize=14, color=C_RED)
        return

    color_current, color_prior = _get_brand_colors(config, ask)
    font = config.font_body

    labels = [r.get("desc", "")[:LABEL_MAX_DUAL] for r in rows]
    current_vals = [r.get("current") or 0 for r in rows]
    prior_vals = [r.get("prior") or 0 for r in rows]

    n = len(labels)
    chart_top = CHART_TOP_STD
    chart_h = max(MIN_CHART_HEIGHT, min(MAX_QOQ_HEIGHT, n * 0.50))
    row_h = chart_h / max(n, 1)

    # Clustered bar
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series(config.period_current, current_vals)
    cd.add_series(config.period_prior, prior_vals)

    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(0.40), Inches(chart_top), Inches(8.0), Inches(chart_h), cd)
    ch = cf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.font.size = Pt(8)
    ch.legend.font.name = font

    set_series_color(ch.series[0], color_current)
    set_series_no_border(ch.series[0])
    enable_data_labels(ch.series[0], color_current, font_name=font)

    set_series_color(ch.series[1], color_prior)
    set_series_no_border(ch.series[1])
    enable_data_labels(ch.series[1], color_prior, font_name=font)

    hide_axis(ch, "val")
    ch.category_axis.tick_labels.font.size = Pt(8)
    ch.category_axis.tick_labels.font.name = font
    invert_cat_axis(ch)
    set_plot_area_gap(ch, BAR_GAP_STD)
    set_overlap(ch, CLUSTERED_OVERLAP)

    # Delta
    deltas = [delta(c, p) for c, p in zip(current_vals, prior_vals)]
    add_delta_table(
        slide, deltas,
        left=8.60, top=chart_top, width=DELTA_COL_WIDTH, row_height=row_h * ROW_SCALE_FACTOR,
        header_text="QoQ Δ", font_name=font,
    )


def render_two_section_bar(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Two vertically stacked single-bar sections (e.g. RYB Rx + TAG Rx)."""
    _slide_chrome(slide, config, ask)

    extra = ask.extra
    font = config.font_body
    chart_top = CHART_TOP_STD

    sections = [("top", extra.get("top", {})), ("bottom", extra.get("bottom", {}))]
    y_offset = chart_top

    for pos, sec_cfg in sections:
        sec_data_key = sec_cfg.get("data_key", "")
        sec_rows = data.get(sec_data_key, [])
        sec_label = _resolve_template(sec_cfg.get("label", ""), config)
        sec_brand = sec_cfg.get("brand", "primary")
        brand = config.brands.get(sec_brand, config.primary)

        # Section label
        textbox(slide, sec_label, 0.40, y_offset - 0.15, 5.0, 0.25,
                fsize=9, bold=True, color=brand.color_current, font=config.font_display)

        if sec_rows:
            labels = [r.get("desc", "")[:LABEL_MAX_DUAL] for r in sec_rows]
            current_vals = [r.get("current") or 0 for r in sec_rows]
            prior_vals = [r.get("prior") for r in sec_rows]
            n = len(labels)
            ch_h = max(1.5, n * 0.55)

            cd = CategoryChartData()
            cd.categories = labels
            cd.add_series(config.period_current, current_vals)

            cf = slide.shapes.add_chart(
                XL_CHART_TYPE.BAR_CLUSTERED,
                Inches(0.40), Inches(y_offset), Inches(5.5), Inches(ch_h), cd)
            ch = cf.chart
            ch.has_legend = False
            s = ch.series[0]
            set_series_color(s, brand.color_current)
            set_series_no_border(s)
            enable_data_labels(s, brand.color_current, font_name=font)
            hide_axis(ch, "val")
            ch.category_axis.tick_labels.font.size = Pt(8)
            ch.category_axis.tick_labels.font.name = font
            invert_cat_axis(ch)
            set_plot_area_gap(ch, BAR_GAP_STD)

            # Delta
            d = [delta(c, p) if p is not None else None
                 for c, p in zip(current_vals, prior_vals)]
            add_delta_table(
                slide, d,
                left=6.0, top=y_offset, width=DELTA_COL_NARROW,
                row_height=ch_h / max(n, 1) * ROW_SCALE_FACTOR,
                header_text="Δ", font_name=font,
            )

            y_offset += ch_h + 0.50
        else:
            y_offset += 1.7


def render_stacked_order(slide, config: ProjectConfig, ask: AskConfig, data: dict, *, namer=None):
    """Stacked bar with ordinal breakdown (1st/2nd/3rd/4th recalled)."""
    _slide_chrome(slide, config, ask)

    rows = data.get(ask.data_key, [])
    rows = _sort_data(rows, ask.sort_by, ask.sort_desc)
    if not rows:
        textbox(slide, "Data not available", 2, 3, 8, 1, fsize=14, color=C_RED)
        return

    color_current, _ = _get_brand_colors(config, ask)
    font = config.font_body

    # Take top 10
    order_top = rows[:min(10, len(rows))]

    labels = [r.get("desc", "")[:LABEL_MAX_STACKED] for r in order_top]

    # Determine ordinals from data keys
    ordinals = ask.extra.get("ordinals", ["1st", "2nd", "3rd", "4th"])
    ordinal_vals = []
    for o in ordinals:
        ordinal_vals.append([r.get(f"{o}_current", 0) for r in order_top])
    total_vals = [r.get("total_current", 0) for r in order_top]

    n = len(labels)
    chart_top = CHART_TOP_STD
    chart_h = max(MIN_CHART_HEIGHT, min(MAX_CHART_HEIGHT, n * 0.42))
    row_h = chart_h / max(n, 1)

    # Build stacked bar chart
    cd = CategoryChartData()
    cd.categories = labels
    ordinal_labels = [f"{o} Recalled" for o in ordinals]
    if len(ordinal_labels) == 4:
        ordinal_labels[3] = f"{ordinals[3]}+ Recalled"
    for label, vals in zip(ordinal_labels, ordinal_vals):
        cd.add_series(label, vals)

    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_STACKED,
        Inches(0.30), Inches(chart_top), Inches(CLUSTERED_BAR_WIDTH), Inches(chart_h), cd)
    ch = cf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.font.size = Pt(7.5)
    ch.legend.font.name = font

    # Color scheme: gradient from dark to light based on brand color
    r, g, b = color_current[0], color_current[1], color_current[2]
    stack_colors = [
        RGBColor(max(0, r - 60), max(0, g - 30), max(0, b - 50)),  # darkest
        color_current,                                                # medium
        RGBColor(min(255, r + 50), min(255, g + 60), min(255, b + 50)),  # light
        RGBColor(min(255, r + 100), min(255, g + 120), min(255, b + 100)),  # palest
    ]
    label_colors = [C_WHITE, C_WHITE,
                    RGBColor(0x30, 0x10, 0x50), RGBColor(0x30, 0x10, 0x50)]

    for idx, series in enumerate(ch.series):
        if idx < len(stack_colors):
            set_series_color(series, stack_colors[idx])
        set_series_no_border(series)
        enable_data_labels(series, label_colors[idx] if idx < len(label_colors) else C_GREY,
                          fsize=6, pos="ctr", font_name=font)
        # Hide labels on small segments
        if idx < len(ordinal_vals):
            for pt_idx, val in enumerate(ordinal_vals[idx]):
                if val <= STACKED_HIDE_THRESHOLD:
                    delete_data_label(series, pt_idx)

    hide_axis(ch, "val")
    ch.category_axis.tick_labels.font.size = Pt(7)
    ch.category_axis.tick_labels.font.name = font
    invert_cat_axis(ch)
    set_plot_area_gap(ch, BAR_GAP_STD)

    # Total column
    add_value_table(
        slide, total_vals,
        left=9.10, top=chart_top, width=0.70, row_height=row_h * ROW_SCALE_FACTOR,
        header_text="Total %", value_color=color_current, font_name=font,
    )

    # QoQ delta column
    total_prior = [r.get("total_prior", 0) for r in order_top]
    qoq_deltas = [delta(c, p) for c, p in zip(total_vals, total_prior)]
    add_delta_table(
        slide, qoq_deltas,
        left=9.90, top=chart_top, width=DELTA_COL_WIDTH, row_height=row_h * ROW_SCALE_FACTOR,
        header_text="QoQ Δ", font_name=font,
    )


# ══════════════════════════════════════════════════════════════════════════════
# REGISTRY — maps slide_type string → renderer function
# ══════════════════════════════════════════════════════════════════════════════

RENDERERS = {
    "cover": render_cover,
    "executive_summary": render_executive_summary,
    "single_bar_with_delta": render_single_bar_with_delta,
    "dual_bar_with_delta": render_dual_bar_with_delta,
    "dual_bar_qoq": render_dual_bar_qoq,
    "clustered_compare": render_clustered_compare,
    "qoq_bar_with_delta": render_qoq_bar_with_delta,
    "two_section_bar": render_two_section_bar,
    "stacked_order": render_stacked_order,
}
