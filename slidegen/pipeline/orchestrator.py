"""
orchestrator.py — Ties config + data + renderers into a complete deck.

Usage:
    from slidegen.pipeline import generate_deck
    generate_deck("projects/jnj_rybrevant/config.yaml")

    # Regenerate a single slide:
    from slidegen.pipeline.orchestrator import regenerate_slide
    regenerate_slide("projects/jnj_rybrevant/config.yaml", slide_index=4)
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import shutil
import tempfile
from datetime import datetime
from pptx import Presentation

from slidegen.pipeline.project_config import load_project_config, ProjectConfig, AskConfig, DataExtractionConfig
from slidegen.pipeline.data_loaders import load_all_data
from slidegen.pipeline.slide_renderers import RENDERERS
from slidegen.pipeline.slide_renderers._shared import render_qual_callout
from slidegen.pptx_utils.deck import (
    load_template as _load_template_impl,
    clear_slide as _clear_slide_impl,
    clear_sections as _clear_sections_impl,
    create_sections as _create_sections_impl,
)

logger = logging.getLogger(__name__)


# ── Shape naming ─────────────────────────────────────────────────────────────

class ShapeNamer:
    """Assigns sequential zrx_ names to shapes on a slide.

    Names follow the pattern zrx_{slide_idx:03d}_{shape_num:03d}, e.g.
    zrx_001_001, zrx_001_002, ... for slide 1.
    """

    def __init__(self, slide_idx: int, ask_id: str = "",
                 data_key: str = "", config_hash: str = "",
                 source_file: str = "", renderer: str = "",
                 last_data_pull: str = ""):
        self._slide_idx = slide_idx
        self._counter = 0
        self._registry: dict[str, dict] = {}
        self._ask_id = ask_id
        self._data_key = data_key
        self._config_hash = config_hash
        self._source_file = source_file
        self._renderer = renderer
        self._last_data_pull = last_data_pull

    def name(self, shape, label: str = "") -> str:
        """Assign a zrx_ name to a shape and record it in the registry."""
        self._counter += 1
        name = f"zrx_{self._slide_idx:03d}_{self._counter:03d}"
        shape.name = name
        self._registry[name] = {"label": label}
        return name

    def name_remaining(self, slide):
        """Name any shapes on the slide that don't yet have a zrx_ prefix."""
        for shape in slide.shapes:
            if not shape.name.startswith("zrx_"):
                self.name(shape, label="chrome")

    @property
    def registry(self) -> dict[str, dict]:
        return self._registry

    @property
    def slide_metadata(self) -> dict:
        """Return per-slide metadata including data_source lineage (PRD §6.2)."""
        now = datetime.now().isoformat()
        meta = {
            "ask_id": self._ask_id,
            "generated_at": now,
            "last_refreshed": now,
            "renderer": self._renderer,
            "shapes": self._registry,
        }
        if self._data_key or self._source_file:
            meta["data_source"] = {
                "data_key": self._data_key,
                "config_hash": self._config_hash,
                "source_file": self._source_file,
            }
        if self._last_data_pull:
            meta["last_data_pull"] = self._last_data_pull
        return meta


def _atomic_json_write(path: str, payload: dict):
    """Write JSON atomically via temp-file + os.replace to prevent corruption."""
    dir_path = os.path.dirname(path)
    os.makedirs(dir_path, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=dir_path, suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2)
        os.replace(tmp, path)
    except BaseException:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def _save_shape_registry(registries: dict, output_dir: str,
                         slide_to_ask: list[int] | None = None):
    """Save the combined shape registry for all slides."""
    path = os.path.join(output_dir, "shape_registry.json")
    payload = {
        "created": datetime.now().isoformat(),
        "slides": registries,
    }
    if slide_to_ask is not None:
        payload["slide_to_ask"] = slide_to_ask
    _atomic_json_write(path, payload)


def _safe_save_pptx(prs, path: str, suffix: str = "_regen") -> str:
    """Save PPTX with PermissionError fallback when file is open in PowerPoint.

    Returns the actual path where the file was saved.
    """
    os.makedirs(os.path.dirname(path), exist_ok=True)
    try:
        prs.save(path)
        # Clean up stale fallback files from prior PermissionError
        alt_path = path.replace(".pptx", f"{suffix}.pptx")
        if os.path.exists(alt_path):
            os.remove(alt_path)
        return path
    except PermissionError:
        alt_path = path.replace(".pptx", f"{suffix}.pptx")
        prs.save(alt_path)
        logger.warning(
            "%s is open in PowerPoint — saved to: %s\n"
            "  Close the original in PowerPoint, then replace it with the new file\n"
            "  (or reopen the new file directly)",
            os.path.basename(path), alt_path
        )
        return alt_path


# ── Speaker notes ─────────────────────────────────────────────────────────────

def _build_speaker_notes(ask: AskConfig, config: ProjectConfig, data: dict) -> str:
    """Build speaker notes text with question codes and descriptions for a slide.

    Collects all data_keys referenced by the ask (primary + any in extra),
    finds the matching extraction configs, and formats the question codes
    and question text used to create the slide.
    """
    # Collect all data keys this ask references
    data_keys = []
    if ask.data_key:
        data_keys.append(ask.data_key)

    extra = ask.extra or {}
    # Dual-bar / compare slides store extra data keys in nested dicts
    for nested_key in ("left", "right"):
        nested = extra.get(nested_key, {})
        if isinstance(nested, dict) and nested.get("data_key"):
            data_keys.append(nested["data_key"])
    # Clustered compare uses primary_key / comp_key
    for ek in ("primary_key", "comp_key"):
        if extra.get(ek):
            data_keys.append(extra[ek])

    # Build extraction lookup: id → DataExtractionConfig
    ext_map: dict[str, DataExtractionConfig] = {
        ex.id: ex for ex in config.extractions
    }

    # Build _sheets lookup for question text: code → desc
    sheets_lookup: dict[str, str] = {}
    sheets_index = data.get("_sheets", {})
    for sheet_name, rows in sheets_index.items():
        if isinstance(rows, list):
            for row in rows:
                code = row.get("code", "")
                desc = row.get("desc", "")
                if code and desc and code not in sheets_lookup:
                    sheets_lookup[code] = desc

    lines = []
    seen_codes = set()

    for dk in data_keys:
        ex = ext_map.get(dk)
        if not ex:
            continue

        params = ex.params or {}
        codes_for_ex = []

        if ex.method == "question_code":
            code = params.get("code", "")
            if code:
                codes_for_ex.append(code)

        elif ex.method == "multi_question_code":
            for entry in params.get("codes", []):
                code = entry.get("code", "")
                if code:
                    codes_for_ex.append(code)

        elif ex.method == "question_code_multi_col":
            code = params.get("code", "")
            if code:
                codes_for_ex.append(code)

        elif ex.method == "row_range":
            sheet_name = ex.sheet
            row_start = params.get("row_start", 0)
            row_end = params.get("row_end", 0)
            lines.append(f"[{dk}] sheet={sheet_name}, rows {row_start}-{row_end}")
            continue

        elif ex.method == "nested_ordinal":
            row_start = params.get("row_start", 0)
            row_end = params.get("row_end", 0)
            lines.append(f"[{dk}] nested_ordinal, rows {row_start}-{row_end}")
            continue

        elif ex.method == "mock":
            lines.append(f"[{dk}] mock data")
            continue

        # Format codes with their question text
        for code in codes_for_ex:
            if code in seen_codes:
                continue
            seen_codes.add(code)
            desc = sheets_lookup.get(code, "")
            if desc:
                lines.append(f"{code}: {desc}")
            else:
                lines.append(f"{code}")

    if not lines:
        return ""

    return "Question codes used:\n" + "\n".join(lines)


def _add_speaker_notes(slide, notes_text: str):
    """Set speaker notes on a slide."""
    if not notes_text:
        return
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = notes_text


# ── Config backup ────────────────────────────────────────────────────────────

def _backup_config(yaml_path: str) -> str:
    """Copy current config to config_history/ with timestamp."""
    project_dir = os.path.dirname(os.path.abspath(yaml_path))
    history_dir = os.path.join(project_dir, "config_history")
    os.makedirs(history_dir, exist_ok=True)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = os.path.splitext(os.path.basename(yaml_path))[0]
    backup_path = os.path.join(history_dir, f"{base}_{ts}.yaml")
    shutil.copy2(yaml_path, backup_path)
    return backup_path


# ── PPTX backup (PRD §10.3) ──────────────────────────────────────────────────

_MAX_PPTX_BACKUPS = 10


def _backup_pptx(pptx_path: str) -> str | None:
    """Create a timestamped backup of the deck before editing.

    Stores up to _MAX_PPTX_BACKUPS most recent copies in a backups/ folder
    alongside the deck. Oldest backups are deleted when the cap is exceeded.

    Returns the backup path, or None if the source doesn't exist.
    """
    if not os.path.exists(pptx_path):
        return None

    backup_dir = os.path.join(os.path.dirname(pptx_path), "backups")
    os.makedirs(backup_dir, exist_ok=True)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = os.path.splitext(os.path.basename(pptx_path))[0]
    backup_path = os.path.join(backup_dir, f"{base}_{ts}.pptx")
    shutil.copy2(pptx_path, backup_path)

    # Prune old backups
    backups = sorted(
        [f for f in os.listdir(backup_dir) if f.endswith(".pptx")],
        key=lambda f: os.path.getmtime(os.path.join(backup_dir, f)),
        reverse=True,
    )
    for old in backups[_MAX_PPTX_BACKUPS:]:
        os.remove(os.path.join(backup_dir, old))

    return backup_path


# ── Config hash ──────────────────────────────────────────────────────────────

def _file_hash(path: str) -> str:
    """Fast MD5 hash of a file."""
    with open(path, "rb") as f:
        return hashlib.md5(f.read()).hexdigest()[:12]


def _config_hash(yaml_path: str) -> str:
    """Fast hash of config file for shape registry lineage."""
    return _file_hash(yaml_path)


# ── Slide/section management (delegated to pptx_utils.deck) ─────────────────

def _clear_slide(slide):
    """Remove all shapes from a slide."""
    _clear_slide_impl(slide)


def _clear_sections(prs):
    """Remove all PowerPoint sections."""
    _clear_sections_impl(prs)


def _create_sections(prs, sections_config, ask_id_to_slide_idx):
    """Create PowerPoint sections from config."""
    _create_sections_impl(prs, sections_config, ask_id_to_slide_idx)


def _resolve_ask_id_to_slide_index(
    ask_id: str, registry_path: str, config: ProjectConfig
) -> int:
    """Resolve an ask_id string to a 0-based slide index.

    Looks up the shape_registry.json for slide metadata matching the ask_id.
    Falls back to scanning config.asks if no registry exists.

    Raises:
        ValueError: If ask_id is not found.
    """
    # Try shape_registry.json first (authoritative — accounts for skipped asks)
    if os.path.exists(registry_path):
        try:
            with open(registry_path, "r", encoding="utf-8") as f:
                reg = json.load(f)
        except (json.JSONDecodeError, OSError) as e:
            logger.warning("Corrupt shape_registry.json — ignoring: %s", e)
            reg = {}
        slides = reg.get("slides", {})
        for slide_num_str, slide_meta in slides.items():
            if slide_meta.get("ask_id") == ask_id:
                return int(slide_num_str) - 1  # registry uses 1-based slide numbers

    # Fallback: find ask_id in config.asks (assumes 1:1 mapping, no skips)
    for i, ask in enumerate(config.asks):
        if ask.id == ask_id:
            return i

    available = [a.id for a in config.asks]
    raise ValueError(
        f"ask_id '{ask_id}' not found in shape registry or config. "
        f"Available: {available}"
    )


def _load_template(config: ProjectConfig):
    """Load template PPTX, returning (Presentation, blank_layout)."""
    return _load_template_impl(config)


def generate_deck(yaml_path: str, output_path: str | None = None,
                   force_fresh: bool = False) -> str:
    """Generate a full slide deck from a YAML project config.

    Args:
        yaml_path: Path to the project YAML file.
        output_path: Override output path. If None, uses config.output_path.
        force_fresh: If True, skip JSON cache and re-extract from Excel/Synapse.

    Returns:
        Path to the generated PPTX file.
    """
    # 1. Load config
    logger.info("Loading config: %s", yaml_path)
    config = load_project_config(yaml_path)

    # 1b. Validate config consistency
    errors = config.validate()
    if errors:
        logger.warning("Config validation errors:")
        for err in errors:
            logger.warning("  - %s", err)
        raise ValueError(
            f"Config validation failed with {len(errors)} error(s):\n"
            + "\n".join(f"  - {e}" for e in errors)
        )

    # 2. Load data (from JSON if available, else extract from Excel)
    logger.info("Loading data...")
    data = load_all_data(config, force_fresh=force_fresh)

    # Extract data pull timestamp for registry lineage
    data_meta = data.get("_meta", {})
    last_data_pull = data_meta.get("extracted_at", "")

    # Print extraction summary
    for key, val in data.items():
        if key.startswith("_"):
            continue
        count = len(val) if isinstance(val, list) else "—"
        logger.info("  %s: %s rows", key, count)

    # 2b. Data completeness gate — abort if majority of extractions failed
    data_warnings = data.get("_warnings", [])
    if data_warnings:
        expected_count = len(config.extractions)
        missing_count = sum(1 for w in data_warnings if "not loaded" in w)
        if expected_count > 0 and missing_count > expected_count * 0.5:
            raise RuntimeError(
                f"Data loading failed: {missing_count}/{expected_count} extractions missing. "
                f"Aborting deck generation.\n"
                + "\n".join(f"  - {w}" for w in data_warnings)
            )

    # 2c. Stage 0q — Index qualitative data (if raw data Excel exists)
    if config.raw_data_source_path and os.path.isfile(config.raw_data_source_path):
        try:
            from slidegen.pipeline.qual_data_loader import index_qualitative
            qual_json = os.path.join(config.context_path, "qualitative_data.json")
            qual_result = index_qualitative(config.raw_data_source_path, qual_json)
            if qual_result:
                n_qs = sum(len(s.get("questions", {})) for s in qual_result.get("sheets", {}).values())
                n_resp = sum(
                    sum(q.get("n_responses", 0) for q in s.get("questions", {}).values())
                    for s in qual_result.get("sheets", {}).values()
                )
                logger.info("Stage 0q: qualitative index — %d questions, %d responses", n_qs, n_resp)
        except Exception as e:
            logger.warning("Stage 0q skipped (qualitative indexing failed): %s", e)

    # 3. Create presentation from template (preserves master logos/fonts)
    logger.info("Creating presentation...")
    prs, blank_layout = _load_template(config)

    # 4. Build slides
    slide_registries = {}
    # Track which ask index produced each slide (skipped asks don't get slides)
    slide_to_ask: list[int] = []
    ask_id_to_slide_idx: dict[str, int] = {}  # for section mapping
    cfg_hash = _config_hash(yaml_path)
    logger.info("Building slides...")
    for ask_idx, ask in enumerate(config.asks):
        renderer = RENDERERS.get(ask.slide_type)
        if renderer is None:
            logger.warning("[ask %d] SKIP — unknown slide_type: %s", ask_idx, ask.slide_type)
            continue

        # Add slide using template layout (inherits master slide logos/chrome)
        slide = prs.slides.add_slide(blank_layout)
        slide_num = len(prs.slides)
        slide_to_ask.append(ask_idx)
        ask_id_to_slide_idx[ask.id] = slide_num - 1  # 0-based

        namer = ShapeNamer(
            slide_num, ask_id=ask.id,
            data_key=getattr(ask, 'data_key', ''),
            config_hash=cfg_hash,
            source_file=config.data_source_path,
            renderer=ask.slide_type,
            last_data_pull=last_data_pull,
        )
        try:
            renderer(slide, config, ask, data, namer=namer)
            render_qual_callout(slide, config, ask)
            namer.name_remaining(slide)
            slide_registries[str(slide_num)] = namer.slide_metadata
            # Add speaker notes with question codes and descriptions
            notes_text = _build_speaker_notes(ask, config, data)
            _add_speaker_notes(slide, notes_text)
            logger.info("[%d] %s (%s)", slide_num, ask.id, ask.slide_type)
        except Exception as e:
            logger.error("Renderer error on %s (slide %d): %s", ask.id, slide_num, e, exc_info=True)
            from slidegen.pptx_utils import textbox, C_RED
            textbox(slide, f"Error: {e}", 1, 3, 10, 1, fsize=12, color=C_RED)
            # Record error in registry for post-run review
            slide_registries[str(slide_num)] = {
                "ask_id": ask.id,
                "slide_type": ask.slide_type,
                "error": str(e),
            }

    # 4b. Create PowerPoint sections (if defined in config)
    if config.sections:
        _create_sections(prs, config.sections, ask_id_to_slide_idx)

    # 5. Save
    out = output_path or config.output_path
    if not out:
        out = os.path.join(os.path.dirname(yaml_path), "output_deck.pptx")

    saved_path = _safe_save_pptx(prs, out)
    _save_shape_registry(slide_registries, os.path.dirname(saved_path),
                         slide_to_ask=slide_to_ask)
    logger.info("Saved: %s", saved_path)
    logger.info("Total slides: %d", len(prs.slides))

    return saved_path


def regenerate_slide(yaml_path: str, slide_index: int | str,
                     output_path: str | None = None) -> str:
    """Regenerate a single slide in an existing deck.

    Clears all shapes from the target slide and re-renders it from
    the current config and data. Other slides are untouched.

    Args:
        yaml_path: Path to project YAML config.
        slide_index: 0-based slide index (int) OR ask_id string to regenerate.
                     When a string is passed, the slide index is resolved from
                     shape_registry.json.
        output_path: Path to existing PPTX. If None, uses config.output_path.

    Returns:
        Path to the updated PPTX file.
    """
    config = load_project_config(yaml_path)

    # Load data (from JSON if available, else extract from Excel)
    data = load_all_data(config)
    data_meta = data.get("_meta", {})
    last_data_pull = data_meta.get("extracted_at", "")

    pptx_path = output_path or config.output_path
    if not pptx_path or not os.path.exists(pptx_path):
        raise FileNotFoundError(f"Deck not found: {pptx_path}")

    # Resolve ask_id → slide_index if a string was passed
    registry_path = os.path.join(os.path.dirname(pptx_path), "shape_registry.json")
    if isinstance(slide_index, str):
        ask_id = slide_index
        slide_index = _resolve_ask_id_to_slide_index(ask_id, registry_path, config)
        logger.info("Resolved ask_id %r → slide index %d", ask_id, slide_index)

    # Backup before editing (PRD §10.3)
    backup = _backup_pptx(pptx_path)
    if backup:
        logger.info("Backup saved: %s", backup)

    prs = Presentation(pptx_path)
    if slide_index < 0 or slide_index >= len(prs.slides):
        raise IndexError(f"Slide index {slide_index} out of range (deck has {len(prs.slides)} slides)")

    slide = prs.slides[slide_index]
    _clear_slide(slide)

    # Resolve the correct ask index — slide_to_ask mapping accounts for
    # asks that were skipped during generate_deck() (unknown slide_type).
    ask_index = slide_index  # default: assume 1:1 mapping
    if os.path.exists(registry_path):
        try:
            with open(registry_path, "r", encoding="utf-8") as f:
                reg = json.load(f)
        except (json.JSONDecodeError, OSError) as e:
            logger.warning("Corrupt shape_registry.json — ignoring: %s", e)
            reg = {}
        slide_to_ask = reg.get("slide_to_ask")
        if slide_to_ask and slide_index < len(slide_to_ask):
            ask_index = slide_to_ask[slide_index]

    if ask_index >= len(config.asks):
        raise IndexError(
            f"Ask index {ask_index} (from slide {slide_index}) out of range "
            f"(config has {len(config.asks)} asks)"
        )
    ask = config.asks[ask_index]
    renderer = RENDERERS.get(ask.slide_type)
    if renderer is None:
        raise ValueError(f"Unknown slide_type: {ask.slide_type}")

    namer = ShapeNamer(
        slide_index + 1, ask_id=ask.id,
        data_key=getattr(ask, 'data_key', ''),
        config_hash=_config_hash(yaml_path),
        source_file=config.data_source_path,
        renderer=ask.slide_type,
        last_data_pull=last_data_pull,
    )
    renderer(slide, config, ask, data, namer=namer)
    render_qual_callout(slide, config, ask)
    namer.name_remaining(slide)
    # Add speaker notes with question codes and descriptions
    notes_text = _build_speaker_notes(ask, config, data)
    _add_speaker_notes(slide, notes_text)

    saved = _safe_save_pptx(prs, pptx_path)
    logger.info("Regenerated slide %d (%s) in %s", slide_index + 1, ask.id, saved)
    return saved


# ── Deck refresh (PRD §9.2) ──────────────────────────────────────────────────

def refresh_deck(yaml_path: str, output_path: str | None = None) -> dict:
    """Refresh all data-driven slides in an existing deck.

    Forces re-extraction of data from Excel/Synapse, then regenerates every
    slide that has a data_source recorded in shape_registry.json.

    Args:
        yaml_path: Path to project YAML config.
        output_path: Path to existing PPTX. If None, uses config.output_path.

    Returns:
        Summary dict: {refreshed: [ask_ids], skipped: [ask_ids], errors: [...]}
    """
    config = load_project_config(yaml_path)
    pptx_path = output_path or config.output_path
    if not pptx_path or not os.path.exists(pptx_path):
        raise FileNotFoundError(f"Deck not found: {pptx_path}")

    registry_path = os.path.join(os.path.dirname(pptx_path), "shape_registry.json")
    if not os.path.exists(registry_path):
        raise FileNotFoundError(f"Shape registry not found: {registry_path}")

    # Load registry to find data-driven slides
    try:
        with open(registry_path, "r", encoding="utf-8") as f:
            reg = json.load(f)
    except json.JSONDecodeError as e:
        raise ValueError(f"Corrupt shape_registry.json — delete and regenerate deck: {e}")

    slides_meta = reg.get("slides", {})

    # Force fresh data extraction
    logger.info("Refreshing deck — forcing data re-extraction...")
    data = load_all_data(config, force_fresh=True)
    data_meta = data.get("_meta", {})
    last_data_pull = data_meta.get("extracted_at", "")

    # Backup before refresh
    backup = _backup_pptx(pptx_path)
    if backup:
        logger.info("Backup saved: %s", backup)

    prs = Presentation(pptx_path)
    slide_to_ask = reg.get("slide_to_ask", [])
    cfg_hash = _config_hash(yaml_path)

    result = {"refreshed": [], "skipped": [], "errors": []}

    for slide_num_str, slide_info in slides_meta.items():
        slide_idx = int(slide_num_str) - 1  # registry uses 1-based
        ask_id = slide_info.get("ask_id", "")

        # Skip slides without data sources (cover, ES, etc.)
        if not slide_info.get("data_source"):
            result["skipped"].append(ask_id or f"slide_{slide_num_str}")
            continue

        if slide_idx < 0 or slide_idx >= len(prs.slides):
            result["errors"].append(f"{ask_id}: slide index {slide_idx} out of range")
            continue

        # Find the ask config
        ask_index = slide_idx
        if slide_to_ask and slide_idx < len(slide_to_ask):
            ask_index = slide_to_ask[slide_idx]

        if ask_index >= len(config.asks):
            result["errors"].append(f"{ask_id}: ask index {ask_index} out of range")
            continue

        ask = config.asks[ask_index]
        renderer_fn = RENDERERS.get(ask.slide_type)
        if renderer_fn is None:
            result["errors"].append(f"{ask_id}: unknown slide_type '{ask.slide_type}'")
            continue

        # Clear and re-render
        slide = prs.slides[slide_idx]
        _clear_slide(slide)

        namer = ShapeNamer(
            slide_idx + 1, ask_id=ask.id,
            data_key=getattr(ask, 'data_key', ''),
            config_hash=cfg_hash,
            source_file=config.data_source_path,
            renderer=ask.slide_type,
            last_data_pull=last_data_pull,
        )
        try:
            renderer_fn(slide, config, ask, data, namer=namer)
            namer.name_remaining(slide)
            notes_text = _build_speaker_notes(ask, config, data)
            _add_speaker_notes(slide, notes_text)
            # Update registry entry for this slide
            slides_meta[slide_num_str] = namer.slide_metadata
            result["refreshed"].append(ask_id)
            logger.info("[%s] %s refreshed", slide_num_str, ask_id)
        except Exception as e:
            logger.error("Refresh error on %s: %s", ask_id, e, exc_info=True)
            result["errors"].append(f"{ask_id}: {e}")
            logger.error("[%s] %s ERROR: %s", slide_num_str, ask_id, e)
            # Record error in registry for post-run review
            slides_meta[slide_num_str] = {
                "ask_id": ask_id,
                "error": str(e),
            }

    # Save updated deck
    pptx_path = _safe_save_pptx(prs, pptx_path, suffix="_refreshed")

    # Save updated registry
    reg["slides"] = slides_meta
    reg["last_refreshed"] = datetime.now().isoformat()
    _atomic_json_write(registry_path, reg)

    logger.info("Refresh complete: %d refreshed, %d skipped, %d errors",
                len(result['refreshed']), len(result['skipped']), len(result['errors']))
    return result


# ── CLI entry point ──────────────────────────────────────────────────────────

if __name__ == "__main__":
    import sys
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = [a for a in sys.argv[1:] if a.startswith("--")]

    if len(args) < 1:
        print("Usage: python -m slidegen.pipeline.orchestrator <project.yaml> [output.pptx] [--fresh]")
        sys.exit(1)

    yaml_path = args[0]
    output = args[1] if len(args) > 1 else None
    fresh = "--fresh" in flags
    generate_deck(yaml_path, output, force_fresh=fresh)
