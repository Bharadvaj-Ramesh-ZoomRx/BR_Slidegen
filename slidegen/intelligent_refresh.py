"""
intelligent_refresh.py — Intelligent slide refresh orchestrated by Claude Code.

Core module for the non-connected slide refresh pipeline. Extracts visual context
from PPTX slides, fetches data from Synapse API, and applies Claude Code's
mapping to refresh charts and tables.

Two-phase pipeline:
  Phase 1 (read):   Extract slide context + fetch data -> structured summary
                     Claude Code reads this output and reasons about mappings
  Phase 2 (refresh): Apply Claude Code's mapping to refresh the deck

Usage:
    python -m slidegen.intelligent_refresh read --pptx path.pptx [--slide 0]
    python -m slidegen.intelligent_refresh refresh --pptx path.pptx --output out.pptx --mapping mapping.json --lineage lineage.json
    python -m slidegen.intelligent_refresh headline --pptx path.pptx --slide 0 --text "New headline text"
"""
from __future__ import annotations

import json
import os
import shutil
import sys
from pathlib import Path

import pandas as pd
import requests
from dotenv import load_dotenv
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

REPO_ROOT = Path(__file__).resolve().parents[1]
load_dotenv(REPO_ROOT / ".env")

_CHART_TYPE_MAP = {
    XL_CHART_TYPE.BAR_STACKED_100: "bar_stacked_100_horizontal",
    XL_CHART_TYPE.BAR_CLUSTERED: "bar_clustered_horizontal",
    XL_CHART_TYPE.COLUMN_CLUSTERED: "column_clustered_vertical",
    XL_CHART_TYPE.COLUMN_STACKED_100: "column_stacked_100_vertical",
    XL_CHART_TYPE.LINE_MARKERS: "line_markers_trended",
    XL_CHART_TYPE.LINE: "line_markers_trended",
    XL_CHART_TYPE.XY_SCATTER: "xy_scatter_abacus",
    XL_CHART_TYPE.DOUGHNUT: "doughnut_default",
}


def _in(emu):
    """Convert EMU to inches, rounded to 2 decimal places."""
    return round(emu / 914400, 2) if emu else 0.0


# ══════════════════════════════════════════════════════════════════════════════
# Phase 1: Read — extract slide context
# ══════════════════════════════════════════════════════════════════════════════

def read_slide_context(pptx_path: str, slide_index: int = 0) -> dict:
    """Extract all visual context from one slide: shapes, positions, text,
    chart data, table data, group shape labels.

    Returns a dict ready to be printed/read by Claude Code.
    """
    prs = Presentation(pptx_path)
    slide = prs.slides[slide_index]

    shapes = []
    for shape in slide.shapes:
        left, top = _in(shape.left), _in(shape.top)
        width, height = _in(shape.width), _in(shape.height)
        base = {
            "name": shape.name, "left": left, "top": top,
            "width": width, "height": height,
        }

        if shape.has_chart:
            chart = shape.chart
            plot = chart.plots[0]
            cats = [str(c) for c in plot.categories] if plot.categories else []
            s_names = [str(s.name) for s in plot.series]
            s_vals = {}
            for s in plot.series:
                try:
                    s_vals[str(s.name)] = [
                        round(float(v), 4) if v else 0
                        for v in s.values
                    ][:4]
                except Exception:
                    pass
            shapes.append({
                **base, "type": "chart",
                "chart_pattern": _CHART_TYPE_MAP.get(chart.chart_type, "unknown"),
                "series_names": s_names, "categories": cats,
                "series_values": s_vals,
            })

        elif shape.has_table:
            tbl = shape.table
            nr, nc = len(tbl.rows), len(tbl.columns)
            headers = (
                [tbl.cell(0, c).text.strip() for c in range(nc)]
                if nr > 0 else []
            )
            sample = []
            for r in range(1, min(nr, 4)):
                sample.append([tbl.cell(r, c).text.strip() for c in range(nc)])
            shapes.append({
                **base, "type": "table",
                "row_count": nr, "col_count": nc,
                "headers": headers, "sample_rows": sample,
            })

        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                shapes.append({**base, "type": "text", "text": text[:300]})

        # Group shape children with text
        try:
            if hasattr(shape, 'shapes'):
                for child in shape.shapes:
                    if child.has_text_frame:
                        t = child.text_frame.text.strip()
                        if t:
                            shapes.append({
                                "name": f"{shape.name}/{child.name}",
                                "left": left, "top": top,
                                "width": width, "height": height,
                                "type": "label", "text": t[:200],
                            })
        except Exception:
            pass

    shapes.sort(key=lambda s: (s["left"], s["top"]))

    return {
        "pptx_path": str(pptx_path),
        "slide_index": slide_index,
        "shapes": shapes,
    }


def format_slide_for_interpretation(context: dict) -> str:
    """Format the slide context as a readable text description for Claude Code
    to interpret."""
    lines = []
    lines.append(f"=== Slide {context.get('slide_index', 0)} from {context.get('pptx_path', '?')} ===")
    lines.append("")

    for s in context.get("shapes", []):
        stype = s.get("type", "unknown")
        pos = f"({s['left']:.1f}\", {s['top']:.1f}\") {s['width']:.1f}\"x{s['height']:.1f}\""

        if stype == "chart":
            lines.append(f"[CHART] {s['name']} at {pos}")
            lines.append(f"  Pattern: {s.get('chart_pattern', '?')}")
            lines.append(f"  Categories: {s.get('categories', [])}")
            lines.append(f"  Series: {s.get('series_names', [])}")
            for sn, sv in s.get("series_values", {}).items():
                lines.append(f"    {sn}: {sv}")
            lines.append("")

        elif stype == "table":
            lines.append(f"[TABLE] {s['name']} at {pos}")
            lines.append(f"  {s.get('row_count', '?')} rows x {s.get('col_count', '?')} cols")
            lines.append(f"  Headers: {s.get('headers', [])}")
            for row in s.get("sample_rows", []):
                lines.append(f"    {row}")
            lines.append("")

        elif stype == "text":
            lines.append(f"[TEXT] {s['name']} at {pos}")
            lines.append(f"  \"{s.get('text', '')}\"")
            lines.append("")

        elif stype == "label":
            lines.append(f"[LABEL] {s['name']} at {pos}")
            lines.append(f"  \"{s.get('text', '')}\"")
            lines.append("")

    return "\n".join(lines)


def format_data_for_interpretation(df: pd.DataFrame) -> str:
    """Format fetched data as a readable summary for Claude Code."""
    if df.empty:
        return "No data fetched."

    lines = []
    lines.append(f"=== Fetched Data: {len(df)} records, {len(df.columns)} columns ===")
    lines.append("")
    lines.append(f"Columns: {sorted(df.columns.tolist())}")
    lines.append("")

    for col in df.columns:
        unique = df[col].dropna().unique()
        if len(unique) <= 20:
            lines.append(f"  {col}: {sorted(str(v) for v in unique)}")

    lines.append("")
    lines.append("First 5 rows:")
    lines.append(df.head(5).to_string(index=False))

    return "\n".join(lines)


# ══════════════════════════════════════════════════════════════════════════════
# Data Fetching
# ══════════════════════════════════════════════════════════════════════════════

def fetch_synapse_data(data_lineage: dict) -> tuple[list[dict], pd.DataFrame]:
    """Fetch records from Synapse API using data lineage dict.

    Returns (records_list, dataframe).
    """
    token = os.environ.get("SYNAPSE_API_TOKEN", "")
    base_url = os.environ.get("SYNAPSE_API_BASE_URL", "https://synapse-api.zoomrx.com")
    headers = {
        "Authorization": token if token.startswith("Bearer") else f"Bearer {token}",
        "accept": "application/json",
        "Content-Type": "application/json",
    }
    payload = {
        "project_id": data_lineage["project_id"],
        "reporting_plan_id": data_lineage["reporting_plan_id"],
        "analysis_ids": data_lineage["analysis_ids"],
        "segment_ids": data_lineage.get("segment_ids", []),
        "setup_type": "DYNAMIC",
        "dynamic_time_period": {
            "latest_n_deliverables": data_lineage.get("dynamic_latest_n", 5),
            "include_live_wave": True,
        },
    }
    resp = requests.post(
        f"{base_url}/api/reports/generate",
        headers=headers, json=payload, timeout=60,
    )
    records = []
    if resp.status_code in (200, 201):
        records = resp.json().get("records", [])

    df = pd.DataFrame(records) if records else pd.DataFrame()
    return records, df


# ══════════════════════════════════════════════════════════════════════════════
# Phase 2: Refresh — apply mapping to update the deck
# ══════════════════════════════════════════════════════════════════════════════

def refresh_slide_from_mapping(
    pptx_path: str,
    output_path: str,
    slide_index: int,
    mapping: dict,
    data_lineage: dict,
) -> dict:
    """Apply a mapping dict to refresh charts + tables on one slide.

    Returns results dict with status per chart/table.

    mapping format:
    {
      "charts": [
        {"chart_name": "Chart 35", "segment_filter": "CARD",
         "series_column": "x_code", "series_name_map": {"IDK": "0"},
         "row_field": "y_label", "value_field": "percentage"}
      ],
      "tables": [
        {"table_name": "Table 25", "segment_filter": "CARD",
         "columns": [
           {"header": "Product", "data_field": "y_label", "format": "string"},
           {"header": "Base", "data_field": "base", "format": "(n = {})"},
           {"header": "", "data_field": null, "format": "spacer"},
           {"header": "Easy", "data_field": "percentage", "format": "{}%",
            "filter": {"x_code": "H"}}
         ]}
      ]
    }
    """
    # Fetch data
    records, df = fetch_synapse_data(data_lineage)
    if df.empty:
        return {"error": "No data returned from Synapse API"}

    # Read source for series order
    src_prs = Presentation(pptx_path)
    src_slide = src_prs.slides[slide_index]
    src_chart_series = {}
    for shape in src_slide.shapes:
        if shape.has_chart:
            src_chart_series[shape.name] = [
                str(s.name) for s in shape.chart.plots[0].series
            ]

    # Clone + refresh
    shutil.copy2(pptx_path, output_path)
    prs = Presentation(output_path)
    slide = prs.slides[slide_index]
    chart_shapes = {s.name: s for s in slide.shapes if s.has_chart}
    table_shapes = {s.name: s for s in slide.shapes if s.has_table}
    ns_c = "http://schemas.openxmlformats.org/drawingml/2006/chart"

    results = {"charts": [], "tables": []}

    # ── Charts ──
    for cm in mapping.get("charts", []):
        name = cm["chart_name"]
        shape = chart_shapes.get(name)
        if not shape:
            results["charts"].append({"name": name, "status": "not_found"})
            continue

        sub = df.copy()
        seg = cm.get("segment_filter")
        if seg and "segment_1" in sub.columns:
            sub = sub[sub["segment_1"] == seg]
        tp = cm.get("time_period")
        if tp and "time_period_name" in sub.columns:
            sub = sub[sub["time_period_name"] == tp]
        if sub.empty:
            results["charts"].append({"name": name, "status": "empty"})
            continue

        row_f = cm.get("row_field", "y_label")
        col_f = cm.get("series_column")
        val_f = cm.get("value_field", "percentage")
        name_map = cm.get("series_name_map", {})

        try:
            if col_f and col_f in sub.columns:
                pivot = sub.pivot_table(
                    index=row_f, columns=col_f, values=val_f, aggfunc="first",
                )
            else:
                pivot = sub.groupby(row_f, sort=False)[val_f].first().to_frame()

            categories = list(pivot.index.astype(str))
            orig_series = src_chart_series.get(name, [])

            series_data = []
            for sn in orig_series:
                data_val = name_map.get(sn, sn)
                matched = None
                for pc in pivot.columns:
                    if str(pc) == data_val or str(pc).lower() == data_val.lower():
                        matched = pc
                        break
                vals = (
                    [0.0 if pd.isna(v) else round(float(v), 4) for v in pivot[matched]]
                    if matched is not None
                    else [0.0] * len(categories)
                )
                series_data.append((sn, vals))

            src_fmts = [
                fc.text
                for fc in shape.chart._chartSpace.iter(f"{{{ns_c}}}formatCode")
            ]
            cd = CategoryChartData()
            cd.categories = categories
            for sn, vals in series_data:
                cd.add_series(sn, vals)
            shape.chart.replace_data(cd)
            for i, fc_el in enumerate(
                shape.chart._chartSpace.iter(f"{{{ns_c}}}formatCode")
            ):
                if i < len(src_fmts):
                    fc_el.text = src_fmts[i]

            results["charts"].append({
                "name": name, "status": "ok", "segment": seg,
                "categories": categories[:3], "n_series": len(series_data),
            })
        except Exception as e:
            results["charts"].append({
                "name": name, "status": "error", "error": str(e),
            })

    # ── Tables ──
    for tm in mapping.get("tables", []):
        name = tm["table_name"]
        shape = table_shapes.get(name)
        if not shape:
            results["tables"].append({"name": name, "status": "not_found"})
            continue

        tbl = shape.table
        sub = df.copy()
        seg = tm.get("segment_filter")
        if seg and "segment_1" in sub.columns:
            sub = sub[sub["segment_1"] == seg]
        if sub.empty:
            results["tables"].append({"name": name, "status": "empty"})
            continue

        columns_spec = tm.get("columns", [])
        if not columns_spec:
            results["tables"].append({"name": name, "status": "no_columns"})
            continue

        try:
            row_key = None
            for cs in columns_spec:
                if cs.get("data_field") and cs.get("format") != "spacer":
                    row_key = cs["data_field"]
                    break
            if not row_key or row_key not in sub.columns:
                results["tables"].append({"name": name, "status": "no_row_key"})
                continue

            keys = sub[row_key].dropna().unique().tolist()
            n_rows = min(len(keys), len(tbl.rows) - 1)

            for ri in range(n_rows):
                row_data = sub[sub[row_key] == keys[ri]]
                if row_data.empty:
                    continue
                for ci, cs in enumerate(columns_spec):
                    if ci >= len(tbl.columns) or ri + 1 >= len(tbl.rows):
                        break
                    fmt = cs.get("format", "string")
                    data_f = cs.get("data_field")
                    col_filter = cs.get("filter")
                    if fmt == "spacer" or not data_f:
                        continue

                    cell_data = row_data
                    if col_filter and isinstance(col_filter, dict):
                        for fk, fv in col_filter.items():
                            if fk in cell_data.columns:
                                cell_data = cell_data[cell_data[fk] == fv]
                    if cell_data.empty or data_f not in cell_data.columns:
                        continue

                    raw = cell_data.iloc[0][data_f]
                    if fmt == "string":
                        text = str(raw)
                    elif "(n" in fmt:
                        text = f"(n = {int(float(raw))})" if raw else ""
                    elif "%" in fmt:
                        v = float(raw)
                        text = f"{v:.0%}" if abs(v) <= 1.0 else f"{int(v)}%"
                    else:
                        text = str(raw)

                    cell = tbl.cell(ri + 1, ci)
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.text = ""
                    if cell.text_frame.paragraphs and cell.text_frame.paragraphs[0].runs:
                        cell.text_frame.paragraphs[0].runs[0].text = text

            results["tables"].append({
                "name": name, "status": "ok", "segment": seg, "rows": n_rows,
            })
        except Exception as e:
            results["tables"].append({
                "name": name, "status": "error", "error": str(e),
            })

    prs.save(output_path)
    return results


# ══════════════════════════════════════════════════════════════════════════════
# Headline Writer
# ══════════════════════════════════════════════════════════════════════════════

def write_headline(pptx_path: str, slide_index: int, headline_text: str) -> None:
    """Write a headline into the slide's headline text shape (top area, largest text).

    Finds the headline shape: text shape in the top 1.5" zone with the longest text.
    Clears all runs, writes the new headline text into the first run preserving formatting.
    Saves the PPTX.
    """
    prs = Presentation(pptx_path)
    slide = prs.slides[slide_index]

    # Find headline shape: text shapes in top 1.5" zone, pick the one with longest text
    TOP_ZONE = 1.5  # inches
    candidates = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        top_in = _in(shape.top)
        if top_in > TOP_ZONE:
            continue
        text = shape.text_frame.text.strip()
        if text:
            candidates.append((shape, len(text)))

    if not candidates:
        raise ValueError(
            f"No text shapes found in top {TOP_ZONE}\" zone of slide {slide_index}"
        )

    # Pick shape with longest existing text (most likely the headline)
    headline_shape = max(candidates, key=lambda x: x[1])[0]

    # Preserve first paragraph's first run formatting, clear everything else
    first_para = headline_shape.text_frame.paragraphs[0]

    # Save formatting from first run if it exists
    saved_font = None
    if first_para.runs:
        run = first_para.runs[0]
        saved_font = {
            "bold": run.font.bold,
            "italic": run.font.italic,
            "size": run.font.size,
            "color_rgb": run.font.color.rgb if run.font.color and run.font.color.rgb else None,
            "name": run.font.name,
        }

    # Clear all runs in all paragraphs
    for para in headline_shape.text_frame.paragraphs:
        for run in para.runs:
            run.text = ""

    # Write new headline text into first run of first paragraph
    if first_para.runs:
        first_para.runs[0].text = headline_text
    else:
        # No runs exist — add one
        run = first_para.add_run()
        run.text = headline_text
        if saved_font:
            if saved_font["bold"] is not None:
                run.font.bold = saved_font["bold"]
            if saved_font["italic"] is not None:
                run.font.italic = saved_font["italic"]
            if saved_font["size"] is not None:
                run.font.size = saved_font["size"]
            if saved_font["color_rgb"] is not None:
                run.font.color.rgb = saved_font["color_rgb"]
            if saved_font["name"] is not None:
                run.font.name = saved_font["name"]

    prs.save(pptx_path)


# ══════════════════════════════════════════════════════════════════════════════
# CLI
# ══════════════════════════════════════════════════════════════════════════════

def main():
    import argparse

    parser = argparse.ArgumentParser(
        description="Intelligent slide refresh — read, refresh, or write headline",
    )
    sub = parser.add_subparsers(dest="command")

    # ── read ──
    p_read = sub.add_parser("read", help="Extract slide context + optionally fetch data")
    p_read.add_argument("--pptx", required=True, help="Source PPTX path")
    p_read.add_argument("--slide", type=int, default=0, help="Slide index (0-based)")
    p_read.add_argument(
        "--lineage", type=str, default=None,
        help="Path to data lineage JSON (if provided, also fetches Synapse data)",
    )

    # ── refresh ──
    p_refresh = sub.add_parser("refresh", help="Apply mapping to refresh the slide")
    p_refresh.add_argument("--pptx", required=True, help="Source PPTX path")
    p_refresh.add_argument("--output", required=True, help="Output PPTX path")
    p_refresh.add_argument("--slide", type=int, default=0, help="Slide index (0-based)")
    p_refresh.add_argument("--mapping", required=True, help="Path to mapping JSON")
    p_refresh.add_argument("--lineage", required=True, help="Path to data lineage JSON")

    # ── headline ──
    p_headline = sub.add_parser("headline", help="Write headline text into a slide")
    p_headline.add_argument("--pptx", required=True, help="PPTX path (modified in place)")
    p_headline.add_argument("--slide", type=int, default=0, help="Slide index (0-based)")
    p_headline.add_argument("--text", required=True, help="Headline text to write")

    args = parser.parse_args()

    if args.command == "read":
        context = read_slide_context(args.pptx, slide_index=args.slide)
        print(format_slide_for_interpretation(context))

        if args.lineage:
            lineage = json.loads(Path(args.lineage).read_text(encoding="utf-8"))
            records, df = fetch_synapse_data(lineage)
            print()
            print(format_data_for_interpretation(df))

    elif args.command == "refresh":
        mapping = json.loads(Path(args.mapping).read_text(encoding="utf-8"))
        lineage = json.loads(Path(args.lineage).read_text(encoding="utf-8"))
        results = refresh_slide_from_mapping(
            args.pptx, args.output, args.slide, mapping, lineage,
        )
        print(json.dumps(results, indent=2, ensure_ascii=False))

    elif args.command == "headline":
        write_headline(args.pptx, args.slide, args.text)
        print(f"Headline written to slide {args.slide}: \"{args.text}\"")

    else:
        parser.print_help()


if __name__ == "__main__":
    main()
