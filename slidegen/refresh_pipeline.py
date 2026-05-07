"""End-to-end deck refresh: spec -> fetch -> stamp -> headlines -> label sync -> badges.

Generic over any deck with Galen Connector tags. Replaces per-deck runner
scripts (`refresh_creon_today.py`, `refresh_aveo_today.py`, etc.) with a
single CLI that auto-derives the output filename from the source deck name.

Usage (CLI):
    python -m slidegen.refresh_pipeline <deck.pptx>
    python -m slidegen.refresh_pipeline <deck.pptx> --out-name MyDeck_v2
    python -m slidegen.refresh_pipeline <deck.pptx> --no-cache --workers 8

Usage (Python):
    from slidegen.refresh_pipeline import run_full_pipeline
    result = run_full_pipeline("path/to/deck.pptx")
    # result["deck"] points at the final annotated PPTX
    # result["tag_mismatch_slides"] lists slides flagged for review

Steps:
    0. Regenerate spec from the deck's Connector tags (so static_time_period_ids,
       include_live_wave overrides, and split-viz positions are picked up live
       from the deck, not a stale committed JSON).
    1. refresh_deck_from_spec — fetch all data sources in parallel, write
       refreshed values into chart cells and table cell_values.
    2. stamp_refresh_notes — per-shape REFRESH_NOTE Connector tags so the
       diagnostic context is visible in the deck.
    3. refresh_headlines (LLM) — rewrite slide headlines based on refreshed
       data deltas. Skipped if LLM_API_KEY is not set.
    4. sync_period_labels — rewrite period banners in text frames adjacent
       to charts so on-slide labels track the chart's new wave window.
    5. Annotate badges — REFRESH and HEADLINE rounded-rect badges on every
       slide with bucket counts.

Outputs (under output_testing/ by default):
    deck_output/{out_name}.pptx              — final annotated deck
    json_output/{out_name}_full_spec.json    — regenerated spec
    json_output/{out_name}_refresh_status.json
    json_output/{out_name}_headline_status.json
    json_output/{out_name}_label_sync_status.json
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]

# Load env files (LLM creds + Synapse creds) — same chain as the per-deck runners.
from dotenv import load_dotenv  # noqa: E402
load_dotenv(REPO_ROOT / ".env", override=False)
_DOCS_ENV = REPO_ROOT / "docs" / ".env"
if _DOCS_ENV.exists():
    load_dotenv(_DOCS_ENV, override=False)
_SHARED_SYNAPSE_ENV = REPO_ROOT / "synapse env" / ".env"
if _SHARED_SYNAPSE_ENV.exists():
    load_dotenv(_SHARED_SYNAPSE_ENV, override=False)

from slidegen.intelligent_refresh import (  # noqa: E402
    refresh_deck_from_spec,
    stamp_refresh_notes,
    walk_shapes_recursive,
)
from slidegen.label_sync import sync_period_labels  # noqa: E402


def _slugify(name: str) -> str:
    """Make a filesystem/git-friendly slug from a deck filename stem.

    Strips parenthesized suffixes (e.g. "(1)") and smart quotes, collapses
    runs of non-alphanumeric chars to single underscores, trims leading/
    trailing underscores. Examples:

        "CREON Share of Voice Study - W33"     -> "CREON_Share_of_Voice_Study_W33"
        "Repatha HCP ATU - Q2'26 Skeleton (1)" -> "Repatha_HCP_ATU_Q2_26_Skeleton"
        "AVEO Wave 5 PET Report v1.0"          -> "AVEO_Wave_5_PET_Report_v1_0"
    """
    s = re.sub(r"\([^)]*\)", " ", name)
    # Replace quotes with separator (not strip) so "Q2'26" -> "Q2_26", not "Q226"
    s = s.replace("’", " ").replace("‘", " ").replace("'", " ")
    s = re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")
    return s or "deck"


def _regen_spec(source_pptx: Path, out_spec: Path) -> None:
    """Regenerate full spec from a deck via tag_reader -> build_full_spec.

    Walks every chart/table shape on every slide, matches by position +
    name, and writes a complete spec JSON the refresher can consume.
    """
    from slidegen.deck_reader.tag_reader import generate_config_specs
    from slidegen.slide_spec.schema import dump_spec
    from tests.build_full_spec import (
        build_data_sources, build_connected_slide, _lineage_ds_key,
    )
    from pptx import Presentation

    specs, _summary = generate_config_specs(str(source_pptx))
    connector_specs = [json.loads(dump_spec(s)) for s in specs]
    data_sources = build_data_sources(connector_specs)
    slide_to_ds = {
        s["slide_index"]: _lineage_ds_key(s.get("data_lineage", {}))
        for s in connector_specs
        if _lineage_ds_key(s.get("data_lineage", {}))
    }

    prs = Presentation(str(source_pptx))
    slide_entries = []
    for spec in connector_specs:
        si = spec["slide_index"]
        ds_key = slide_to_ds.get(si)
        if not ds_key:
            continue
        slide_entry = build_connected_slide(spec, ds_key)
        if si < len(prs.slides):
            pptx_slide = prs.slides[si]
            # walk_shapes_recursive descends into Group shapes so a chart
            # or table nested in a Group still has its connector tag
            # routed to the right physical shape.
            all_leaf = list(walk_shapes_recursive(pptx_slide.shapes))
            chart_shapes = [s for s in all_leaf if s.has_chart]
            table_shapes = [s for s in all_leaf if s.has_table]
            # Match by CLOSEST position (not first-within-tolerance) so dense
            # slides like Repatha ATU 11/12 — where n-size tables sit 0.13in
            # below value tables — don't pick the wrong shape. Also record
            # the shape's intrinsic XML id so the refresher can look up
            # exactly the right shape regardless of name collisions.
            for comp in slide_entry["components"]:
                pos = comp.get("position", {})
                cl = float(pos.get("left", 0) or 0)
                ct = float(pos.get("top", 0) or 0)
                cands = chart_shapes if comp["type"] == "chart" else table_shapes
                if not cands:
                    continue
                best_shape = None
                best_dist = float("inf")
                for shape in cands:
                    sl = (shape.left or 0) / 914400
                    st = (shape.top or 0) / 914400
                    d = abs(sl - cl) + abs(st - ct)
                    if d < best_dist:
                        best_dist = d
                        best_shape = shape
                # Accept if the closest shape is within 0.5in (loose enough
                # for minor xfrm drift, tight enough that we never grab a
                # genuinely different shape elsewhere on the slide).
                if best_shape is not None and best_dist < 0.5:
                    comp["name"] = best_shape.name
                    try:
                        comp["shape_id"] = int(best_shape.shape_id)
                    except Exception:
                        pass
        slide_entries.append(slide_entry)

    full_spec = {
        "source_deck": str(source_pptx),
        "data_sources": data_sources,
        "slides": slide_entries,
    }
    out_spec.write_text(
        json.dumps(full_spec, indent=2, ensure_ascii=False), encoding="utf-8",
    )


def _check_synapse_creds() -> tuple[bool, str]:
    if not os.getenv("SYNAPSE_API_URL"):
        return False, "SYNAPSE_API_URL not set in .env"
    if not os.getenv("SYNAPSE_API_KEY"):
        return True, "SYNAPSE_API_KEY not in .env — will try synapse-cli fallback chain"
    return True, "ok"


_NOTES_DIVIDER = "─" * 60
_NOTES_HEADER = f"\n\n{_NOTES_DIVIDER}\n[Refresh status —"


def _append_refresh_notes(slide, slide_status: dict | None,
                          headline_update: dict | None, today: str) -> None:
    """Append a per-component refresh report to the slide's speaker notes.

    Existing notes (the analyst's question text, sample sizes, etc.) are
    PRESERVED. The new content is appended below a divider so it's visually
    separated. If a previous run wrote a refresh report (detected by the
    same divider+header), that block is replaced — we don't accumulate.
    """
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    existing = tf.text or ""

    # Strip any prior refresh-status block written by an earlier run, so
    # repeated refreshes don't pile on. Marker is the literal divider+header.
    idx = existing.find(_NOTES_HEADER)
    if idx >= 0:
        existing = existing[:idx].rstrip()

    lines = [_NOTES_DIVIDER, f"[Refresh status — {today}]"]
    if slide_status is None:
        lines.append("  Slide is non-connected — no automatic refresh.")
    else:
        comps = (slide_status.get("charts") or []) + (slide_status.get("tables") or [])
        if not comps:
            lines.append("  No components processed for this slide.")
        else:
            buckets: dict[str, list[str]] = {}
            for c in comps:
                key = c.get("status", "missing")
                buckets.setdefault(key, []).append(c.get("name", "?"))
            order = [
                "ok", "ok_with_dynamic_added", "static_pinned_skipped",
                "ambiguous_tag", "alignment_failed", "tag_mismatch",
                "no_data_for_shifted_window", "empty", "not_found",
            ]
            seen = set()
            for k in order:
                if k in buckets:
                    seen.add(k)
                    lines.append(f"  {k:24s} {', '.join(buckets[k])}")
            for k in buckets:
                if k not in seen:
                    lines.append(f"  {k:24s} {', '.join(buckets[k])}")
            # Surface diagnostic notes (tag_mismatch detail, drift, etc.)
            for c in comps:
                for n in (c.get("notes") or []):
                    detail = (n.get("detail") or "").strip()
                    if detail:
                        lines.append(f"  note ({c.get('name')}): {detail[:140]}")

    if headline_update:
        st = headline_update.get("status", "")
        if st == "updated":
            old = (headline_update.get("old_headline") or "").strip()[:80]
            new = (headline_update.get("new_headline") or "").strip()[:80]
            lines.append(f"  headline: rewritten ({old!r} -> {new!r})")
        elif st in ("unchanged", "no_chart", "no_headline"):
            lines.append(f"  headline: {st}")

    # Append new paragraphs via lxml so existing speaker-notes content
    # (Q text, sample sizes, formatting) stays untouched. Setting
    # tf.text would clobber paragraph-level formatting on existing
    # notes — we only want to APPEND a divider + report block.
    from lxml import etree as _et
    NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    txBody = tf._txBody
    QA = lambda tag: f"{{{NS_A}}}{tag}"

    # First, remove any prior refresh block: walk paragraphs, find one
    # whose text starts with the divider marker, and delete it + every
    # paragraph after it (the prior report extends to end-of-notes).
    paras = txBody.findall(QA("p"))
    cut_from = None
    for i, p in enumerate(paras):
        ptext = "".join(t.text or "" for t in p.iter(QA("t")))
        if _NOTES_DIVIDER in ptext:
            cut_from = i
            break
    if cut_from is not None:
        for p in paras[cut_from:]:
            txBody.remove(p)

    # Append: blank separator + each new line as its own paragraph.
    def _add_para(text: str):
        p = _et.SubElement(txBody, QA("p"))
        if text:
            r = _et.SubElement(p, QA("r"))
            rPr = _et.SubElement(r, QA("rPr"))
            rPr.set("lang", "en-US")
            rPr.set("dirty", "0")
            t = _et.SubElement(r, QA("t"))
            t.text = text
        else:
            _et.SubElement(p, QA("endParaRPr")).set("lang", "en-US")

    _add_para("")  # blank spacer between existing and new block
    for line in lines:
        _add_para(line)


def _bucket_label(r_label: str) -> str:
    """Map a refresh-badge label to a slide-bucket key."""
    if "NON-CONNECTED" in r_label:
        return "non_connected"
    if "tag-mismatch" in r_label:
        return "tag_mismatch"
    if "alignment-failed" in r_label:
        return "alignment_failed"
    # Mixed slide: some components refreshed, others static-pinned.
    # Badge text is "REFRESH N ok / M static". Must check before the plain
    # static check below so it doesn't get swallowed by "static_review".
    if r_label.startswith("REFRESH ") and " ok" in r_label and "static" in r_label:
        return "ok_with_static_pinned"
    if "STATIC" in r_label or "static-pinned" in r_label:
        return "static_review"
    if "(partial)" in r_label:
        return "ok_with_partial"
    if "dynamic-added" in r_label:
        return "ok_with_dynamic_added"
    if r_label.startswith("REFRESH ok"):
        return "ok"
    if "no new data" in r_label:
        return "no_data"
    return "other"


def run_full_pipeline(
    source_pptx: str | Path,
    *,
    out_name: str | None = None,
    deck_out_dir: str | Path | None = None,
    json_out_dir: str | Path | None = None,
    force_fresh: bool = False,
    max_workers: int = 6,
) -> dict:
    """Run the full 5-step refresh pipeline against any deck.

    Returns a result dict with paths to every artifact and the slide
    bucket counts. See module docstring for output layout.
    """
    source_pptx = Path(source_pptx).resolve()
    if not source_pptx.exists():
        raise FileNotFoundError(f"Source deck missing: {source_pptx}")

    today = datetime.now().strftime("%Y-%m-%d")
    if out_name is None:
        out_name = f"{_slugify(source_pptx.stem)}_{today}_v1"

    deck_out_dir = Path(deck_out_dir) if deck_out_dir else (
        REPO_ROOT / "output_testing" / "deck_output"
    )
    json_out_dir = Path(json_out_dir) if json_out_dir else (
        REPO_ROOT / "output_testing" / "json_output"
    )
    deck_out_dir.mkdir(parents=True, exist_ok=True)
    json_out_dir.mkdir(parents=True, exist_ok=True)

    print(f"=== Refresh: {source_pptx.name} ({today}) ===\n")

    creds_ok, creds_msg = _check_synapse_creds()
    if not creds_ok:
        print(f"[ERROR] {creds_msg}")
        print("        Add SYNAPSE_API_URL and SYNAPSE_API_KEY to .env, or run")
        print("        `synapse login` to populate the JWT cache.")
        raise RuntimeError(creds_msg)
    if "fallback" in creds_msg:
        print(f"[warn] {creds_msg}")
    if not os.getenv("LLM_API_KEY"):
        print("[warn] LLM_API_KEY not set — Step 3 headlines will skip")

    spec_path = json_out_dir / f"{out_name}_full_spec.json"
    final_pptx = deck_out_dir / f"{out_name}.pptx"
    refresh_status_path = json_out_dir / f"{out_name}_refresh_status.json"
    headline_status_path = json_out_dir / f"{out_name}_headline_status.json"
    label_sync_status_path = json_out_dir / f"{out_name}_label_sync_status.json"

    work_dir = Path(tempfile.mkdtemp(prefix=f"{out_name}_"))
    refreshed_pptx = work_dir / f"{out_name}_step1.pptx"

    print(f"  source:     {source_pptx.relative_to(REPO_ROOT)}")
    print(f"  spec:       {spec_path.relative_to(REPO_ROOT)}")
    print(f"  final deck: {final_pptx.relative_to(REPO_ROOT)}")
    print()

    # ── Step 0 (implicit): Regenerate spec ──
    print("  Regenerating spec from source deck...")
    _regen_spec(source_pptx, spec_path)
    print(f"  spec -> {spec_path.relative_to(REPO_ROOT)}\n")

    # ── Step 1: Refresh data ──
    print("[1/5] refresh_deck_from_spec...")
    result = refresh_deck_from_spec(
        spec_path=str(spec_path),
        pptx_path=str(source_pptx),
        output_path=str(refreshed_pptx),
        force_fresh=force_fresh,
        max_workers=max_workers,
    )
    refresh_status_path.write_text(
        json.dumps(result, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )
    n_slides = len(result.get("slides", []))
    n_ok = sum(
        1 for s in result.get("slides", [])
        for c in (s.get("charts") or []) + (s.get("tables") or [])
        if c.get("status") == "ok"
    )
    n_with_notes = sum(
        1 for s in result.get("slides", [])
        for c in (s.get("charts") or []) + (s.get("tables") or [])
        if c.get("notes")
    )
    print(f"      {n_slides} slides processed, {n_ok} components ok, "
          f"{n_with_notes} with notes")
    print(f"      status -> {refresh_status_path.relative_to(REPO_ROOT)}")

    # ── Step 2: Stamp REFRESH_NOTE Connector tags ──
    print("\n[2/5] stamp_refresh_notes (per-shape Connector tags)...")
    stamp_results = stamp_refresh_notes(str(refreshed_pptx), result)
    n_stamped = sum(1 for v in stamp_results.values() if v == "stamped")
    n_skipped = sum(1 for v in stamp_results.values() if v != "stamped")
    print(f"      {n_stamped} shapes stamped, {n_skipped} skipped")

    # ── Step 3: Headlines (LLM) ──
    after_headlines = refreshed_pptx
    if os.getenv("LLM_API_KEY"):
        print("\n[3/5] refresh_headlines (LLM rewrite)...")
        from slidegen.headline_refresh import refresh_headlines
        try:
            with_headlines_pptx = work_dir / f"{out_name}_step3.pptx"
            updates = refresh_headlines(
                source_pptx=str(source_pptx),
                refreshed_pptx=str(refreshed_pptx),
                spec_path=str(spec_path),
                out_pptx=str(with_headlines_pptx),
            )
            updates_payload = {
                "updates": [
                    {
                        "slide_idx": u.slide_idx,
                        "old_headline": u.old_headline,
                        "new_headline": u.new_headline,
                        "chart_summary": u.chart_summary,
                        "status": u.status,
                    }
                    for u in updates
                ]
            }
            headline_status_path.write_text(
                json.dumps(
                    updates_payload, indent=2, ensure_ascii=False, default=str,
                ),
                encoding="utf-8",
            )
            n_updates = len(updates)
            n_updated = sum(1 for u in updates if u.status == "updated")
            print(f"      {n_updates} slides analyzed, "
                  f"{n_updated} headlines rewritten")
            print(f"      status -> {headline_status_path.relative_to(REPO_ROOT)}")
            after_headlines = with_headlines_pptx
        except Exception as exc:
            print(f"      [warn] headline rewrite failed: {exc}")
    else:
        print("\n[3/5] skipped (LLM_API_KEY not set)")

    # ── Step 4: Sync period labels ──
    print("\n[4/5] sync_period_labels (text frames adjacent to charts)...")
    after_label_sync = work_dir / f"{out_name}_step4.pptx"
    label_sync_diag = sync_period_labels(
        source_pptx=str(source_pptx),
        refreshed_pptx=str(after_headlines),
        out_pptx=str(after_label_sync),
        sidecar_path=str(label_sync_status_path),
    )
    n_shifted = label_sync_diag["totals"]["slides_with_shift"]
    n_edits = label_sync_diag["totals"]["edits"]
    n_conflicts = label_sync_diag["totals"]["conflicts"]
    print(f"      {n_shifted} slides shifted, {n_edits} run-level edits, "
          f"{n_conflicts} conflicts")
    print(f"      status -> {label_sync_status_path.relative_to(REPO_ROOT)}")
    after_headlines = after_label_sync  # feeds the badge step below

    # ── Step 5: Annotate badges ──
    print("\n[5/5] slide-level annotation badges...")
    from pptx import Presentation
    sys.path.insert(0, str(REPO_ROOT / "scripts"))
    from annotate_refresh_and_headline import (  # noqa: E402
        _refresh_badge, _headline_badge, _add_badge,
    )

    refresh_status = json.loads(refresh_status_path.read_text(encoding="utf-8"))
    refresh_by_idx = {s["slide_index"]: s for s in refresh_status.get("slides", [])}

    headline_by_idx: dict[int, dict] = {}
    if headline_status_path.exists():
        hl_status = json.loads(headline_status_path.read_text(encoding="utf-8"))
        headline_by_idx = {u["slide_idx"]: u for u in hl_status.get("updates", [])}

    pres = Presentation(str(after_headlines))
    slide_w = pres.slide_width

    counts: dict[str, int] = {}
    for s_idx, slide in enumerate(pres.slides):
        rs = refresh_by_idx.get(s_idx)
        hl = headline_by_idx.get(s_idx)
        r_color, r_label = _refresh_badge(rs)
        h_color, h_label = _headline_badge(hl)
        _add_badge(slide, r_color, r_label, slide_w, row=0)
        _add_badge(slide, h_color, h_label, slide_w, row=1)
        # Append per-component status to speaker notes so the analyst
        # can see WHICH components are refreshed vs static-pinned vs
        # tag-mismatch when the visible badge says e.g. "4 ok / 8 static".
        # Existing notes (Q text, sample sizes, etc.) preserved above
        # a divider.
        try:
            _append_refresh_notes(slide, rs, hl, today)
        except Exception:
            pass
        bucket = _bucket_label(r_label)
        counts[bucket] = counts.get(bucket, 0) + 1

    pres.save(str(final_pptx))

    try:
        shutil.rmtree(work_dir)
    except Exception:
        pass

    print(f"      annotated -> {final_pptx.relative_to(REPO_ROOT)}")

    # ── Summary ──
    print()
    print("=== Summary ===")
    print(f"  output deck: {final_pptx.relative_to(REPO_ROOT)}")
    print(f"  size:        {final_pptx.stat().st_size / 1024 / 1024:.1f} MB")
    print(f"  refresh JSON:    {refresh_status_path.relative_to(REPO_ROOT)}")
    if headline_status_path.exists():
        print(f"  headline JSON:   {headline_status_path.relative_to(REPO_ROOT)}")
    print(f"  label sync JSON: {label_sync_status_path.relative_to(REPO_ROOT)}")
    print("  slide buckets:")
    for k, v in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"    {k:25s} {v}")

    tag_mismatch_slides = [
        s_idx for s_idx in range(len(pres.slides))
        if any(
            n.get("kind") == "tag_mismatch"
            for c in (refresh_by_idx.get(s_idx, {}).get("charts") or [])
            + (refresh_by_idx.get(s_idx, {}).get("tables") or [])
            for n in (c.get("notes") or [])
        )
    ]
    if tag_mismatch_slides:
        print()
        print(f"  [REVIEW NEEDED] {len(tag_mismatch_slides)} slide(s) "
              f"with tag_mismatch:")
        for si in tag_mismatch_slides[:20]:
            shapes = [
                f"{c.get('name')} ({n.get('detail', '')[:80]})"
                for c in (refresh_by_idx.get(si, {}).get("charts") or [])
                for n in (c.get("notes") or [])
                if n.get("kind") == "tag_mismatch"
            ]
            print(f"    slide {si}: {shapes}")
        if len(tag_mismatch_slides) > 20:
            print(f"    ... and {len(tag_mismatch_slides) - 20} more")
        print("  Per slide: fix the Connector tag, or move to non-connected path.")

    return {
        "deck": str(final_pptx),
        "spec": str(spec_path),
        "refresh_status": str(refresh_status_path),
        "headline_status": str(headline_status_path) if headline_status_path.exists() else None,
        "label_sync_status": str(label_sync_status_path),
        "slide_buckets": counts,
        "tag_mismatch_slides": tag_mismatch_slides,
    }


def main():
    parser = argparse.ArgumentParser(
        prog="python -m slidegen.refresh_pipeline",
        description=(
            "End-to-end deck refresh: spec -> fetch -> stamp -> headlines "
            "-> label sync -> badges. Writes the final annotated PPTX to "
            "output_testing/deck_output/ by default."
        ),
    )
    parser.add_argument("deck", help="Path to source PPTX deck")
    parser.add_argument(
        "--out-name", default=None,
        help="Output base name (default: <slugified deck stem>_<date>_v1)",
    )
    parser.add_argument(
        "--deck-out-dir", default=None,
        help="Where to write the final PPTX (default: output_testing/deck_output/)",
    )
    parser.add_argument(
        "--json-out-dir", default=None,
        help="Where to write sidecar JSONs (default: output_testing/json_output/)",
    )
    parser.add_argument(
        "--no-cache", action="store_true",
        help="Bypass the on-disk Synapse cache (force fresh API fetch)",
    )
    parser.add_argument(
        "--workers", type=int, default=6,
        help="Max parallel data_source fetches (default: 6)",
    )
    args = parser.parse_args()

    run_full_pipeline(
        source_pptx=args.deck,
        out_name=args.out_name,
        deck_out_dir=args.deck_out_dir,
        json_out_dir=args.json_out_dir,
        force_fresh=args.no_cache,
        max_workers=args.workers,
    )


if __name__ == "__main__":
    main()
