"""
deck_assembler.py — Assembles a list of SlideSpecs into a PPTX deck.

Takes list[SlideSpec] + optional template PPTX path. Calls slide-creator per
spec (stubbed for now since slide-creator's Python implementation is not yet
built). Handles insertion/replacement/reordering. Outputs PPTX + updated
shape_registry.json.

Usage:
    from slidegen.deck_assembler import assemble_deck

    result = assemble_deck(specs, output_path="output/deck.pptx")
"""
from __future__ import annotations

import json
import os
from dataclasses import dataclass, field, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from pptx import Presentation
from pptx.util import Inches, Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from slidegen.slide_spec.schema import (
    SlideSpec, ChartComponent, LabelTableComponent,
    DeltaColumnComponent, ValueTableComponent,
    TextboxComponent, CalloutComponent, ImageComponent,
    dump_spec, SPEC_VERSION,
)
from slidegen.slide_spec.validator import validate_spec, SpecValidationError


@dataclass
class AssemblyResult:
    """Result of deck assembly."""
    output_path: str = ""
    registry_path: str = ""
    slides_rendered: int = 0
    slides_failed: int = 0
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


class DeckAssemblyError(Exception):
    """Raised when deck assembly fails."""
    pass


# ── Stub slide creator ───────────────────────────────────────────────────────
# NOTE: slide-creator's Python implementation is NOT YET BUILT.
# This is a minimal stub that creates a blank slide per spec and applies the
# headline text, so deck-assembler is testable end-to-end.


def _stub_render_slide(prs: Presentation, spec: SlideSpec, layout) -> Any:
    """Stub slide-creator: creates a blank slide with headline text.

    When the real slide-creator is built, this will be replaced by:
        from slidegen.slide_creator import render_slide
        return render_slide(spec, prs, namer)
    """
    slide = prs.slides.add_slide(layout)

    # Apply headline
    if spec.headline and spec.headline.text:
        shape = slide.shapes.add_textbox(
            Inches(0.20), Inches(0.16),
            Inches(12.08), Inches(1.21),
        )
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = spec.headline.text
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xF7, 0x58, 0x24)  # ZoomRx red

    # Apply subheadline
    if spec.subheadline and spec.subheadline.text:
        shape = slide.shapes.add_textbox(
            Inches(0.20), Inches(1.30),
            Inches(12.08), Inches(0.40),
        )
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = spec.subheadline.text
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    # Apply footer
    if spec.footer and spec.footer.text:
        shape = slide.shapes.add_textbox(
            Inches(0.20), Inches(7.10),
            Inches(12.70), Inches(0.25),
        )
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = spec.footer.text
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xA0, 0xA0, 0xA0)

    # Stub annotation
    shape = slide.shapes.add_textbox(
        Inches(1.0), Inches(3.0),
        Inches(11.0), Inches(2.0),
    )
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    component_types = [c.type for c in spec.components]
    run.text = (
        f"[STUB] Slide {spec.slide_index + 1}: {spec.slide_id}\n"
        f"Layout: {spec.layout}\n"
        f"Components: {component_types}\n"
        f"(slide-creator not yet implemented)"
    )
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

    return slide


# ── Main assembly ────────────────────────────────────────────────────────────


def assemble_deck(
    specs: list[SlideSpec],
    output_path: str | Path,
    template_path: Optional[str | Path] = None,
    registry_path: Optional[str | Path] = None,
    validate_specs: bool = True,
) -> AssemblyResult:
    """Assemble a deck from a list of SlideSpecs.

    Args:
        specs: List of SlideSpecs, ordered by slide_index.
        output_path: Path to write the output PPTX.
        template_path: Optional template PPTX (for slide master/layouts).
        registry_path: Path for shape_registry.json (default: alongside output).
        validate_specs: If True, validate each spec before rendering.

    Returns:
        AssemblyResult with paths and counts.
    """
    output_path = Path(output_path)
    result = AssemblyResult(output_path=str(output_path))

    if not specs:
        raise DeckAssemblyError("No specs provided — cannot assemble empty deck")

    # Sort specs by slide_index
    sorted_specs = sorted(specs, key=lambda s: s.slide_index)

    # Load or create presentation
    if template_path and Path(template_path).exists():
        prs = Presentation(str(template_path))
        # Clear existing slides
        for _ in range(len(prs.slides)):
            rId = prs.slides._sldIdLst[0].get(
                '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
            )
            prs.part.drop_rel(rId)
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
    else:
        prs = Presentation()
        prs.slide_width = Emu(12192000)   # 13.33 inches
        prs.slide_height = Emu(6858000)   # 7.50 inches

    # Get blank layout
    blank_layout = None
    for layout in prs.slide_layouts:
        if layout.name == "Blank":
            blank_layout = layout
            break
    if blank_layout is None:
        blank_layout = (
            prs.slide_layouts[6] if len(prs.slide_layouts) > 6
            else prs.slide_layouts[0]
        )

    # Build shape registry
    registry = {
        "version": "2.0",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "deck_path": str(output_path),
        "slides": {},
        "shapes": {},
    }

    # Render each spec
    for spec in sorted_specs:
        # Validate spec
        if validate_specs:
            errors = validate_spec(spec)
            if errors:
                result.errors.append(
                    f"Slide {spec.slide_index} ({spec.slide_id}): "
                    f"{len(errors)} validation error(s): {errors[0]}"
                )
                result.slides_failed += 1
                continue

        # Render slide (stub for now)
        try:
            slide = _stub_render_slide(prs, spec, blank_layout)
            result.slides_rendered += 1

            # Update registry
            registry["slides"][str(spec.slide_index)] = {
                "slide_id": spec.slide_id,
                "layout": spec.layout,
                "headline": spec.headline.text if spec.headline else "",
                "brand": spec.brand,
                "section": spec.section,
                "component_types": [c.type for c in spec.components],
                "rendered_at": datetime.now(timezone.utc).isoformat(),
            }

            # Add data lineage to registry
            if spec.data_lineage:
                dl = spec.data_lineage
                registry["slides"][str(spec.slide_index)]["data_lineage"] = {
                    "project_id": dl.project_id,
                    "reporting_plan_id": dl.reporting_plan_id,
                    "analysis_ids": dl.analysis_ids,
                    "data_source": dl.data_source,
                    "extraction_method": dl.extraction_method,
                    "question_codes": dl.question_codes,
                    "last_data_pull": dl.last_data_pull,
                    "config_hash": dl.config_hash,
                }

        except Exception as e:
            result.errors.append(
                f"Slide {spec.slide_index} ({spec.slide_id}): render failed: {e}"
            )
            result.slides_failed += 1

    # Write output PPTX
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))

    # Write registry
    if registry_path is None:
        registry_path = output_path.parent / "shape_registry.json"
    else:
        registry_path = Path(registry_path)

    registry_path.parent.mkdir(parents=True, exist_ok=True)
    result.registry_path = str(registry_path)
    registry_path.write_text(
        json.dumps(registry, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    return result


def replace_slide(
    specs: list[SlideSpec],
    target_index: int,
    new_spec: SlideSpec,
) -> list[SlideSpec]:
    """Replace a spec at a given index. Returns a new list."""
    result = []
    replaced = False
    for spec in specs:
        if spec.slide_index == target_index:
            new_spec_copy = new_spec
            new_spec_copy.slide_index = target_index
            result.append(new_spec_copy)
            replaced = True
        else:
            result.append(spec)

    if not replaced:
        result.append(new_spec)
        result.sort(key=lambda s: s.slide_index)

    return result


def insert_slide(
    specs: list[SlideSpec],
    new_spec: SlideSpec,
    after_index: int,
) -> list[SlideSpec]:
    """Insert a spec after a given index. Renumbers subsequent slides."""
    insert_at = after_index + 1
    new_spec.slide_index = insert_at

    result = []
    for spec in specs:
        if spec.slide_index >= insert_at:
            spec.slide_index += 1
        result.append(spec)

    result.append(new_spec)
    result.sort(key=lambda s: s.slide_index)
    return result


def reorder_slides(
    specs: list[SlideSpec],
    new_order: list[str],
) -> list[SlideSpec]:
    """Reorder specs by slide_id. new_order is a list of slide_id strings."""
    by_id = {s.slide_id: s for s in specs}
    result = []
    for i, sid in enumerate(new_order):
        if sid not in by_id:
            raise DeckAssemblyError(f"slide_id {sid!r} not found in specs")
        spec = by_id[sid]
        spec.slide_index = i
        result.append(spec)

    # Append any specs not in new_order
    remaining = [s for s in specs if s.slide_id not in set(new_order)]
    for spec in remaining:
        spec.slide_index = len(result)
        result.append(spec)

    return result
