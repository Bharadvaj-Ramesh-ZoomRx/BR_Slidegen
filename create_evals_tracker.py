import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# Colors
C_HEADER = 'FF1F3864'
C_STAGE  = 'FF2E75B6'
C_PASS   = 'FF92D050'
C_PROG   = 'FFFFC000'
C_NONE   = 'FFD9D9D9'
C_CRIT   = 'FFFF5050'
C_ALT    = 'FFF2F2F2'

thin = Side(border_style='thin', color='FFBFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def hcell(ws, row, col, text, bg=C_HEADER, fc='FFFFFFFF', align='center'):
    c = ws.cell(row=row, column=col, value=text)
    c.fill = PatternFill('solid', fgColor=bg)
    c.font = Font(color=fc, bold=True, size=11)
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    c.border = BORDER
    return c

def dcell(ws, row, col, text, bg=None, bold=False, align='left', fc='FF000000'):
    c = ws.cell(row=row, column=col, value=text)
    if bg:
        c.fill = PatternFill('solid', fgColor=bg)
    c.font = Font(bold=bold, color=fc, size=10)
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    c.border = BORDER
    return c

def col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def status_color(status):
    s = status.upper()
    if s == 'COVERED':     return C_PASS
    if s == 'IN PROGRESS': return C_PROG
    if s == 'NOT STARTED': return C_NONE
    if s == 'CRITICAL GAP': return C_CRIT
    return None

# ────────────────────────────────────────────────────────────────────────────
# SHEET 1: Dashboard — coverage % per pipeline stage
# ────────────────────────────────────────────────────────────────────────────
ws1 = wb.active
ws1.title = 'Dashboard'
ws1.sheet_view.showGridLines = False

hcell(ws1, 1, 1, 'REFRESH WORKFLOW — EVALS COVERAGE TRACKER', align='center')
ws1.merge_cells('A1:E1')
ws1.row_dimensions[1].height = 32

hcell(ws1, 2, 1, 'Purpose: tracks which parts of the refresh pipeline have automated evals and which do not.', bg='FFFFFFFF', fc='FF404040', align='left')
ws1.merge_cells('A2:E2')
ws1.row_dimensions[2].height = 24

for col, h in enumerate(['Pipeline Stage', 'Total Evals', 'Covered', 'Coverage %', 'Priority'], 1):
    hcell(ws1, 4, col, h)

stages = [
    ('Stage 0 — deck-reader (tag extraction, SlideSpec generation)',        6, 1, 'P0 (blocker for all downstream)'),
    ('Stage 1 — prior-wave-context-builder',                                3, 0, 'P2 (nice to have)'),
    ('Stage 2 — synapse-read (data fetch via API or banner plan)',          4, 0, 'P0 (blocked on token)'),
    ('Stage 3 — slide-plan-gen-refresh (DIFF: UPDATE/ADD/DELETE)',          5, 0, 'P1 (not built yet by Vijay)'),
    ('Stage 4 — slide-updater (in-place chart data replacement)',           6, 0, 'P0 (Vijay built, partially tested)'),
    ('Stage 4 — headline-writer (regenerate headlines on data change)',     4, 0, 'P1 (not built yet)'),
    ('Stage 5 — spec-validator (reject incomplete specs)',                  3, 0, 'P2'),
    ('Stage 5 — slide-creator / deck-assembler (rendering)',                4, 0, 'P1 (shared with new-deck)'),
    ('End-to-End — full refresh run (input deck + new data → output deck)', 5, 0, 'P0 (gold standard test)'),
]

for r, (stage, total, covered, priority) in enumerate(stages, 5):
    dcell(ws1, r, 1, stage)
    dcell(ws1, r, 2, total, align='center')
    cc = dcell(ws1, r, 3, covered, align='center')
    if covered == total and total > 0:
        cc.fill = PatternFill('solid', fgColor=C_PASS)
    elif covered > 0:
        cc.fill = PatternFill('solid', fgColor=C_PROG)
    else:
        cc.fill = PatternFill('solid', fgColor=C_NONE)
    pct = f'{(covered/total*100):.0f}%' if total > 0 else '—'
    pc = dcell(ws1, r, 4, pct, align='center', bold=True)
    if pct == '0%':
        pc.fill = PatternFill('solid', fgColor=C_CRIT)
    dcell(ws1, r, 5, priority)

total_total = sum(s[1] for s in stages)
total_covered = sum(s[2] for s in stages)
hcell(ws1, 14, 1, 'TOTAL', bg=C_STAGE)
dcell(ws1, 14, 2, total_total, align='center', bold=True)
dcell(ws1, 14, 3, total_covered, align='center', bold=True)
dcell(ws1, 14, 4, f'{(total_covered/total_total*100):.0f}%', align='center', bold=True, bg=C_CRIT)
dcell(ws1, 14, 5, '')

col_widths(ws1, [55, 12, 10, 12, 38])
ws1.freeze_panes = 'A5'

# ────────────────────────────────────────────────────────────────────────────
# SHEET 2: Eval Catalog — detailed spec of every eval
# ────────────────────────────────────────────────────────────────────────────
ws2 = wb.create_sheet('Eval Catalog')
ws2.row_dimensions[1].height = 38

headers = ['#', 'Pipeline Stage', 'Eval Name', 'What It Checks (Plain English)',
           'Input', 'Expected Output / Pass Criterion', 'Status', 'Priority',
           'Dependencies / Blockers', 'Notes']
for col, h in enumerate(headers, 1):
    hcell(ws2, 1, col, h)

evals = [
    # ── Stage 0: deck-reader ─────────────────────────────────────────────
    (1, 'deck-reader', 'Tag count per slide',
     'Runs deck-reader on a known deck and checks the count of tagged vs untagged shapes per slide against a committed golden file.',
     'ATU deck (99 slides), UAT deck (27 slides)',
     'Tagged/untagged counts exactly match golden snapshot for each slide',
     'COVERED', 'P0',
     'None — unblocked',
     'Validated 2026-04-21: positive test PASS on ATU; mutation test (3 corruptions) FAIL with all 3 reported; eval bug (short-circuit) caught and fixed. File: tests/evals/deck_reader/test_tag_counts.py'),
    (2, 'deck-reader', 'ReportConfig / PivotConfig / MappingConfig resolution',
     'Verifies that every tagged shape has its ReportConfigHash correctly resolved to the Custom XML Part containing the full JSON config, and that PivotConfig + MappingConfig parse without errors.',
     'ATU deck, UAT deck',
     'resolved = tagged for both ReportConfig and MappingConfig (allow up to 2% PivotConfig slippage per Vijay commit)',
     'NOT STARTED', 'P0',
     'None',
     'ATU: 197/199 pivot, 199/199 mapping resolved. UAT: golden baseline needed.'),
    (3, 'deck-reader', 'DataLineage field completeness',
     'Checks that every tagged slide has project_id, reporting_plan_id, analysis_ids, and survey_id populated in data_lineage. Flags any slide with missing fields.',
     'ATU deck',
     'All tagged slides: 4 required lineage fields populated',
     'NOT STARTED', 'P0',
     'None',
     'Current state: 55/55 tagged slides complete on ATU.'),
    (4, 'deck-reader', 'SlideSpec schema v1.2 compliance',
     'Validates every generated SlideSpec against the v1.2 schema. Catches new fields added to deck-reader that are not declared in schema.py.',
     'Any deck',
     'All specs parse via load_spec() without raising schema errors',
     'NOT STARTED', 'P0',
     'None',
     'Quick win — uses existing load_spec() from slide_spec/schema.py.'),
    (5, 'deck-reader', 'Golden snapshot regression',
     'Full spec JSON diff against committed golden. Catches ANY behavioral change in deck-reader between commits.',
     'ATU deck, UAT deck',
     'Generated specs == golden specs byte-for-byte (after normalization)',
     'NOT STARTED', 'P0',
     'None',
     'Strongest regression protection. Intentional changes require updating golden file.'),
    (6, 'deck-reader', 'Tier 2 inference confidence',
     'For untagged slides, inspects data_inference output and flags any inference with confidence below threshold (0.7). Measures Tier 2 reliability.',
     'ATU deck untagged slides (44 on ATU)',
     '>= 70% of untagged slides produce inferences at confidence ≥ 0.7',
     'NOT STARTED', 'P2',
     'data_inference.py is a prototype — low confidence is expected',
     'Sets a quality bar for Tier 2 before productionization.'),

    # ── Stage 1: prior-wave-context-builder ─────────────────────────────
    (7, 'prior-wave-context', 'Prior wave file detection',
     'Given a folder with prior wave files, verifies the skill picks up all relevant files (pptx reports, ES markdown, readout docx) and skips irrelevant ones.',
     'Fixture folder with mixed files',
     'File list matches expected set',
     'NOT STARTED', 'P2',
     'Requires /prior-wave-context skill running',
     'Lower priority — Bharadvaj ownership focuses on Tier 1 refresh path.'),
    (8, 'prior-wave-context', 'Finding extraction accuracy',
     'Checks extracted findings from prior wave ES against known ground truth (what the ES actually says).',
     'Prior wave ES fixture',
     'Extracted findings cover >= 90% of ground-truth topics',
     'NOT STARTED', 'P2',
     '/prior-wave-context skill',
     'Qualitative — uses LLM-as-judge or manual annotation.'),
    (9, 'prior-wave-context', 'Hand-off to Stage 2',
     'Verifies prior_wave_context.md produced by Stage 1 is parseable by hypothesis generator.',
     'Stage 1 output',
     'File loads without error; all expected sections present',
     'NOT STARTED', 'P2',
     '',
     'Mechanical check.'),

    # ── Stage 2: synapse-read ───────────────────────────────────────────
    (10, 'synapse-read', 'Data fetch for known analysis ID',
     'Given a known analysis_id + project_id, fetches data from Synapse and verifies the flat records structure (required fields present, row count matches expected).',
     'ATU analysis IDs (e.g., 338175)',
     'Returns list of dicts with expected keys; row count within tolerance',
     'NOT STARTED', 'P0',
     'BLOCKED on Synapse API token',
     'Once token available, this is the first eval to run.'),
    (11, 'synapse-read', 'Wave filtering correctness',
     'For an analysis with dynamic_latest_n=2, verifies the fetch returns exactly 2 most recent waves and no others.',
     'ATU lineage with dynamic_latest_n=2',
     'Returned records cover exactly 2 time periods',
     'NOT STARTED', 'P0',
     'BLOCKED on Synapse token',
     ''),
    (12, 'synapse-read', 'Segment filter application',
     'When lineage specifies segment_ids, verifies fetched records are filtered to those segments only.',
     'ATU lineage with segments',
     'Records scoped to specified segment_ids',
     'NOT STARTED', 'P0',
     'BLOCKED on Synapse token',
     ''),
    (13, 'synapse-read', 'Error handling — expired token',
     'With an expired token, verifies the fetch fails with a clear error and does not silently return empty results.',
     'Expired token fixture',
     'Raises TokenExpiredError with clear message',
     'NOT STARTED', 'P1',
     '',
     'Quick to build once any token path works.'),

    # ── Stage 3: slide-plan-gen-refresh ─────────────────────────────────
    (14, 'slide-plan-gen-refresh', 'UPDATE detection',
     'Given prior specs + new data where values changed, verifies the diff engine marks those slides as UPDATE.',
     'Prior specs (Q1) + new records (Q2)',
     'Affected slides classified UPDATE with delta per field',
     'NOT STARTED', 'P1',
     'Component does not yet exist — must build first',
     'Blocking for change classification.'),
    (15, 'slide-plan-gen-refresh', 'ADD detection',
     'When new data contains analyses not present in prior specs, verifies the engine proposes ADD with a suggested slide.',
     'Prior specs + extended new records',
     'New analyses tagged ADD',
     'NOT STARTED', 'P1',
     'Component does not yet exist',
     ''),
    (16, 'slide-plan-gen-refresh', 'DELETE detection',
     'When new data is missing an analysis present in prior specs, verifies the engine proposes DELETE.',
     'Prior specs + truncated new records',
     'Missing analyses tagged DELETE',
     'NOT STARTED', 'P1',
     'Component does not yet exist',
     ''),
    (17, 'slide-plan-gen-refresh', 'No-op detection',
     'When data is unchanged, verifies no slides are flagged for change.',
     'Prior specs + identical new records',
     'All slides classified UNCHANGED',
     'NOT STARTED', 'P1',
     '',
     'Important regression guard — unchanged slides must not rerender.'),
    (18, 'slide-plan-gen-refresh', 'Idempotency',
     'Running the diff engine twice with same inputs produces identical output.',
     'Same inputs, two runs',
     'output1 == output2',
     'NOT STARTED', 'P1',
     '',
     'Protects against nondeterminism (random ordering, timestamps, etc.).'),

    # ── Stage 4: slide-updater ──────────────────────────────────────────
    (19, 'slide-updater', 'Same-wave fidelity (round-trip)',
     'Takes source PPTX, wipes chart data, refreshes using same Q1 data, compares to source. This is Vijay Stage 3 turned into a repeatable eval.',
     'Source deck + SlideSpecs',
     '>= 95% of slides byte-equivalent in chart data',
     'IN PROGRESS', 'P0',
     'Vijay test_spec_refresh_pipeline.py already covers this on UAT — port to framework',
     'UAT result: 25/27 = 92.6% visually correct. Make this automated.'),
    (20, 'slide-updater', 'Cross-wave update (Q1→Q2)',
     'Refresh Q1 deck with Q2 data, verify chart values match Q2 data (not Q1).',
     'Q1 deck + Q2 banner plan / API',
     'Every refreshed chart reflects Q2 values within tolerance',
     'NOT STARTED', 'P0',
     'BLOCKED on Synapse token OR analysis→question mapping',
     'The critical test. Until this works, refresh is not production-ready.'),
    (21, 'slide-updater', 'Layout preservation',
     'Chart types, shape positions, fonts, colors preserved after refresh.',
     'Source vs refreshed PPTX',
     'Shape-level diff: only chart data changed; layout unchanged',
     'NOT STARTED', 'P0',
     '',
     'Vijay noted 2/27 slides had layout issues — this would catch those.'),
    (22, 'slide-updater', 'Tables refresh correctly',
     'Label tables and value tables paired with charts get updated with new data.',
     'Source vs refreshed PPTX',
     'Table values match new data',
     'NOT STARTED', 'P0',
     '',
     'Vijay commit: tables restored from source after refresh.'),
    (23, 'slide-updater', 'Split visualization handling',
     'Slides with splitGroupID / splitOrder / rowsPerObject render split correctly.',
     'Source with split viz fixtures',
     'Split order preserved; rows distributed correctly',
     'NOT STARTED', 'P1',
     '',
     'Vijay commit calls this out explicitly.'),
    (24, 'slide-updater', 'Formatting restoration',
     'Per-element formatCode (%, $, #) restored correctly on value cells after replace_data() wipes them.',
     'Source with mixed format codes',
     'All cells retain correct number formatting',
     'NOT STARTED', 'P0',
     '',
     'Vijay commit mentions this as a specific fix.'),

    # ── Stage 4: headline-writer ────────────────────────────────────────
    (25, 'headline-writer', 'Freshness when data changes',
     'When underlying data changes significantly (delta > threshold), headline updates. When data is stable, headline unchanged.',
     'Prior + new data with known deltas',
     'Headline regenerated iff max_abs_delta > threshold',
     'NOT STARTED', 'P1',
     'Component not built yet',
     'The PRD v1.1 "headline freshness" principle — core correctness.'),
    (26, 'headline-writer', 'Factual grounding',
     'Regenerated headline mentions values/directions consistent with the actual data (no hallucinated numbers).',
     'New data + generated headline',
     'All numeric claims in headline match source data within tolerance',
     'NOT STARTED', 'P1',
     '',
     'Critical for client-ready decks.'),
    (27, 'headline-writer', 'Length + style constraints',
     'Headlines conform to length limit, no forbidden words, proper capitalization.',
     'Any generated headline',
     'Length ≤ N chars; no banned terms; sentence-case',
     'NOT STARTED', 'P2',
     '',
     'Simple pattern assertions.'),
    (28, 'headline-writer', 'Arc consistency',
     'Headlines across slides in the same story arc stay thematically consistent.',
     'Deck with arc tagging',
     'Arc-level coherence score (manual or LLM-judge)',
     'NOT STARTED', 'P2',
     '',
     'Qualitative — lower priority.'),

    # ── Stage 5: spec-validator ─────────────────────────────────────────
    (29, 'spec-validator', 'Reject incomplete specs',
     'Specs missing required fields (headline, chart type, data) are rejected before rendering.',
     'Malformed spec fixtures',
     'Validator raises SpecInvalidError with clear field name',
     'NOT STARTED', 'P2',
     '',
     'Defensive — guards the rendering pipeline.'),
    (30, 'spec-validator', 'Accept minimal valid spec',
     'Smallest valid spec passes without error.',
     'Minimal spec fixture',
     'Validator returns success',
     'NOT STARTED', 'P2',
     '',
     ''),
    (31, 'spec-validator', 'Schema evolution safety',
     'Older-version specs either pass or produce a helpful error ("please re-run deck-reader").',
     'Old-version spec fixture',
     'Either succeeds or clear error',
     'NOT STARTED', 'P2',
     '',
     ''),

    # ── Stage 5: slide-creator / deck-assembler ─────────────────────────
    (32, 'slide-creator', 'Renderer mapping',
     'Every chart_pattern in a spec maps to a valid renderer in RENDERERS registry.',
     'Any spec',
     'No KeyError raised during rendering',
     'NOT STARTED', 'P1',
     '',
     'Shared with new-deck workflow — tests Pradeep domain too.'),
    (33, 'slide-creator', 'Brand config application',
     'Colors, fonts, sizes match BRAND{} config per project.',
     'Rendered slide + BRAND config',
     'Font families and colors match spec',
     'NOT STARTED', 'P1',
     '',
     ''),
    (34, 'deck-assembler', 'Slide order correctness',
     'Final PPTX slides ordered per spec (arc-first, then intra-arc).',
     'List of SlideSpecs + assembled deck',
     'Slide order matches spec order',
     'NOT STARTED', 'P1',
     '',
     ''),
    (35, 'deck-assembler', 'shape_registry.json completeness',
     'Every rendered shape appears in shape_registry with correct lineage + timestamps.',
     'Assembled deck + registry',
     'Registry has entry per named shape; last_refreshed populated',
     'NOT STARTED', 'P1',
     '',
     'Registry is the audit trail.'),

    # ── End-to-End ──────────────────────────────────────────────────────
    (36, 'end-to-end', 'Full refresh on ATU deck',
     'Run the complete refresh workflow (deck-reader → synapse-read → slide-updater → assembler) on ATU Q1 deck with Q2 data. Compare output to a frozen expected deck.',
     'ATU Q1 deck + Q2 data source',
     'Output deck byte-level match to expected (chart data + headlines)',
     'NOT STARTED', 'P0',
     'BLOCKED on Synapse token OR banner plan mapping',
     'Gold standard end-to-end test.'),
    (37, 'end-to-end', 'Full refresh on UAT deck (regression)',
     'Port Vijay 3-stage test into evals. Lock in his 25/27 baseline and catch any regression.',
     'UAT source deck',
     '>= 25/27 slides visually correct (or improve)',
     'NOT STARTED', 'P0',
     'Vijay already proved this once — just needs automation',
     'Quickest path to end-to-end eval coverage.'),
    (38, 'end-to-end', 'Idempotency — same inputs, same output',
     'Two runs with same inputs produce identical output decks.',
     'Same deck + same data, two runs',
     'Byte-level diff = 0',
     'NOT STARTED', 'P1',
     '',
     'Nondeterminism bugs are subtle — this eval catches them.'),
    (39, 'end-to-end', 'Performance baseline',
     'Full refresh run completes within a time budget. Regressions flag performance drops.',
     'ATU deck, Q2 data',
     'Total runtime < 5 minutes (initial target)',
     'NOT STARTED', 'P2',
     '',
     'Tune threshold after baseline established.'),
    (40, 'end-to-end', 'Graceful failure modes',
     'When Synapse is down, token is expired, or banner plan is missing, the pipeline fails with clear errors and no partial output.',
     'Fault-injection fixtures',
     'Correct error raised; no half-written deck on disk',
     'NOT STARTED', 'P1',
     '',
     'Production readiness.'),
]

for r, row in enumerate(evals, 2):
    bg = None if r % 2 == 0 else C_ALT
    for c, val in enumerate(row, 1):
        cell = dcell(ws2, r, c, val, bg=bg)
        if c == 7:  # Status column
            sc = status_color(str(val))
            if sc:
                cell.fill = PatternFill('solid', fgColor=sc)
                cell.font = Font(bold=True, size=10)
        if c == 8:  # Priority
            if 'P0' in str(val):
                cell.font = Font(bold=True, color='FFC00000', size=10)
            elif 'P1' in str(val):
                cell.font = Font(bold=True, color='FFED7D31', size=10)

col_widths(ws2, [5, 22, 38, 55, 32, 45, 14, 12, 32, 42])
ws2.freeze_panes = 'E2'

# ────────────────────────────────────────────────────────────────────────────
# SHEET 3: Legend / How to Use
# ────────────────────────────────────────────────────────────────────────────
ws3 = wb.create_sheet('Legend')
ws3.sheet_view.showGridLines = False

hcell(ws3, 1, 1, 'HOW TO USE THIS TRACKER', align='center')
ws3.merge_cells('A1:B1')
ws3.row_dimensions[1].height = 28

rows = [
    ('', ''),
    ('Status values', ''),
    ('  COVERED',      'Eval is written, passing, and runs in CI / tests/evals/'),
    ('  IN PROGRESS',  'Eval partially exists or is being built now'),
    ('  NOT STARTED',  'Eval is designed but no code yet'),
    ('  CRITICAL GAP', 'Eval is needed and should be built ASAP'),
    ('', ''),
    ('Priority levels', ''),
    ('  P0', 'Must have. Blocks production / trust. Build first.'),
    ('  P1', 'Should have. Important but not blocking.'),
    ('  P2', 'Nice to have. Defensive or qualitative.'),
    ('', ''),
    ('How to update', ''),
    ('  1', 'When you finish an eval, change its Status to COVERED on the Eval Catalog sheet.'),
    ('  2', 'Update the Covered count on the Dashboard sheet for that stage.'),
    ('  3', 'If you identify a new eval, add a new row to Eval Catalog; increment Total on Dashboard.'),
    ('  4', 'Commit the updated xlsx alongside the eval code so the tracker stays in sync.'),
    ('', ''),
    ('Folder convention', ''),
    ('  tests/evals/deck_reader/',           'evals for Stage 0'),
    ('  tests/evals/synapse_read/',          'evals for Stage 2'),
    ('  tests/evals/slide_updater/',         'evals for Stage 4 refresh engine'),
    ('  tests/evals/headline_writer/',       'evals for Stage 4 headline'),
    ('  tests/evals/slide_plan_gen_refresh/','evals for Stage 3 diff engine'),
    ('  tests/evals/end_to_end/',            'evals for full refresh runs'),
    ('  tests/evals/fixtures/',              'shared input decks, golden files, expected outputs'),
    ('  tests/evals/run_evals.py',           'CLI runner — python tests/evals/run_evals.py [--stage X]'),
]

for r, (k, v) in enumerate(rows, 2):
    bold = k in ('Status values', 'Priority levels', 'How to update', 'Folder convention')
    dcell(ws3, r, 1, k, bold=bold, bg='FFE7E6E6' if bold else None)
    dcell(ws3, r, 2, v)

col_widths(ws3, [30, 90])

wb.save('evals_coverage_tracker.xlsx')
print('Saved: evals_coverage_tracker.xlsx')
